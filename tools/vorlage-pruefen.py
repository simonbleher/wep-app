#!/usr/bin/env python3
"""Prüft die Excel-Vorlage des BV Pankow (docs/WEP_Vorlage.xlsx) auf Fehler.

Das Skript beantwortet zwei Fragen:

1. Welche der im Fehlerbericht (docs/VORLAGE-FEHLERBERICHT.md) beschriebenen Funde sind in einer
   Datei **noch vorhanden**? Damit lässt sich eine neue Fassung der Vorlage prüfen
   (Ausgabe „VORHANDEN“ / „behoben“ / „nicht automatisch prüfbar“ je Fund).
2. Wie lauten die Rechenergebnisse der Vorlage bei erfundenen Testeingaben, und stimmen sie mit einer
   unabhängigen Handrechnung (Python) überein?

Technik: Die Datei wird **nie verändert**. Für die Rechentests schreibt das Skript Testeingaben in
Kopien (openpyxl), lässt sie mit LibreOffice (headless) neu berechnen und vergleicht jede Zelle mit der
Handrechnung. Die statischen Prüfungen lesen nur Formeln, Zahlenformate, Sperren, Auswahllisten und Texte.

Aufruf (aus dem Repository-Ordner):
  python tools/vorlage-pruefen.py                          # prüft docs/WEP_Vorlage.xlsx
  python tools/vorlage-pruefen.py NEUE_VORLAGE.xlsx        # prüft eine andere Fassung
  python tools/vorlage-pruefen.py --ohne-rechentests       # nur statische Prüfungen (ohne LibreOffice)
  python tools/vorlage-pruefen.py --bericht docs/VORLAGE-FEHLERBERICHT.md [--pdf docs/VORLAGE-FEHLERBERICHT.pdf]
  python tools/vorlage-pruefen.py --korrekturvorschlag docs/WEP_Vorlage_Korrekturvorschlag.xlsx

Voraussetzungen: Python 3.10+, `pip install openpyxl lxml`; für Rechentests LibreOffice (`soffice`),
für --pdf zusätzlich `pip install playwright` (Chromium). Siehe docs/VORLAGE-FEHLERBERICHT.md, Anhang.
Der Exit-Code ist die Anzahl der noch vorhandenen Funde mit Gewissheit „sicher“ (0 = alles behoben).
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import html
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Callable

import openpyxl

__version__ = "1.0"

REPO = Path(__file__).resolve().parent.parent
STANDARD_VORLAGE = REPO / "docs" / "WEP_Vorlage.xlsx"

# Blattnamen der Vorlage (beachte das Leerzeichen am Ende bei „Deckbl UPV “)
FO, DB, BL, MG, AA = "Formalien", "Deckbl UPV ", "Baulichkeiten", "Mängelliste", "Außenanlagen"
AW1, AW2, BE, SK2, T1 = "Aufwuchs I", "Aufwuchs II", "Beseitigungskosten", "Skizze Außenanlage-Aufwuchs", "Tabelle1"

AUSWIRKUNGEN = ["falscher Betrag", "falsche oder fehlende Angabe", "Bedienung", "kosmetisch"]
GEWISSHEITEN = ["sicher", "wahrscheinlich", "bitte BV fragen"]


# ---------------------------------------------------------------------------------------------------
# 1. Referenzdaten aus der Richtlinie (BewRL, LV Berlin, gültig ab 01.01.2023)
# ---------------------------------------------------------------------------------------------------

# Richtlinie S. 10, „Restnutzungsjahre der Baulichkeiten“, Spalte „Abschr in %“ (Index = Restnutzungsjahre).
# Von Hand aus dem Bild der Seite übertragen und gegen Tabelle1 der Vorlage geprüft.
RN_IV = [100, 98, 96, 94, 92, 90, 86.2, 82.4, 78.6, 74.8, 71, 67.6, 64.2, 60.8, 57.4, 54, 51.2, 48.4, 45.6, 42.8,
         40, 37.8, 35.6, 33.4, 31.2, 29, 27.4, 25.8, 24.2, 22.6, 21, 20.1, 19.2, 18.3, 17.4, 16.5, 15.2, 13.9,
         12.6, 11.3, 10]
RN_III = [100, 98, 96, 94, 92, 90, 89, 88, 87, 86, 85, 82.2, 79.4, 76.6, 73.8, 71, 68.8, 66.6, 64.4, 62.2, 60,
          57.8, 55.6, 53.4, 51.2, 49, 47.2, 45.4, 43.6, 41.8, 40, 38.4, 36.8, 35.2, 33.6, 32, 30.8, 29.6, 28.4,
          27.2, 26, 25, 24, 23, 22, 21, 20.4, 19.8, 19.2, 18.6, 18, 17.5, 17, 16.5, 16, 15.5, 14.4, 13.3, 12.2,
          11.1, 10]
RN_I = [100, 98, 96, 94, 92, 90, 89.8, 89.6, 89.3, 88.9, 88.5, 88.2, 86.4, 84.6, 82.8, 81, 79, 77, 75, 73, 71,
        69.2, 67.4, 65.6, 63.8, 62, 60.4, 58.8, 57.2, 55.6, 54, 52.6, 51.2, 49.8, 48.4, 47, 45.6, 44.2, 42.8,
        41.4, 40, 38.8, 37.6, 36.4, 35.2, 34, 33, 32, 31, 30, 29, 28.2, 27.4, 26.6, 25.8, 25, 24.2, 23.4, 22.6,
        21.8, 21, 20.3, 19.6, 18.9, 18.2, 17.5, 17.3, 17.1, 16.9, 16.7, 16.5, 16.3, 16.1, 15.9, 15.7, 15.5, 14.4,
        13.3, 12.2, 11.1, 10]
RN_TABELLEN = {"I": RN_I, "II/III": RN_III, "IV": RN_IV}

# Richtlinie S. 5: NHK 2015 je Bauklasse/Ausführung (€/m²); Spalte der Restnutzungstabelle in Klammern
NHK = {"I.1": 515, "I.2": 595, "I.3": 704, "II.1": 350, "II.3": 490, "III.1": 420, "III.2": 555, "III.3": 672,
       "IV.1": 265, "IV.2": 325, "IV.3": 420}
# Richtlinie S. 19: Abriss je m³ umbauten Raumes
ABRISS = {"I": 180, "II": 180, "III": 150, "IV": 120}

# Richtlinie S. 6–8: Mängeltabelle – Bauteil → (Abzug maximal in %, Mängel)
MAENGEL_RICHTLINIE: dict[str, tuple[float, list[str]]] = {
    "Dach": (20, ["Sparren – Abstand nicht fachgerecht, größer als 60 - 80 cm", "Sparren – Risse durch Belastung",
                  "Dach nicht horizontal", "Regenrinne oder Fallrohr nicht fachgerecht",
                  "Fallrohr ohne Ableitung von der Laube", "Dachoberfläche beschädigt",
                  "Dachoberfläche nicht fachgerecht", "keine Traufkante", "Ortgang fehlerhaft oder fehlt",
                  "Dachfenster / Dachluke fehlerhaft", "muffiger Geruch unter dem Dach",
                  "Feuchtigkeit in der Dachkonstruktion", "Schornstein fehlerhaft",
                  "Schornstein ohne Kehrbescheinigung", "Moose und Flechten"]),
    "Decken": (10, ["Decke hängt durch", "Risse", "Feuchtigkeitsflecken", "loser Putz / Hohlstellen",
                    "Verkleidung defekt", "Anstrich"]),
    "Wände": (13, ["Rohmauerwerk", "Hohlstellen im Putz", "Risse", "defekte Verkleidung", "Feuchtigkeitsflecken",
                   "Horizontalsperre defekt", "Wände nicht lotrecht", "keine fachgerechte Ausführung",
                   "offene Fugen"]),
    "Fußboden": (10, ["Hohlstellen im Estrich", "Risse", "Dielung gibt nach", "Fußboden nicht waagerecht",
                      "Feuchtigkeitsflecken", "Belag verschlissen"]),
    "Fundament": (20, ["Fundamenttiefe nicht erkennbar", "Setzungen", "Brüche",
                       "kein Spritzwassersockel (mindestens 25 cm)", "Spritzwassersockel fehlerhaft",
                       "nicht maßgerecht", "keine fachgerechte Ausführung (Festigkeit)"]),
    "Ausbau außen": (5, ["Anstrich fehlerhaft", "Verkleidung fehlerhaft", "Regenrinne fehlerhaft",
                         "Fallrohre fehlerhaft"]),
    "Ausbau innen": (7, ["Rückflussverhinderer fehlt", "Anstrich, Tapeten, Verkleidung fehlerhaft",
                         "Fußbodenbelag fehlerhaft", "Treppen, Leitern fehlerhaft", "Fliesen fehlerhaft",
                         "Wasseranschluss fehlerhaft (nur bei vorhandener Abwassersammelanlage)"]),
    "Bauelemente": (5, ["Fenster – Anstrich defekt", "Fenster – Glasdichtung fehlerhaft (Kit)", "Fenster klemmt",
                        "Fenster undicht", "Fensterläden nicht ordnungsgemäß", "Fensterläden – Anstrich defekt",
                        "Tür nicht für Außenbereich geeignet", "Tür klemmt", "Tür undicht", "Tür – Anstrich defekt",
                        "Einbruchschäden"]),
    "Schädlinge": (10, ["Schimmel", "Marder, Ratten", "Hausschwamm", "Taubenkot"]),
    "Sonderabfall": (10, ["Asbest", "Dachpappe"]),
}
# Zeilen dieser Bauteile im Blatt „Mängelliste“ der Vorlage (Spalten C und D)
MAENGEL_ZEILEN = {"Dach": (5, 12), "Decken": (13, 15), "Wände": (16, 20), "Fußboden": (21, 23),
                  "Fundament": (24, 27), "Ausbau außen": (28, 29), "Ausbau innen": (30, 33),
                  "Bauelemente": (34, 39), "Schädlinge": (40, 41), "Sonderabfall": (42, 42)}
MAENGEL_MAX_ZELLEN = {"Dach": "E5", "Decken": "E13", "Wände": "E16", "Fußboden": "E21", "Fundament": "E24",
                      "Ausbau außen": "E28", "Ausbau innen": "E30", "Bauelemente": "E34", "Schädlinge": "E40",
                      "Sonderabfall": "E42"}

# Bekannte Tippfehler der Vorlage (Falschschreibung → richtig). Gesucht wird in allen Textzellen.
TIPPFEHLER = {
    "horizentral": "horizontal", "maßgrecht": "maßgerecht", "Gladichtung": "Glasdichtung",
    "nich ordnungsgemäß": "nicht ordnungsgemäß", "unddicht": "undicht", "enstsprechend": "entsprechend",
    "Sepzember": "September", "entfernene": "entfernende", "Anfalll": "Anfall", "Aussenhecke": "Außenhecke",
    "lotgerecht": "lotrecht",
}


# ---------------------------------------------------------------------------------------------------
# 2. Hilfen: Mappe lesen, runden, Zellen vergleichen
# ---------------------------------------------------------------------------------------------------

def rd(x: float, n: int = 2) -> float:
    """Kaufmännisch runden wie Excel ROUND (0,5 immer weg von null)."""
    q = Decimal(1).scaleb(-n)
    return float(Decimal(repr(round(x, 10))).quantize(q, rounding=ROUND_HALF_UP))


def norm(formel: Any) -> str:
    """Formeltext ohne „=“, Leerzeichen und $ – zum Vergleichen."""
    if not isinstance(formel, str):
        return ""
    return formel.lstrip("=").replace("$", "").replace(" ", "").upper()


def ist_formel(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


FUNKTIONEN_DE = {"SUM": "SUMME", "ROUND": "RUNDEN", "IF": "WENN", "VLOOKUP": "SVERWEIS", "IFERROR": "WENNFEHLER"}


def de(formel: str) -> str:
    """Formel in deutscher Excel-Schreibweise (Funktionsnamen deutsch, Semikolon, Dezimalkomma) für den Bericht.
    Die Datei selbst speichert Formeln englisch; Texte in Anführungszeichen bleiben unverändert."""
    teile = re.split(r'("[^"]*")', formel)
    for i in range(0, len(teile), 2):
        s = re.sub(r"[A-Z]+(?=\()", lambda m: FUNKTIONEN_DE.get(m.group(0), m.group(0)), teile[i])
        s = s.replace(",", ";")
        teile[i] = re.sub(r"(?<=\d)\.(?=\d)", ",", s)
    return "".join(teile)


def fx(v: Any) -> str:
    """Zellinhalt für den Bericht: Formeln deutsch, alles andere als Text."""
    return de(v) if ist_formel(v) else str(v)


def fast_gleich(a: Any, b: Any, tol: float = 0.0051) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tol
    return a == b


class Mappe:
    """Die zu prüfende Datei (nur lesen)."""

    def __init__(self, pfad: Path):
        self.pfad = Path(pfad)
        self.wb = openpyxl.load_workbook(self.pfad)
        self.zip = zipfile.ZipFile(self.pfad)

    def blatt(self, name: str):
        if name not in self.wb.sheetnames:
            raise KeyError(f"Blatt „{name}“ fehlt")
        return self.wb[name]

    def roh(self, blatt: str, zelle: str) -> Any:
        return self.blatt(blatt)[zelle].value

    def formel(self, blatt: str, zelle: str) -> str:
        return norm(self.roh(blatt, zelle))

    def typ_name(self, text: str) -> str:
        """Schreibweise eines Bauteil-Typs wie in der Auswahlliste (Tabelle1 J54:J75, teils mit Leerzeichen am Ende)."""
        ws = self.blatt(T1)
        for r in range(54, 76):
            v = ws[f"J{r}"].value
            if isinstance(v, str) and v.strip() == text:
                return v
        raise KeyError(f"Typ „{text}“ nicht in Tabelle1")

    def texte(self, nur_sichtbar: bool = False):
        for ws in self.wb:
            if nur_sichtbar and ws.sheet_state != "visible":
                continue
            for zeile in ws.iter_rows():
                for c in zeile:
                    if isinstance(c.value, str) and not c.value.startswith("="):
                        yield ws.title, c.coordinate, c.value

    def blattschutz(self) -> dict[str, str]:
        """Blatt → Zustand des Blattschutzes aus dem XML („aktiv“, „aus (Element ohne sheet=1)“, „nicht vorhanden“)."""
        rels = {}
        wbx = self.zip.read("xl/workbook.xml").decode("utf8")
        relx = self.zip.read("xl/_rels/workbook.xml.rels").decode("utf8")
        for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', relx):
            rels[m.group(1)] = m.group(2)
        for m in re.finditer(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', relx):
            rels.setdefault(m.group(2), m.group(1))
        res = {}
        for m in re.finditer(r'<sheet name="([^"]+)"[^>]*?r:id="([^"]+)"', wbx):
            name = html.unescape(m.group(1))
            ziel = rels.get(m.group(2), "")
            teil = "xl/" + ziel.lstrip("/").removeprefix("xl/")
            xml = self.zip.read(teil).decode("utf8")
            sp = re.search(r"<sheetProtection[^>]*>", xml)
            if not sp:
                res[name] = "nicht vorhanden"
            elif re.search(r'\bsheet="(1|true)"', sp.group(0)):
                res[name] = "aktiv"
            else:
                res[name] = "aus"
        return res


# ---------------------------------------------------------------------------------------------------
# 3. Rechentests: Testeingaben in Kopien, LibreOffice rechnet neu, Handrechnung vergleicht
# ---------------------------------------------------------------------------------------------------

def finde_libreoffice(vorgabe: str | None = None) -> str | None:
    kandidaten = [vorgabe] if vorgabe else []
    kandidaten += [shutil.which("soffice"), shutil.which("libreoffice"),
                   r"C:\Program Files\LibreOffice\program\soffice.exe",
                   r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                   "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for k in kandidaten:
        if k and Path(k).exists():
            return k
        if k and shutil.which(k):
            return shutil.which(k)
    return None


@dataclass
class Szenario:
    name: str
    beschreibung: str
    eingaben: dict[tuple[str, str], Any]
    lesen: list[tuple[str, str]]
    ergebnis: dict[tuple[str, str], Any] = field(default_factory=dict)
    fehlerzellen: list[str] = field(default_factory=list)   # Zellen mit Fehlerwert in sichtbaren Blättern


class Lauf:
    """Alle Rechenszenarien einer Datei. Ergebnisse stehen nach `rechne()` in `szenario.ergebnis`."""

    def __init__(self, mappe: Mappe, lo: str):
        self.mappe, self.lo = mappe, lo
        self.szenarien: dict[str, Szenario] = {}

    def neu(self, name: str, beschreibung: str, eingaben: dict, lesen: list) -> Szenario:
        s = Szenario(name, beschreibung, eingaben, lesen)
        self.szenarien[name] = s
        return s

    def rechne(self) -> None:
        tmp = Path(tempfile.mkdtemp(prefix="vorlage-pruefen-"))
        try:
            quelle = tmp / "quelle"
            quelle.mkdir()
            ausgabe = tmp / "ausgabe"
            ausgabe.mkdir()
            dateien = []
            for s in self.szenarien.values():
                wb = openpyxl.load_workbook(self.mappe.pfad)
                for (blatt, zelle), wert in s.eingaben.items():
                    wb[blatt][zelle] = wert
                p = quelle / f"{s.name}.xlsx"
                wb.save(p)
                dateien.append(str(p))
            profil = (tmp / "profil").as_uri()
            cmd = [self.lo, f"-env:UserInstallation={profil}", "--headless", "--calc",
                   "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(ausgabe), *dateien]
            subprocess.run(cmd, capture_output=True, timeout=900, check=False)
            for s in self.szenarien.values():
                p = ausgabe / f"{s.name}.xlsx"
                if not p.exists():
                    raise RuntimeError(f"LibreOffice hat {s.name} nicht berechnet")
                w2 = openpyxl.load_workbook(p, data_only=True)
                s.ergebnis = {(b, z): w2[b][z].value for b, z in s.lesen}
                s.fehlerzellen = [f"{ws.title}!{c.coordinate}" for ws in w2 if ws.sheet_state == "visible"
                                  for zeile in ws.iter_rows() for c in zeile
                                  if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def __getitem__(self, name: str) -> dict[tuple[str, str], Any]:
        return self.szenarien[name].ergebnis


BASIS_FORMALIEN = {
    (FO, "E3"): "Kleingartenanlage Musterwiese", (FO, "E4"): "12", (FO, "E6"): "Erika Musterfrau",
    (FO, "E12"): "Max Mustermann", (FO, "E13"): "Maria Beispiel", (FO, "E14"): dt.date(2026, 10, 1),
    (FO, "E15"): dt.date(2026, 10, 5), (FO, "E16"): 300, (FO, "E17"): "2026-999",
}

# Bauteile der Vorlage: Kürzel → (Zellen Typ, Restnutzungsjahre, Mängel %, m², Ergebniszelle Betrag)
BAUTEIL_ZELLEN = {
    "L": ("G12", "I13", "E14", "A17", "E17"), "A1": ("G21", "I22", "E23", "A26", "E26"),
    "A2": ("G30", "I31", "E32", "A35", "E35"), "S": ("G47", "I48", "E49", "A52", "E52"),
}
# Fall T1: Bauklasse-Spalte der Restnutzungstabelle und NHK-Stufe je Bauteil
T1_BAUTEILE = {
    "L": ("Doppelt Holz 10,01 - 14,99 cm", "III.1", "II/III", 20, 10, 20),
    "A1": ("Stein 10,00 - 12,00 cm", "II.1", "II/III", 15, 5, 3),
    "A2": ("Doppelwand 4,01 - 8,00 cm", "IV.2", "IV", 10, 0, 1),
    "S": ("Blockhaus bis 3,00 cm", "IV.1", "IV", 8, 20, 4),
}


def abwasser_abzug(alter: float, modus: str) -> float:
    """Altersabzug von Grube/Kanal in Prozent: „linear“ = Alter × 3 %, „restwert“ = (1 − 0,97^Alter) × 100 (siehe Fund F29)."""
    return min(100.0, alter * 3) if modus == "linear" else (1 - 0.97 ** alter) * 100


def grube_betrag(alter: float, m3: float, grundwert: float, modus: str) -> float:
    return rd(rd(grundwert - grundwert * abwasser_abzug(alter, modus) / 100) * m3)


def kanal_betrag(alter: float, kosten: float, modus: str) -> float:
    e79 = min(2200.0, kosten)
    return rd(e79 - e79 * abwasser_abzug(alter, modus) / 100)


def altersabzug_tabelle(spalte: str, rn: int) -> float:
    """Abschreibung in % laut Richtlinie. Jenseits des Tabellenendes: 0 (so rechnet die Vorlage; siehe Fund F09)."""
    tab = RN_TABELLEN[spalte]
    return tab[rn] if 0 <= rn < len(tab) else 0.0


def handrechnung_bauteil(nhk: float, mangel: float, rn: int, spalte: str, m2: float, index: float) -> dict[str, float]:
    h14 = rd(nhk - nhk * mangel / 100)
    abzug = altersabzug_tabelle(spalte, rn)
    h15 = rd(h14 - h14 * abzug / 100)
    h16 = rd(h15 * (index / 100))
    return {"nhk": nhk, "nach_maengel": h14, "abzug": abzug, "nach_alter": h15, "je_m2": h16, "betrag": rd(m2 * h16)}


def baue_szenarien(m: Mappe, lauf: Lauf) -> dict[str, Any]:
    """Legt alle Szenarien an und gibt die Handrechnungen (Sollwerte) zurück."""
    index = float(m.roh(BL, "H7") or 0)
    modus = "restwert" if "0.97" in m.formel(BL, "C75") else "linear"
    soll: dict[str, Any] = {"index": index, "modus": modus}

    def formalien(extra: dict) -> dict:
        return {**BASIS_FORMALIEN, **extra}

    # T0: leere Vorlage
    lesen0 = [(DB, f"E{i}") for i in range(21, 29)] + [(DB, "C7"), (DB, "C11"), (BL, "K81"), (BL, "H104"), (BL, "I38"),
                                                       (BE, "K24"), (MG, "F46"), (AA, "J3"), (BL, "C108")]
    lauf.neu("T0", "leere Vorlage ohne Eingaben", {}, lesen0)

    # T1: typischer Fall mit allen Bauteilen, Gewächshaus, Grube und Kanal
    e = dict(BASIS_FORMALIEN)
    lesen1: list[tuple[str, str]] = []
    teile = {}
    for k, (typ, stufe, spalte, rn, mangel, m2) in T1_BAUTEILE.items():
        zt, zrn, zm, za, zb = BAUTEIL_ZELLEN[k]
        e[(BL, zt)] = m.typ_name(typ)
        e[(BL, zrn)], e[(BL, zm)], e[(BL, za)] = rn, mangel, m2
        lesen1 += [(BL, zb)]
        teile[k] = handrechnung_bauteil(NHK[stufe], mangel, rn, spalte, m2, index)
    e.update({(BL, "D41"): 160, (BL, "G40"): 5, (BL, "I40"): 4, (BL, "F44"): 10,
              (BL, "C74"): 10, (BL, "I75"): 6, (BL, "F74"): 470, (BL, "G78"): 10, (BL, "C79"): 2000,
              (BL, "F104"): 0, (BL, "G104"): 0})
    lesen1 += [(BL, z) for z in ("K17", "K26", "K35", "K52", "K44", "K53", "K61", "K71", "K75", "K79", "K81", "F101",
                                 "C33", "E33", "H32", "H33", "C75", "F75", "H79", "E79")] + [(DB, "E21"), (DB, "E28")]
    lauf.neu("T1", "typischer Fall: Laube, Anbau A1, Anbau A2, Schuppen, Gewächshaus, Grube 10 J., Kanal 10 J.", e, lesen1)
    h41 = rd(160 - 160 * 0.1)
    i42 = rd(h41 - h41 * 20 / 100)
    i43 = rd(i42 * (index / 100))
    gwh = rd(10 * i43)
    grube = grube_betrag(10, 6, 470, modus)
    kanal = kanal_betrag(10, 2000, modus)
    summe_bau = sum(t["betrag"] for t in teile.values())
    soll["T1"] = {"teile": teile, "gwh": gwh, "grube": grube, "kanal": kanal,
                  "K53": rd(summe_bau + gwh), "K81": rd(summe_bau + gwh + grube + kanal),
                  "F101": sum(T1_BAUTEILE[k][5] for k in ("L", "A1", "A2", "S"))}

    # T2: alle Zeilen voll (Außenanlagen, Aufwuchs I und II, Beseitigungskosten)
    e = dict(BASIS_FORMALIEN)
    e.update({(BL, "F104"): 0, (BL, "G104"): 0})
    aa = aw1 = aw2 = 0.0
    for r in range(15, 46):
        a, k = 1 + (r % 5) * 0.5, 2.5 + (r % 11) * 1.75 + 0.05 * (r % 3)
        e[(AA, f"A{r}")], e[(AA, f"K{r}")] = a, k
        aa += rd(a * k)
    for r in range(18, 46):
        a, k = 1 + (r % 4) * 0.5, 1.2 + (r % 9) * 2.15
        e[(AW1, f"A{r}")], e[(AW1, f"K{r}")] = a, k
        aw1 += rd(a * k)
    for r in range(10, 46):
        a, k = 2 + (r % 3) * 0.5, 0.8 + (r % 7) * 3.1
        e[(AW2, f"A{r}")], e[(AW2, f"K{r}")] = a, k
        aw2 += rd(a * k)
    be1 = be1_roh = be2 = be2_roh = 0.0
    faktoren = [1, 1.5, 2, 1.25]
    for r in list(range(12, 24)) + list(range(34, 45)):
        f, h, j = 1 + r % 3, round(10.05 + r * 0.7, 2), faktoren[r % 4]
        e[(BE, f"F{r}")], e[(BE, f"H{r}")], e[(BE, f"J{r}")] = f, h, j
        if r < 30:
            be1 += rd(f * h * j)
            be1_roh += f * h * j
        else:
            be2 += rd(f * h * j)
            be2_roh += f * h * j
    lauf.neu("T2", "alle Zeilen voll: Außenanlagen 31, Aufwuchs I 28, Aufwuchs II 36, Beseitigung 12 + 11 Zeilen", e,
             [(AA, "L46"), (AW1, "L46"), (AW2, "L46"), (BE, "K24"), (BE, "K45"), (BE, "K21")] +
             [(DB, f"E{i}") for i in range(21, 29)])
    soll["T2"] = {"AA": rd(aa), "AW1": rd(aw1), "AW2": rd(aw1 + aw2), "BE1": rd(be1), "BE1_roh": be1_roh,
                  "BE2": rd(be2), "BE2_roh": be2_roh, "K21": rd((1 + 21 % 3) * round(10.05 + 21 * 0.7, 2) * faktoren[21 % 4])}
    soll["T2"]["E28"] = rd(soll["T2"]["AA"] + soll["T2"]["AW2"] - soll["T2"]["BE1"] - soll["T2"]["BE2"])

    # T3/T4: überdachter Laubenvorplatz a) integriert, b) angebaut (Laube wie in T1)
    laube = {(BL, "G12"): m.typ_name("Doppelt Holz 10,01 - 14,99 cm"), (BL, "I13"): 20, (BL, "E14"): 10, (BL, "A17"): 20,
             (BL, "F104"): 0, (BL, "G104"): 0}
    lesen3 = [(BL, z) for z in ("G63", "I63", "C64", "E64", "I64", "I67", "G68", "I68", "C70", "I70", "K71", "K81", "E15")]
    lauf.neu("T3", "Laubenvorplatz a) integriert, 275 €/m² bis, 6 m²", formalien({**laube, (BL, "E63"): 275, (BL, "F71"): 6}), lesen3)
    i63 = rd(275 - 275 * 10 / 100)
    abzug_laube = altersabzug_tabelle("II/III", 20)
    i64 = rd(i63 - i63 * abzug_laube / 100)
    i70 = rd(i64 * (index / 100))
    soll["T3"] = {"I63": i63, "I64": i64, "I70": i70, "K71": rd(6 * i70)}
    lauf.neu("T4", "Laubenvorplatz b) angebaut, hochwertig 175 €/m², 3 %/Jahr × 8 Jahre, 6 m²",
             formalien({**laube, (BL, "B67"): "hochwertig", (BL, "E66"): 3, (BL, "H66"): 8, (BL, "G67"): 0.1,
                        (BL, "F71"): 6}), lesen3)
    i67 = rd(175 - 175 * 0.1)
    i68 = rd(i67 - i67 * 24 / 100)
    i70b = rd(i68 * (index / 100))
    soll["T4"] = {"I67": i67, "G68": 24, "I68": i68, "I70": i70b, "K71": rd(6 * i70b)}

    # T5: Randfälle Bau (Auswahl leer, Restnutzungsjahre jenseits der Tabelle, Mängel über 100 %, alte Grube/Kanal)
    e = dict(BASIS_FORMALIEN)
    e.update({(BL, "G12"): None, (BL, "I13"): 20, (BL, "E14"): 10, (BL, "A17"): 20,
              (BL, "G21"): m.typ_name("Stein 10,00 - 12,00 cm"), (BL, "I22"): 61, (BL, "E23"): 0, (BL, "A26"): 3,
              (BL, "G30"): m.typ_name("Doppelwand 4,01 - 8,00 cm"), (BL, "I31"): 41, (BL, "E32"): 0, (BL, "A35"): 1,
              (BL, "G47"): m.typ_name("Blockhaus bis 3,00 cm"), (BL, "I48"): 8, (BL, "E49"): 120, (BL, "A52"): 4,
              (BL, "C74"): 40, (BL, "I75"): 6, (BL, "F74"): 470, (BL, "G78"): 40, (BL, "C79"): 2000,
              (BL, "F104"): 0, (BL, "G104"): 0})
    lauf.neu("T5", "Randfälle: Auswahl leer, Restnutzungsjahre 61 (III) und 41 (IV), Mängel 120 %, Grube/Kanal 40 Jahre", e,
             [(BL, z) for z in ("F12", "C14", "E17", "K53", "K81", "E24", "E26", "E33", "E35", "H49", "E52", "C75", "F75", "K75",
                               "H79", "K79")])

    # T6/T7: Kanal – Alter 0 und Kosten bis 74 €
    lauf.neu("T6", "Kanal erstellt im Jahr der Wertermittlung (Alter 0), Kosten 2.000 €",
             formalien({(BL, "G78"): 0, (BL, "C79"): 2000, (BL, "F104"): 0, (BL, "G104"): 0}),
             [(BL, "E79"), (BL, "H79"), (BL, "K79")])
    lauf.neu("T7", "Kanal Alter 10, Kosten 70 €",
             formalien({(BL, "G78"): 10, (BL, "C79"): 70, (BL, "F104"): 0, (BL, "G104"): 0}),
             [(BL, "E79"), (BL, "H79"), (BL, "K79")])

    # T8: Aufwuchs II mit Abzug (negativer Preis), Beseitigung mit Rundung, kleine Parzelle
    e = formalien({(BL, "F104"): 0, (BL, "G104"): 0, (AW1, "A18"): 10, (AW1, "K18"): 10, (AW2, "A10"): 1, (AW2, "K10"): -20,
                   (BE, "F34"): 3, (BE, "H34"): 10.05, (BE, "J34"): 1.5, (BE, "F35"): 3, (BE, "H35"): 10.05, (BE, "J35"): 1.5,
                   (BE, "H12"): 0, (FO, "E16"): 20})
    lauf.neu("T8", "Aufwuchs I 100 €, Aufwuchs II mit Abzug −20 €; Beseitigung 3 × 10,05 € × 150 %; Parzelle 20 m²", e,
             [(AW1, "L46"), (AW2, "L46"), (DB, "E23"), (BE, "K34"), (BE, "K35"), (BE, "K45"), (DB, "E27"), (AA, "J3")])

    # T9: falsche Eingaben bei den Restnutzungsjahren
    e = formalien({(BL, "G12"): m.typ_name("Doppelt Holz 10,01 - 14,99 cm"), (BL, "I13"): "zwanzig", (BL, "E14"): 10, (BL, "A17"): 20,
                   (BL, "G21"): m.typ_name("Stein 10,00 - 12,00 cm"), (BL, "I22"): -3, (BL, "E23"): 0, (BL, "A26"): 3,
                   (BL, "G30"): m.typ_name("Doppelwand 4,01 - 8,00 cm"), (BL, "I31"): 7.5, (BL, "E32"): 0, (BL, "A35"): 1,
                   (BL, "F104"): 0, (BL, "G104"): 0})
    lauf.neu("T9", "Restnutzungsjahre als Text, negativ und mit Komma (7,5)", e,
             [(BL, "E15"), (BL, "E17"), (BL, "E24"), (BL, "E26"), (BL, "E33"), (BL, "E35"), (BL, "K53"), (DB, "E21")])
    soll["T9"] = {"E33_7_5": altersabzug_tabelle("IV", 7)}
    return soll


# ---------------------------------------------------------------------------------------------------
# 4. Katalog der Funde (Text für den Bericht) und Prüfungen (automatisch, je Fund)
# ---------------------------------------------------------------------------------------------------

@dataclass
class Fund:
    nr: str
    blatt: str
    zellen: str
    titel: str
    ist: str              # was steht da
    problem: str          # was ist falsch
    auswirkung: str       # eine von AUSWIRKUNGEN
    gewissheit: str       # eine von GEWISSHEITEN
    beleg: str            # Richtlinie / Testfall (fest); die Zahlen aus dem Lauf hängt der Bericht an
    vorschlag: str        # Korrekturvorschlag (konkrete Formel / Text)
    top: int = 0          # >0: gehört zu den wichtigsten Funden (Reihenfolge)


@dataclass
class Befund:
    vorhanden: bool | None   # True = Fehler in der Datei vorhanden, False = nicht (mehr), None = nicht automatisch prüfbar
    beleg: str = ""


def eur(x: Any) -> str:
    if isinstance(x, (int, float)):
        return f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".").replace("-", "−")
    return str(x)


class Ctx:
    def __init__(self, m: Mappe, lauf: Lauf | None, soll: dict[str, Any]):
        self.m, self.lauf, self.soll = m, lauf, soll

    def r(self, szenario: str, blatt: str, zelle: str) -> Any:
        if self.lauf is None:
            raise LookupError("keine Rechentests")
        return self.lauf[szenario][(blatt, zelle)]


PRUEFUNGEN: dict[str, Callable[[Ctx], Befund]] = {}


def pruefung(nr: str):
    def deko(fn):
        def sicher(ctx: Ctx) -> Befund:
            try:
                return fn(ctx)
            except KeyError as ex:
                return Befund(None, f"nicht prüfbar: {ex}")
            except LookupError:
                return Befund(None, "nicht prüfbar (ohne Rechentests)")
        PRUEFUNGEN[nr] = sicher
        return fn
    return deko


def ist_fehlerwert(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("#")


@pruefung("F01")
def _f01(c: Ctx) -> Befund:
    fehler = c.m.formel(BL, "C33") != "H32"
    b = f"C33 enthält {fx(c.m.roh(BL, 'C33'))}, soll H32."
    if c.lauf:
        ist, soll = c.r("T1", BL, "E35"), c.soll["T1"]["teile"]["A2"]["betrag"]
        b += f" Testfall T1: Anbau A2 (E35) {eur(ist)}, Handrechnung {eur(soll)}."
    return Befund(fehler, b)


@pruefung("F02")
def _f02(c: Ctx) -> Befund:
    terme = norm(c.m.roh(BL, "K81")).split("+")
    fehler = "K52" in terme
    b = f"K81 enthält {fx(c.m.roh(BL, 'K81'))}; K53 (= K61) enthält K52 bereits."
    if c.lauf:
        s = c.soll["T1"]
        b += (f" Testfall T1: Baulichkeiten insgesamt (K81) {eur(c.r('T1', BL, 'K81'))}, "
              f"Handrechnung {eur(s['K81'])}; Zwischensumme K53 {eur(c.r('T1', BL, 'K53'))}, Handrechnung {eur(s['K53'])}.")
    return Befund(fehler, b)


@pruefung("F03")
def _f03(c: Ctx) -> Befund:
    fehler = c.m.roh(BL, "C64") in (None, "")
    b = "C64 ist leer, I64 und E64 rechnen damit."
    if c.lauf:
        b += (f" Testfall T3 (275 €/m², 6 m²): Wert des Vorplatzes K71 {eur(c.r('T3', BL, 'K71'))}, "
              f"Handrechnung {eur(c.soll['T3']['K71'])}.")
    return Befund(fehler, b)


@pruefung("F04")
def _f04(c: Ctx) -> Befund:
    ohne = [f"K{r}" for r in list(range(12, 24)) + list(range(34, 45)) if not ist_formel(c.m.roh(BE, f"K{r}"))]
    b = f"Zellen ohne Formel in K12:K23 / K34:K44: {', '.join(ohne) or 'keine'}."
    if c.lauf:
        s = c.soll["T2"]
        b += (f" Testfall T2: Summe Anlage 4 Abschnitt 1 (K24) {eur(c.r('T2', BE, 'K24'))}, "
              f"Handrechnung {eur(s['BE1'])}.")
    return Befund(bool(ohne), b)


@pruefung("F05")
def _f05(c: Ctx) -> Befund:
    fehler = not ist_formel(c.m.roh(BL, "G68"))
    b = f"G68 enthält {fx(c.m.roh(BL, 'G68'))}; E66 und H66 werden von keiner Formel gelesen."
    if c.lauf:
        s = c.soll["T4"]
        b += (f" Testfall T4 (hochwertig 175 €/m², 3 % × 8 Jahre, 6 m²): K71 {eur(c.r('T4', BL, 'K71'))}, "
              f"Handrechnung {eur(s['K71'])}.")
    return Befund(fehler, b)


@pruefung("F06")
def _f06(c: Ctx) -> Befund:
    f = c.m.formel(BL, "C75")
    fehler = "MIN(" not in f and "0.97" not in f
    b = f"C75 enthält {fx(c.m.roh(BL, 'C75'))}."
    if c.lauf:
        b += (f" Testfall T5 (Alter 40 J., 6 m³): Altersabzug {c.r('T5', BL, 'C75')} %, Wert K75 {eur(c.r('T5', BL, 'K75'))}.")
        fehler = fehler and isinstance(c.r("T5", BL, "K75"), (int, float)) and c.r("T5", BL, "K75") < 0
    return Befund(fehler, b)


@pruefung("F07")
def _f07(c: Ctx) -> Befund:
    f = c.m.formel(BL, "H79")
    fehler = "MIN(" not in f and "0.97" not in f
    b = f"H79 enthält {fx(c.m.roh(BL, 'H79'))}."
    if c.lauf:
        b += f" Testfall T5 (Alter 40 J., 2.000 €): Abzug {c.r('T5', BL, 'H79')}, Wert K79 {eur(c.r('T5', BL, 'K79'))}."
        fehler = fehler and isinstance(c.r("T5", BL, "K79"), (int, float)) and c.r("T5", BL, "K79") < 0
    return Befund(fehler, b)


@pruefung("F08")
def _f08(c: Ctx) -> Befund:
    b = f"E79 enthält {fx(c.m.roh(BL, 'E79'))}; K79 enthält {fx(c.m.roh(BL, 'K79'))}."
    fehler = "C79>74" in c.m.formel(BL, "K79") or "G78>0" in c.m.formel(BL, "E79")
    if c.lauf:
        b += (f" Testfall T6 (Alter 0, 2.000 €): K79 {eur(c.r('T6', BL, 'K79'))} (Handrechnung {eur(kanal_betrag(0, 2000, c.soll['modus']))}); "
              f"T7 (Alter 10, 70 €): K79 {eur(c.r('T7', BL, 'K79'))} (Handrechnung {eur(kanal_betrag(10, 70, c.soll['modus']))}, "
              f"{'linear' if c.soll['modus'] == 'linear' else 'vom Restwert'} gerechnet).")
        fehler = (c.r("T6", BL, "K79") == 0) or (c.r("T7", BL, "K79") == 0)
    return Befund(fehler, b)


@pruefung("F09")
def _f09(c: Ctx) -> Befund:
    ws = c.m.blatt(T1)
    leer = [spalte for spalte, col, ende in (("IV", "E", 43), ("II/III", "D", 63)) if ws[f"{col}{ende + 1}"].value is None]
    b = f"Tabelle1: Spalte ohne Werte nach dem Tabellenende (Bauklasse IV ab 41 J., II/III ab 61 J.): {', '.join(leer) or 'keine'}."
    fehler = bool(leer)
    if c.lauf:
        b += (f" Testfall T5: Restnutzungsjahre 61 (Bauklasse II) → Altersabzug E24 = {c.r('T5', BL, 'E24')} %, "
              f"41 (Bauklasse IV) → E33 = {c.r('T5', BL, 'E33')} % (Richtlinie: Tabelle endet bei 10 %).")
        fehler = c.r("T5", BL, "E24") == 0 or c.r("T5", BL, "E33") == 0
    return Befund(fehler, b)


@pruefung("F10")
def _f10(c: Ctx) -> Befund:
    fehler = "MAX(" in c.m.formel(DB, "E23")
    b = f"E23 enthält {fx(c.m.roh(DB, 'E23'))}."
    if c.lauf:
        b += (f" Testfall T8 (Aufwuchs I 100 €, Aufwuchs II mit Abzug: 80 €): E23 {eur(c.r('T8', DB, 'E23'))}, "
              f"Handrechnung {eur(c.r('T8', AW2, 'L46'))}.")
        fehler = fehler and not fast_gleich(c.r("T8", DB, "E23"), c.r("T8", AW2, "L46"))
    return Befund(fehler, b)


@pruefung("F11")
def _f11(c: Ctx) -> Befund:
    ohne = [f"K{r}" for r in list(range(12, 24)) + list(range(34, 45))
            if ist_formel(c.m.roh(BE, f"K{r}")) and "ROUND(" not in c.m.formel(BE, f"K{r}")]
    b = f"Formeln ohne RUNDEN: {len(ohne)} von 23 Zeilen."
    fehler = bool(ohne)
    if c.lauf:
        b += (f" Testfall T8 (3 × 10,05 € × 150 %): K34 = {_bild(c.r('T8', BE, 'K34'))} (Handrechnung 45,23 €), "
              f"Summe K45 = {_bild(c.r('T8', BE, 'K45'))}.")
    return Befund(fehler, b)


@pruefung("F12")
def _f12(c: Ctx) -> Befund:
    gefunden = []
    for blatt, zelle in ((BL, "F104"), (BL, "G104"), (BL, "H105"), (BL, "I19"), (BL, "I28"), (BL, "I37"), (BE, "H12"),
                         (MG, "F44"), (BE, "A26")):
        v = c.m.roh(blatt, zelle)
        if v not in (None, "", 0) and not ist_formel(v):
            gefunden.append(f"{blatt}!{zelle} = {str(v)[:30]!r}")
    b = f"Vorbelegt: {'; '.join(gefunden) or 'nichts'}."
    fehler = bool(gefunden)
    if c.lauf:
        b += f" Testfall T0 (leere Vorlage): Gesamtbetrag mit Abzügen (Deckblatt E28) = {eur(c.r('T0', DB, 'E28'))}."
        fehler = bool(gefunden) or c.r("T0", DB, "E28") != 0
    return Befund(fehler, b)


def _dv(m: Mappe, blatt: str, zelle: str):
    ws = m.blatt(blatt)
    for dv in ws.data_validations.dataValidation:
        if zelle in str(dv.sqref).split():
            return dv
    return None


@pruefung("F13")
def _f13(c: Ctx) -> Befund:
    dv = _dv(c.m, BL, "K77")
    return Befund(dv is not None and "-100" in (dv.formula1 or ""), f"Auswahlliste K77: {dv.formula1 if dv else 'keine'}.")


@pruefung("F14")
def _f14(c: Ctx) -> Befund:
    dv = _dv(c.m, BL, "E63")
    return Befund(dv is not None and "225" in (dv.formula1 or ""), f"Auswahlliste E63: {dv.formula1 if dv else 'keine'}.")


@pruefung("F15")
def _f15(c: Ctx) -> Befund:
    ohne_pruefung = [z for z in ("E14", "E23", "E32", "E49") if _dv(c.m, BL, z) is None]
    b = f"Eingabezellen ohne Grenzwert-Prüfung: {', '.join(ohne_pruefung) or 'keine'}; G67 erwartet einen Bruchteil, E14 Prozent."
    fehler = bool(ohne_pruefung)
    if c.lauf:
        b += (f" Testfall T5 (Mängel 120 %): NHK nach Mängelabzug H49 = {c.r('T5', BL, 'H49')}, Betrag E52 = {eur(c.r('T5', BL, 'E52'))}.")
    return Befund(fehler, b)


@pruefung("F16")
def _f16(c: Ctx) -> Befund:
    f = c.m.formel(AA, "J3")
    fehler = "MAX(" not in f or "ROUND(" not in f
    b = f"J3 enthält {fx(c.m.roh(AA, 'J3'))}."
    if c.lauf:
        b += f" Testfall T8 (Parzelle 20 m²): J3 = {c.r('T8', AA, 'J3')} m²; T0 (leer): {c.r('T0', AA, 'J3')} m²."
    return Befund(fehler, b)


@pruefung("F17")
def _f17(c: Ctx) -> Befund:
    t = str(c.m.roh(BE, "M29"))
    return Befund("1,80" in t, f"M29: „{t}“; A51 und M32 nennen 1,25 m.")


def _norm_text(s: str) -> str:
    s = s.lower().replace("–", "-").replace("ß", "ss")
    return re.sub(r"[^a-z0-9äöü]+", "", s)


def _aehnlich(a: str, b: str) -> float:
    import difflib
    return difflib.SequenceMatcher(None, _norm_text(a), _norm_text(b)).ratio()


def maengel_vergleich(m: Mappe) -> tuple[list[str], list[str], list[str]]:
    """(fehlende Richtlinien-Positionen, überzählige Positionen der Vorlage, doppelte Positionen)."""
    ws = m.blatt(MG)
    fehlt, extra, doppelt = [], [], []
    for bauteil, (_, liste) in MAENGEL_RICHTLINIE.items():
        a, b = MAENGEL_ZEILEN[bauteil]
        vorlage = [str(ws[f"{col}{r}"].value) for r in range(a, b + 1) for col in "CD"
                   if ws[f"{col}{r}"].value not in (None, "")]
        gesehen: set[str] = set()
        for v in vorlage:
            k = _norm_text(v)
            if k in gesehen:
                doppelt.append(f"{bauteil}: {v.strip()}")
            gesehen.add(k)
        for rl in liste:
            if not any(_aehnlich(rl, v) >= 0.8 for v in vorlage):
                fehlt.append(f"{bauteil}: {rl}")
        for v in vorlage:
            if not any(_aehnlich(rl, v) >= 0.8 for rl in liste):
                extra.append(f"{bauteil}: {v.strip()}")
    return fehlt, extra, doppelt


@pruefung("F18")
def _f18(c: Ctx) -> Befund:
    fehlt, extra, doppelt = maengel_vergleich(c.m)
    abw = []
    for bauteil, zelle in MAENGEL_MAX_ZELLEN.items():
        if not fast_gleich(c.m.roh(MG, zelle), MAENGEL_RICHTLINIE[bauteil][0] / 100, 0.0001):
            abw.append(f"{bauteil} {fx(c.m.roh(MG, zelle))}")
    b = (f"Fehlende Positionen ({len(fehlt)}): {'; '.join(fehlt) or '–'}. Nicht in der Richtlinie ({len(extra)}): "
         f"{'; '.join(extra) or '–'}. Doppelt ({len(doppelt)}): {'; '.join(doppelt) or '–'}. "
         f"Abzug-maximal-Werte weichen ab: {', '.join(abw) or 'keine'}.")
    return Befund(bool(fehlt or extra or doppelt or abw), b)


def _verweist_auf(m: Mappe, blatt: str, muster: str) -> bool:
    ws = m.blatt(blatt)
    return any(ist_formel(c.value) and muster in c.value for z in ws.iter_rows() for c in z)


@pruefung("F19")
def _f19(c: Ctx) -> Befund:
    verkn = _verweist_auf(c.m, BL, "Mängelliste") or _verweist_auf(c.m, BL, "Maengel")
    fmt = c.m.blatt(MG)["F46"].number_format
    return Befund(not verkn, f"Kein Verweis von Baulichkeiten auf die Mängelliste; Zahlenformat der Summe F46: „{fmt}“.")


def _text_enthaelt(m: Mappe, blatt: str, *woerter: str) -> bool:
    return any(w.lower() in t.lower() for b, _, t in m.texte() if b == blatt for w in woerter)


@pruefung("F20")
def _f20(c: Ctx) -> Befund:
    ok = _text_enthaelt(c.m, BL, "begründ", "nachgewiesen", "geschätzt")
    return Befund(not ok, "Im Blatt Baulichkeiten kein Feld/Hinweis zu Begründung der Restnutzungsdauer oder Nachweis des Baujahrs.")


@pruefung("F21")
def _f21(c: Ctx) -> Befund:
    return Befund(not _text_enthaelt(c.m, BL, "Keller"), "Im Blatt Baulichkeiten kommt „Keller“ nicht vor.")


@pruefung("F22")
def _f22(c: Ctx) -> Befund:
    ok = _text_enthaelt(c.m, BL, "Trockentoilette", "Elektr", "WC")
    return Befund(not ok, "In der Baubeschreibung (A82:K100) kein Raster für Wasser/WC/Trockentoilette/Elektro.")


@pruefung("F23")
def _f23(c: Ctx) -> Befund:
    fehlt = []
    if _dv(c.m, BL, "F74") is None:
        fehlt.append("Auswahl Grundwert 470/370/180")
    if not _text_enthaelt(c.m, BL, "Sanierung"):
        fehlt.append("Sanierungsdatum")
    if "MIN(" not in c.m.formel(BL, "C75") and "16" not in c.m.formel(BL, "C75"):
        fehlt.append("16-Jahre-Regel für Tanks/monolithische Gruben")
    return Befund(bool(fehlt), f"Fehlt: {', '.join(fehlt) or 'nichts'}.")


@pruefung("F24")
def _f24(c: Ctx) -> Befund:
    f = c.m.formel(BL, "F101")
    return Befund("F71" not in f, f"F101 enthält {fx(c.m.roh(BL, 'F101'))}; I38 und F101 werden nirgends weiterverwendet, keine Prüfung gegen 24 m².")


@pruefung("F25")
def _f25(c: Ctx) -> Befund:
    d41 = c.m.roh(BL, "D41")
    fehler = d41 != 160 or not _text_enthaelt(c.m, BL, "12 m", "Firsthöhe")
    return Befund(fehler, f"D41 = {d41!r} (Richtlinie 160 €/m²); Grenzen 12 m² / 2,20 m Firsthöhe stehen nicht im Blatt; F41 (Mängel) fest {fx(c.m.roh(BL, 'F41'))}.")


@pruefung("F26")
def _f26(c: Ctx) -> Befund:
    return Befund(None, "Fachfrage an den BV (Richtlinie S. 11: „bei allen Anlagen mindestens 3 % jährlich“); nicht automatisch prüfbar.")


@pruefung("F27")
def _f27(c: Ctx) -> Befund:
    kopf = " ".join(str(c.m.roh(AA, z) or "") for z in ("A13", "B13", "C13", "K13", "K14", "L14"))
    return Befund("alter" not in kopf.lower(), f"Spaltenüberschriften Außenanlagen: {kopf.strip()!r}; Aufwuchs analog.")


@pruefung("F28")
def _f28(c: Ctx) -> Befund:
    text = str(c.m.roh(BL, "I7"))
    jahr = re.search(r"(20\d\d)", text)
    heute = dt.date.today().year
    fehler = bool(jahr) and int(jahr.group(1)) < heute
    return Befund(fehler, f"H7 enthält {fx(c.m.roh(BL, 'H7'))}, Beschriftung I7: „{text}“; heute ist {heute}.")


@pruefung("F29")
def _f29(c: Ctx) -> Befund:
    f1, f2 = c.m.formel(BL, "C75"), c.m.formel(BL, "H79")
    linear = f1 == "C74*3" or f2 == "G78/100*3"
    return Befund(linear, f"C75 enthält {fx(c.m.roh(BL, 'C75'))}, H79 enthält {fx(c.m.roh(BL, 'H79'))} (linear: Alter × 3 %); die App rechnet (1 − 0,97^Alter), PR #95.")


@pruefung("F30")
def _f30(c: Ctx) -> Befund:
    ohne_bezug = not _verweist_auf(c.m, DB, "H105") and not _verweist_auf(c.m, BL, "H105")
    zeile12 = c.m.roh(BE, "H12")
    return Befund(ohne_bezug or zeile12 not in (None, ""),
                  f"H105 enthält {c.m.roh(BL, 'H105')!r} wird nirgends verrechnet; K77 bietet nur −100 € an (Auswahlliste); Beseitigungskosten H12 = {zeile12!r}.")


@pruefung("F31")
def _f31(c: Ctx) -> Befund:
    t = str(c.m.roh(BE, "M30"))
    return Befund("10%" in t.replace(" ", ""), f"M30: „{t}“ – in der Richtlinie nicht zu finden.")


@pruefung("F32")
def _f32(c: Ctx) -> Befund:
    schutz = c.m.blattschutz()
    haupt = (FO, DB, BL, MG, AA, AW1, AW2, BE)
    aus = [b for b in haupt if schutz.get(b) != "aktiv"]
    hinweis = "geschützt" in str(c.m.roh(FO, "H23") or "")
    return Befund(bool(aus) and hinweis, f"Blattschutz nicht aktiv bei: {', '.join(aus) or 'keinem'} (Formalien H23 verspricht ihn).")


def gesperrte_eingabefelder(m: Mappe) -> list[str]:
    """Zellen, die erkennbar zur Eingabe gedacht sind (Auswahlliste, Eingabe-Format, Vorbelegung), aber gesperrt sind."""
    kandidaten = {
        FO: ["E17"], AA: ["B15", "C15", "C20", "A8"], BE: ["B12", "F12", "H12", "J12", "F34", "H34", "J34", "B34"],
        BL: ["I19", "I28", "I37"],
    }
    res = []
    for blatt, zellen in kandidaten.items():
        ws = m.blatt(blatt)
        for z in zellen:
            if ws[z].protection.locked:
                res.append(f"{blatt}!{z}")
    return res


@pruefung("F33")
def _f33(c: Ctx) -> Befund:
    g = gesperrte_eingabefelder(c.m)
    return Befund(bool(g), f"Gesperrt (wären bei aktivem Blattschutz nicht bedienbar): {', '.join(g) or 'keine'}.")


@pruefung("F34")
def _f34(c: Ctx) -> Befund:
    offen = []
    for blatt in (FO, DB, BL, MG, AA, AW1, AW2, BE):
        ws = c.m.blatt(blatt)
        for z in ws.iter_rows():
            for cell in z:
                if ist_formel(cell.value) and cell.protection.locked is False:
                    offen.append(f"{blatt}!{cell.coordinate}")
    return Befund(bool(offen), f"Formelzellen, die nicht gesperrt sind: {', '.join(offen) or 'keine'}.")


@pruefung("F35")
def _f35(c: Ctx) -> Befund:
    b = "Auswahlfelder G12/G21/G30/G47 und Restnutzungsjahre I13/I22/I31/I48 ohne Eingabeprüfung für leer/Text/negativ."
    fehler = True
    if c.lauf:
        ohne_pruefung = [z for z in ("I13", "I22", "I31", "I48") if _dv(c.m, BL, z) is None]
        fehler = ist_fehlerwert(c.r("T5", BL, "K53")) or (ist_fehlerwert(c.r("T9", BL, "E15")) and bool(ohne_pruefung))
        b += (f" Testfall T5 (Typ leer): K53 = {c.r('T5', BL, 'K53')}, K81 = {c.r('T5', BL, 'K81')}; "
              f"T9 (Restnutzungsjahre „zwanzig“): E15 = {c.r('T9', BL, 'E15')}, Deckblatt E21 = {c.r('T9', DB, 'E21')}.")
    return Befund(fehler, b)


@pruefung("F36")
def _f36(c: Ctx) -> Befund:
    ohne_wenn = [z for z in ("C7", "C8") if "IF(" not in c.m.formel(DB, z)]
    b = f"Deckblatt C7 enthält {fx(c.m.roh(DB, 'C7'))} (ohne WENN-Abfrage)."
    if c.lauf:
        b += f" Testfall T0 (leer): C7 zeigt {c.r('T0', DB, 'C7')!r} statt leer; das Datum C11 ist der Wert 0 (= 00.01.1900)."
    return Befund(bool(ohne_wenn), b)


@pruefung("F37")
def _f37(c: Ctx) -> Befund:
    ws = c.m.blatt(MG)
    formate = {ws[z].number_format for z in ("F5", "F13", "F44", "F46")}
    fehler = formate <= {"0", "General"}
    return Befund(fehler, f"Zahlenformate der Spalte „Abzug konkret“: {sorted(formate)}; Eingaben wie 0,05 erscheinen als „0“.")


def druckbild_pruefung(m: Mappe) -> list[str]:
    probleme = []
    ws = m.blatt(BL)
    erste_seite = sum((ws.row_dimensions[r].height or (ws.sheet_format.defaultRowHeight or 15)) for r in range(1, 55))
    skala = (ws.page_setup.scale or 100) / 100
    nutzbar = 841.9 - (ws.page_margins.top + ws.page_margins.bottom) * 72
    if erste_seite * skala > nutzbar:
        probleme.append(f"Baulichkeiten Seite 1 (Zeilen 1–54): {erste_seite * skala:.0f} pt Inhalt, {nutzbar:.0f} pt nutzbar → "
                        "Zeile 54 („Seite 2“) rutscht auf eine leere Seite")
    for blatt in (FO, AW2):
        if not m.blatt(blatt).print_area:
            probleme.append(f"{blatt}: kein Druckbereich")
    return probleme


@pruefung("F38")
def _f38(c: Ctx) -> Befund:
    p = druckbild_pruefung(c.m)
    return Befund(bool(p), "; ".join(p) or "keine Auffälligkeiten")


def fehlerzellen_versteckt(m: Mappe) -> dict[str, int]:
    res = {}
    for ws in m.wb:
        n = sum(1 for z in ws.iter_rows() for cell in z if ist_formel(cell.value) and "#REF!" in cell.value)
        if n:
            res[f"{ws.title}{' (versteckt)' if ws.sheet_state != 'visible' else ''}"] = n
    return res


@pruefung("F39")
def _f39(c: Ctx) -> Befund:
    fz = fehlerzellen_versteckt(c.m)
    versteckt = [w.title for w in c.m.wb if w.sheet_state != "visible"]
    f_aufl = norm(c.m.roh("Beseitigungsauflagen", "H22")) if "Beseitigungsauflagen" in c.m.wb.sheetnames else ""
    f_aufl_de = de("=" + f_aufl) if f_aufl else ""
    b = (f"{len(versteckt)} versteckte Blätter; Zellen mit #BEZUG! je Blatt: "
         f"{', '.join(f'{k} {v}' for k, v in fz.items()) or 'keine'}; Beseitigungsauflagen H22 {f_aufl_de or '–'}.")
    return Befund(bool(fz) or "G33" in f_aufl, b)


@pruefung("F40")
def _f40(c: Ctx) -> Befund:
    teile = []
    for z, soll in (("B19", "=B12"), ("B28", "=B21"), ("B37", "=B30")):
        if not ist_formel(c.m.roh(BL, z)):
            teile.append(f"{z} = {c.m.roh(BL, z)!r} statt {soll}")
    if c.m.formel(BL, "J52") != "B47":
        teile.append(f"J52 enthält {fx(c.m.roh(BL, 'J52'))} (zeigt beim Schuppen das Kürzel von A1; soll =$B$47 sein)")
    if c.m.formel(BL, "E34") != "H7/100":
        teile.append(f"E34 enthält {fx(c.m.roh(BL, 'E34'))} (Muster der anderen Bauteile: =H7/100)")
    if any(ist_formel(c.m.roh(BL, z)) and "$" not in str(c.m.roh(BL, z)) for z in ("E16", "E25", "E43", "E51")):
        teile.append("H7 in E16/E25/E43/E51 ohne $ (beim Kopieren verschiebt sich der Bezug)")
    if c.m.roh(BL, "H42") not in (None, ""):
        teile.append(f"H42 enthält {c.m.roh(BL, 'H42')!r} ohne Funktion")
    if c.m.roh(BL, "N79"):
        teile.append(f"N79/N80: Rest-Summe „{fx(c.m.roh(BL, 'N79'))}“ (Formel {fx(c.m.roh(BL, 'N80'))})")
    return Befund(bool(teile), "; ".join(teile) or "nichts auffällig")


@pruefung("F41")
def _f41(c: Ctx) -> Befund:
    d103 = str(c.m.roh(T1, "D103") or "")
    probleme = []
    if "J19:J109" in d103:
        probleme.append("D103: verschachtelte WENN-Formel endet mit kaputtem Bezug (J19:J109)")
    if ist_formel(c.m.roh(T1, "F5")):
        probleme.append(f"F5 enthält {fx(c.m.roh(T1, 'F5'))}")
    ungenutzt = [n for n in ("BKL.Beschreibung", "Jürgen", "rose", "ungeprüft", "uwe.2") if n in c.m.wb.defined_names]
    if ungenutzt:
        probleme.append("ungenutzte Namen: " + ", ".join(ungenutzt))
    return Befund(bool(probleme), "; ".join(probleme) or "keine")


@pruefung("F42")
def _f42(c: Ctx) -> Befund:
    treffer = []
    for blatt, zelle, text in c.m.texte():
        for falsch in TIPPFEHLER:
            if falsch.lower() in text.lower():
                treffer.append(f"{blatt}!{zelle} „{falsch}“")
    wir = [z for z in ("H6", "H7", "H8", "H10", "H15", "H16", "H17", "H19") if str(c.m.roh(FO, z) or "").strip(". ") == "wir"]
    if wir:
        treffer.append("Formalien " + ", ".join(wir) + " „wir.“")
    return Befund(bool(treffer), f"{len(treffer)} Stellen: {'; '.join(treffer[:40])}")


@pruefung("F43")
def _f43(c: Ctx) -> Befund:
    n = [str(c.m.roh(BE, f"A{r}")) for r in range(34, 45)]
    doppelt = len(n) != len(set(n))
    return Befund(doppelt, f"Nummern in A34:A44: {', '.join(n)}; A24 trägt „{fx(c.m.roh(BE, 'A24'))}“ in der Summenzeile.")


@pruefung("F44")
def _f44(c: Ctx) -> Befund:
    t = str(c.m.roh(SK2, "A1"))
    return Befund("Baulichkeiten" in t, f"Skizze Außenanlage-Aufwuchs A1: „{t}“.")


@pruefung("F45")
def _f45(c: Ctx) -> Befund:
    t = str(c.m.roh(BL, "L62"))
    return Befund("E58" in t or "B62" in t, "L62 nennt „E58 bzw. B62“; die Felder sind E63 und B67.")


@pruefung("F46")
def _f46(c: Ctx) -> Befund:
    a34, a38 = str(c.m.roh(DB, "A34")), str(c.m.roh(DB, "A38"))
    widerspruch = "direkt an die Wertermittler" in a34 and "Bezirksverband" in a38
    return Befund(widerspruch, "Deckblatt A34: Widerspruch „direkt an die Wertermittler“; A38: „an den Bezirksverband“.")


# Katalog (Reihenfolge = Nummer). Die Texte sind die Grundlage des Fehlerberichts.
FUNDE: list[Fund] = [
    # ---------------------------------------------------------------------------- falscher Betrag
    Fund("F01", "Baulichkeiten", "C33", "Anbau A2 rechnet mit dem Ergebnis von Anbau A1",
         "=$H$23 (Wert von A1 nach Mängelabzug)",
         "In der Rechenkette von Anbau A2 zeigt C33 auf die Zeile von A1 (H23) statt auf die eigene (H32). A2 wird deshalb mit den NHK und dem Mängelabzug von A1 weitergerechnet; nur der Altersabzug ist der eigene.",
         "falscher Betrag", "sicher", "Vergleich der vier Bauteil-Blöcke (Laube/A1/A2/Schuppen sind gleich aufgebaut); Testfall T1.",
         "=$H$32", top=5),
    Fund("F02", "Baulichkeiten", "K81", "Schuppen wird in „Baulichkeiten insgesamt“ doppelt gezählt",
         "=K52+K61+K71+K75+K77+K79",
         "K61 ist der Übertrag der Zwischensumme K53 =SUMME(K14:K52); die enthält den Schuppen (K52) schon. K81 addiert K52 noch einmal. Bei jedem Protokoll mit Schuppen ist die Summe um den Schuppenwert zu hoch (und damit der Gesamtbetrag auf dem Deckblatt).",
         "falscher Betrag", "sicher", "Testfall T1.", "=K61+K71+K75+K77+K79", top=1),
    Fund("F03", "Baulichkeiten", "C64", "Überdachter Laubenvorplatz a) „integriert“: Wert geht verloren",
         "C64 ist leer (I64 und E64 rechnen damit)",
         "Zwischen I63 (NHK nach Mängelabzug) und I64 (Altersabzug) fehlt die Verbindung. C64 bleibt leer, I64 ist 0, der integrierte Laubenvorplatz wird mit 0 € bewertet, wenn b) „angebaut“ nicht ausgefüllt ist.",
         "falscher Betrag", "sicher", "Testfall T3.", "=I63", top=3),
    Fund("F04", "Beseitigungskosten", "K21", "Zeile 9 (Abschnitt 1) fehlt in der Summe",
         "K21 hat keine Formel (K12:K20 und K22:K23 haben sie)",
         "Menge × Einzelpreis × Faktor der 9. Zeile werden nicht berechnet und fehlen in K24, damit auch in den Beseitigungskosten Bau auf dem Deckblatt (E26): Abzug zu niedrig.",
         "falscher Betrag", "sicher", "Testfall T2.", "=RUNDEN(F21*H21*J21;2)  (wie alle anderen Zeilen)"),
    Fund("F05", "Baulichkeiten", "G68 (E66, H66)", "Laubenvorplatz b) „angebaut“: Altersabzug fest 100 %",
         "G68 = 100 (Zahl, gesperrt); die Felder „Altersminderung pro Jahr“ (E66) und „Alter“ (H66) werden von keiner Formel gelesen",
         "Der Altersabzug des angebauten Vorplatzes ist fest 100 %. I68 ist deshalb immer 0, C70 fällt auf den Wert von a) zurück: ein angebauter Vorplatz wird nie bewertet, was auch immer in Güte, Mängel, Alter eingetragen wird. Beim Gewächshaus ist der entsprechende Abzug richtig als Formel gelöst (E42 = I40*G40).",
         "falscher Betrag", "wahrscheinlich", "Testfall T4; Vergleich mit Gewächshaus (E42); Richtlinie S. 11 (Alterswertminderung mindestens 3 % jährlich).",
         "=MIN(100;E66*H66)   (Abzug in %, höchstens 100; wie beim Gewächshaus E42)", top=2),
    Fund("F06", "Baulichkeiten", "C75 (F75, K75)", "Abwassergrube: Altersabzug kann über 100 % steigen (negativer Wert)",
         "=C74*3",
         "Der Abzug wächst linear ohne Obergrenze: ab 34 Jahren ist er über 100 %, der Wert der Grube wird negativ und mindert die Baulichkeiten. Die Regel „Tanks und monolithische Betongruben nach 16 Jahren nicht weiter abschreiben“ (Hinweis M104) ist in keiner Formel umgesetzt.",
         "falscher Betrag", "sicher", "Richtlinie S. 13–14; Testfall T5.", "=MIN(100;C74*3)  (linear) – oder nach BV-Entscheidung F29: =(1-0,97^C74)*100"),
    Fund("F07", "Baulichkeiten", "H79 (K79)", "Kanalanschluss: Altersabzug kann über 100 % steigen (negativer Wert)",
         "=G78/100*3",
         "Wie F06: ab 34 Jahren ist H79 über 100 %, der Kanalanschluss wird negativ bewertet.",
         "falscher Betrag", "sicher", "Richtlinie S. 14; Testfall T5.", "=MIN(1;G78*3%)  – oder nach BV-Entscheidung F29: =1-0,97^G78"),
    Fund("F08", "Baulichkeiten", "E79, K79", "Kanalanschluss: unklare Bedingungen „Alter > 0“ und „Kosten > 74“",
         "E79 =WENN(G78>0;MIN(2200;C79);0); K79 =WENN(C79>74;RUNDEN(E79-E79*H79;2);0)",
         "Ein im Jahr der Wertermittlung gebauter Anschluss (Alter 0) wird mit 0 € bewertet, obwohl er nach der Richtlinie den vollen Erstellungspreis (bis 2.200 €) bringt. Kosten bis 74 € ergeben immer 0 €; woher die 74 stammen, ist nicht erkennbar.",
         "falscher Betrag", "bitte BV fragen", "Richtlinie S. 14 (Erstellungspreis bis 2.200 €, 3 % jährlich); Testfälle T6, T7.",
         "E79 =WENN(G78=\"\";0;MIN(2200;C79)); K79 =RUNDEN(E79-E79*H79;2)"),
    Fund("F09", "Tabelle1 → Baulichkeiten", "Tabelle1 D64:D83, E44:E83; Baulichkeiten E15, E24, E33, E50",
         "Restnutzungsjahre jenseits der Tabelle ergeben 0 % Altersabzug",
         "Die Restnutzungstabelle endet bei Bauklasse IV nach 40 Jahren, bei II/III nach 60 Jahren; darüber sind die Zellen leer",
         "Wer mehr Restnutzungsjahre einträgt (z. B. 45 bei Bauklasse IV), bekommt statt des Abzugs von 10 % (letzter Wert) 0 %: Der Betrag ist stillschweigend viel zu hoch. Es erscheint weder Fehler noch Warnung.",
         "falscher Betrag", "wahrscheinlich", "Richtlinie S. 10 (Tabelle endet bei 40/60/80 Jahren; S. 9: bei mangelfreiem Zustand bis 20 Jahre); Testfall T5.",
         "Tabelle1 D64:D83 und E44:E83 mit 10 füllen (letzter Tabellenwert) – oder in E15 usw. mit MIN(…;Tabellenende) begrenzen"),
    Fund("F10", "Deckbl UPV", "E23", "Aufwuchs: MAX() statt Wert aus Aufwuchs II",
         "=MAX('Aufwuchs I'!L46;'Aufwuchs II'!L46)",
         "Aufwuchs II enthält die Summe von Aufwuchs I als Übertrag, ist also immer die Gesamtsumme. MAX() liefert aber den größeren der beiden Werte: Steht in Aufwuchs II ein Abzug (negativer Preis), gewinnt Aufwuchs I und der Abzug fehlt.",
         "falscher Betrag", "sicher", "Testfall T8.", "='Aufwuchs II'!L46"),
    Fund("F11", "Beseitigungskosten", "K12:K23, K34:K44", "Zeilensummen nicht auf Cent gerundet",
         "=(F12*H12*J12)",
         "Alle anderen Blätter runden die Zeilensumme (=RUNDEN(A*K;2)). Hier nicht: Bei Faktoren wie 150 % entstehen halbe Cent (z. B. 45,225 €). Angezeigt wird gerundet, gerechnet mit dem Rest, die Summe kann um 1 Cent von der Summe der angezeigten Zeilen abweichen.",
         "falscher Betrag", "sicher", "Testfall T8.", "=RUNDEN(F12*H12*J12;2) in allen 23 Zeilen"),
    Fund("F12", "Baulichkeiten, Beseitigungskosten, Mängelliste", "Baulichkeiten F104, G104, H105, I19, I28, I37; Beseitigungskosten F12:F21, H12, J12:J21, A26; Mängelliste F44",
         "Beispielwerte in der leeren Vorlage",
         "5 m³ × 50 € „nicht beräumt“, 150 € Fäkalienabfuhr, 1/2/3 m² „nicht berechnete Baufläche“, Begründungstext, 0,05 Mängelabzug",
         "Die leere Vorlage ist nicht leer: 5 m³ × 50 € „nicht beräumt“ (250 €), 150 € Fäkalienabfuhr, 1/2/3 m² „nicht berechnete Baufläche“, ein Begründungstext zu bitumenhaltigen Stoffen und 5 % Mängelabzug in der Mängelliste (angezeigt als „0“). Wer sie nicht löscht, zieht Abzüge ab, die es nicht gibt. Das Deckblatt zeigt in der leeren Vorlage −400,00 €.",
         "falscher Betrag", "sicher", "Testfall T0.", "Zellen leeren (Mengen/Preise/Texte); ggf. nur Zeile 12 „Einbehalt Fäkalienabfuhr“ als Hinweis stehen lassen", top=4),
    Fund("F13", "Baulichkeiten", "K77 (Auswahlliste)", "Abzug „nicht geleerte Abwassergrube“: Liste bietet −100 € statt −150 €",
         "Datenüberprüfung K77: Liste „-100“",
         "Die Richtlinie sagt: Bei nicht geleerten Abwassersammelanlagen ist die Bewertung um 150 € zu senken. Die Vorlage bietet nur −100 € an; H105 zeigt dagegen 150 €.",
         "falscher Betrag", "sicher", "Richtlinie S. 14.", "Liste auf „-150“ ändern"),
    Fund("F14", "Baulichkeiten", "E63 (Auswahlliste)", "Laubenvorplatz a): Liste enthält 225 €, die Richtlinie nicht",
         "Datenüberprüfung E63: „0,275,225,175“",
         "Die Richtlinie kennt für Laubenvorplätze bis 275 € (verlängertes Dach), bis 175 € (hochwertig angebaut) und bis 80 € (einfach angebaut). 225 € kommt nicht vor.",
         "falscher Betrag", "bitte BV fragen", "Richtlinie S. 11.", "Liste „0,275“ für a) integriert; die Werte 175/80 gehören zu b)"),
    Fund("F15", "Baulichkeiten, Mängelliste", "E14, E23, E32, E49 (G63, G67); Mängelliste F5:F44, F46", "Mängelabzug: nicht verknüpft, keine Obergrenze, uneinheitliche Einheit",
         "Mängel-Prozente werden von Hand in E14/E23/E32/E49 eingetragen (Zahl = Prozent), in G67 als Bruchteil (0,1 = 10 %), in der Mängelliste ebenfalls als Bruchteil",
         "Die Mängelliste hat „Abzug konkret“ und eine Summe, die Baulichkeiten lesen sie aber nicht: Dieselbe Zahl muss doppelt (und in anderer Einheit) eingegeben werden. 0,1 in E14 bedeutet 0,1 %, nicht 10 %. Es gibt keine Obergrenze: Die Höchstwerte der Mängeltabelle summieren sich auf 120 %; bei Mängel über 100 % wird der NHK-Betrag negativ.",
         "falscher Betrag", "sicher", "Richtlinie S. 6–8 (Abzug maximal je Bauteil); Testfall T5.", "Datenüberprüfung 0–100 für E14/E23/E32/E49; Verweis aus der Mängelliste (Summe × 100)"),
    Fund("F16", "Außenanlagen", "J3", "Versiegelbare Fläche kann negativ oder ungerundet sein",
         "=(E3-24)/100*6",
         "Bei kleiner oder fehlender Parzellenfläche erscheint ein negativer Wert (−1,44 m² in der leeren Vorlage, −0,24 m² bei 20 m²); keine Rundung. Die Regel selbst stimmt mit der Richtlinie überein (6 % der Fläche ohne 24 m²), gilt aber nur, wenn die Laube genau 24 m² hat („bzw. genehmigte bebaute Grundfläche“).",
         "falscher Betrag", "sicher", "Richtlinie S. 14 (Wege- und Platzbefestigung); Testfall T0, T8.", "=MAX(0;RUNDEN((E3-24)/100*6;2))"),
    # ---------------------------------------------------------------------------- falsche oder fehlende Angabe
    Fund("F17", "Beseitigungskosten", "M29", "Textbaustein „Aussenhecke auf zul. Höhe (1,80m)“ widerspricht der Richtlinie",
         "„Aussenhecke auf zul. Höhe (1,80m) schneiden“",
         "Die Richtlinie lässt Hecken höchstens 1,25 m hoch zu; dieselbe Seite der Vorlage (A51, M32) nennt 1,25 m.",
         "falsche oder fehlende Angabe", "sicher", "Richtlinie S. 14 (Einfriedungen) und S. 18 (Hecken); Vorlage A51, M32.", "„Außenhecke auf zul. Höhe (1,25 m) schneiden“"),
    Fund("F18", "Mängelliste", "A5:D43, E-Spalte", "Mängeltabelle weicht von der Richtlinie ab",
         "siehe Beleg",
         "Es fehlen Mängel der Richtlinie (u. a. Schornstein, Moose und Flechten, Risse bei Decken, Belag verschlissen, Fliesen fehlerhaft); „Feuchtigkeit in der Dachkonstruktion“ steht zweimal; „unzulässige Höhe der Laube“ steht nicht in der Richtlinie; bei Sonderabfall fehlt „bei Weiternutzung“.",
         "falsche oder fehlende Angabe", "sicher", "Richtlinie S. 6–8.", "Zeilen ergänzen/streichen laut Beleg; Prozent-Höchstwerte sind richtig"),
    Fund("F19", "Mängelliste ↔ Baulichkeiten", "F46; Baulichkeiten E14 …", "Mängelabzug ist nicht begründet und nicht verknüpft",
         "Mängelliste hat nur „Abzug konkret“ und die Summe F46",
         "Die Richtlinie verlangt, dass die Senkung der NHK „rechnerisch dargestellt und erläutert“ wird, die Mängel „genau zu bezeichnen“ sind und die Abzüge „textlich zu erläutern“. Es gibt keine Spalte für Erläuterung/Begründung, die Summe landet nicht im Blatt Baulichkeiten (siehe F15), das Zahlenformat der Summe zeigt „0“.",
         "falsche oder fehlende Angabe", "wahrscheinlich", "Richtlinie S. 4 und S. 8.", "Spalte „Erläuterung“ ergänzen; E14 =Mängelliste!F46*100 (Laube) bzw. Teil-Summen je Bauteil"),
    Fund("F20", "Baulichkeiten", "H13, I13 (und I22, I31, I48), C13", "Keine Begründung der Restnutzungsdauer, kein Nachweis „Baujahr nachgewiesen/geschätzt“",
         "Felder „Baujahr“ und „Restnutzungsjahre“ ohne Zusatz",
         "Die Richtlinie verlangt, die Restnutzungsdauer „im Einzelnen zu begründen“ und das Jahr der Errichtung nachzuweisen oder als geschätzt zu kennzeichnen. Die Restnutzungsjahre sind außerdem reine Handeingabe (keine Rechnung aus dem Baujahr).",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 9.", "Feld „Begründung / Nachweis Baujahr“ je Bauteil"),
    Fund("F21", "Baulichkeiten", "–", "Kein Bewertungsblock für Kellerräume",
         "„Keller“ kommt nur in einem Beseitigungs-Textbaustein vor",
         "Die Richtlinie lässt vertraglich zulässige Kellerräume bewerten (Innenmaße, Bauklasse II, Restnutzungsjahre nicht über denen der Laube). Dafür gibt es keinen Block.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 11 c).", "Block „2.x Keller“ oder Hinweis, den Keller als Anbau (A1/A2) einzutragen"),
    Fund("F22", "Baulichkeiten", "A82:K100", "Baubeschreibung ohne die Pflichtangaben der Richtlinie",
         "freies Textfeld",
         "Die Richtlinie verlangt für die Baubeschreibung Angaben zu Wasseranschluss, WC, Trockentoilette und Elektrizitätsanschluss (auch „fehlt“). Das Textfeld nennt sie nicht.",
         "falsche oder fehlende Angabe", "wahrscheinlich", "Richtlinie S. 24.", "Vier Zeilen vorbelegen: Wasser / WC / Trockentoilette / Elektro: vorhanden – fehlt"),
    Fund("F23", "Baulichkeiten", "F74, E76, A76:D76", "Abwasser: Grundwert, Sanierung und Tank-Regel nicht abgebildet",
         "F74 = 470 (Zahl), E76 nur „fehlt“",
         "Die Richtlinie kennt drei Grundwerte (Schachtringe 180, sanierte Anlagen 370, Kunststoff/monolithisch 470) und die Altersregel 16 Jahre für Tanks und Betongruben. Vorbelegt ist nur 470; ein Feld für Anlagenart und Sanierungsdatum fehlt.",
         "falsche oder fehlende Angabe", "wahrscheinlich", "Richtlinie S. 13–14.", "Auswahlliste „470,370,180“ an F74; Zusatzzeile Anlagenart/Sanierungsdatum"),
    Fund("F24", "Baulichkeiten", "F101, I38", "„Bewertete bebaute Fläche“ ohne Laubenvorplatz und ohne Prüfung gegen 24 m²",
         "F101 =A17+A26+A35+A52",
         "Ein überdachter Laubenvorplatz zählt zur Laube und darf zusammen mit ihr 24 m² nicht überschreiten. F71 fehlt in der Summe. Weder F101 noch I38 werden weiterverwendet oder geprüft (besser als Hinweis, nicht als Sperre).",
         "falsche oder fehlende Angabe", "wahrscheinlich", "Richtlinie S. 3 (Nr. 2, 3) und S. 11 d).", "=A17+A26+A35+A52+F71 und Hinweiszelle „über 24 m²“"),
    Fund("F25", "Baulichkeiten", "D41, F41, A40:K44", "Gewächshaus: NHK 160 € nicht vorbelegt, Grenzen fehlen",
         "D41 leer, F41 = 10 % Mängel (fest)",
         "Die Richtlinie nennt NHK 160 €/m², höchstens 12 m² Grundfläche und 2,20 m Firsthöhe. Die Vorlage nennt die Grenzen nicht. Der Mängelabzug von 10 % ist als Vorbelegung ohne Grundlage.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 11.", "D41 = 160 vorbelegen, Hinweis „max. 12 m², 2,20 m First“"),
    Fund("F26", "Baulichkeiten", "Schuppenblock A47:K53", "Schuppen: Mindestabzug 3 % jährlich nicht geprüft",
         "Altersabzug nur aus der Restnutzungstabelle",
         "Für „sonstige bauliche Anlagen“ (Gewächshaus, Schuppen, Keller, Laubenvorplatz) verlangt die Richtlinie „mindestens 3 % jährlich“. Beim Schuppen rechnet die Vorlage nur mit der Restnutzungstabelle der Bauklasse, die bei hoher Restnutzung weniger abzieht.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 11.", "Entscheidung des BV: Tabelle oder 3 %/Jahr (größerer Abzug)"),
    Fund("F27", "Außenanlagen, Aufwuchs I/II", "Spalten A–L", "Keine Spalte für Alter/Zustand, keine Hinweise auf Höchstwerte",
         "Spalten: Anzahl, Einheit, Art, Einzelpreis, Gesamtpreis",
         "Die Richtlinie verlangt für alle Außenanlagen eine Alterswertminderung von mindestens 3 % jährlich und Abschläge je Zustand; die Vorlage hat dafür kein Feld, der Abzug steckt unsichtbar im Einzelpreis. Höchstwerte, Mengengrenzen (Rasen 100 m², Hecke 1,25 m) und die Pauschalen (Wasseranschluss 500 €) stehen nicht im Blatt.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 12 und S. 16–18.", "Spalten „Alter/Zustand“ und „Höchstwert lt. Richtlinie“ ergänzen oder bewusst weglassen (BV)"),
    Fund("F28", "Baulichkeiten", "H7, I7", "Baupreisindex „für 2025“ ist veraltet?",
         "H7 = 179,2 mit Beschriftung „für 2025“",
         "Die Richtlinie nimmt den Index des IV. Quartals des Vorjahres für das ganze Folgejahr. Für Wertermittlungen im Jahr 2026 braucht es den Wert vom Herbst 2025; die Vorlage trägt noch „für 2025“. Der Wert kann nicht gegen die Statistik geprüft werden (kein Netz in der Prüfung).",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 3 (Nr. 6) und S. 11; Amt für Statistik Berlin-Brandenburg.", "Jährlich aktualisieren; Quelle/Quartal in I7 nennen"),
    Fund("F29", "Baulichkeiten", "C75, D75, H79 (und AnlIV Blatt 3 B49)", "Altersabzug 3 % je Jahr: linear oder vom Restwert?",
         "C75 =C74*3 (linear), H79 =G78/100*3",
         "Die Vorlage zieht 3 % vom Grundwert je Jahr ab (16 Jahre = 48 %). Die App rechnet nach Vorgabe des Wertermittlers „vom Restwert“ (16 Jahre = 38,57 %, PR #95). Die Richtlinie sagt „von den Bis-Werten ausgehend … jährlicher Abzug von 3 %“ – das lässt sich beides lesen. Die Rechnung der App und die der Vorlage weichen damit ab.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 13 (Abwassersammelanlagen) und S. 14 (Kanalisationsanschlüsse).", "BV legt fest; dann dieselbe Formel in Vorlage und App"),
    Fund("F30", "Baulichkeiten, Beseitigungskosten", "K77, H105, Beseitigungskosten Zeile 12", "„Nicht geleerte Grube“ an drei Stellen, 150 €/100 € und Gefahr des doppelten Abzugs",
         "K77 (−100 per Liste), H105 = 150 (angezeigt, nicht verrechnet), Beseitigungskosten H12 = 150 (vorbelegt)",
         "Die Richtlinie nennt den Betrag zweimal: „Bewertung um 150 € senken“ (S. 14) und „Abfuhrkosten pauschal 150 €“ als Beseitigungskosten (S. 20). In der Vorlage gibt es drei Stellen mit unterschiedlichen Beträgen; H105 wird nirgends gerechnet.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 14 und S. 20.", "BV legt fest, ob beide Abzüge gewollt sind; H105 mit K77 verbinden"),
    Fund("F31", "Beseitigungskosten", "M30", "Textbaustein „Beetflächen auf 10 % Parzellenfläche erhöhen“ ohne Fundstelle",
         "„Beetflächen auf 10% Parzellenfläche erhöhen“",
         "In der Richtlinie steht 10 % nur bei Hochbeeten (S. 15). Ob der Baustein einen anderen Hintergrund hat (Gartenordnung, Kleingartengesetz: ein Drittel), ist nicht erkennbar.",
         "falsche oder fehlende Angabe", "bitte BV fragen", "Richtlinie S. 15, S. 21 („Kosten für das Anlegen von Beeten“).", "Quelle nennen oder Text ändern"),
    # ---------------------------------------------------------------------------- Bedienung
    Fund("F32", "alle Eingabeblätter", "Blattschutz", "Blattschutz ist nicht aktiv, obwohl das Formalien-Blatt ihn verspricht",
         "Formalien H23: „Die Blätter sind geschützt …“; im Dateiinhalt steht bei den meisten Blättern ein Schutz-Element ohne Schalter („sheet=1“ fehlt), beim Deckblatt gar keins",
         "Alle Formelzellen sind frei überschreibbar; ein versehentlich überschriebener Wert fällt nicht auf. Nur das versteckte Blatt Tabelle1 ist (mit Passwort) geschützt.",
         "Bedienung", "sicher", "Dateiinhalt (xl/worksheets/sheet*.xml); Formalien H23–H27.", "Blattschutz ohne Passwort wieder einschalten (nach Korrektur von F33/F34)"),
    Fund("F33", "Formalien, Außenanlagen, Beseitigungskosten, Baulichkeiten", "Formalien E17; Außenanlagen A8:L11, B15:B18, C15:J45; Beseitigungskosten B12:K44; Baulichkeiten I19, I28, I37",
         "Eingabefelder sind als „gesperrt“ markiert",
         "Zellschutz „gesperrt“ an Zellen, in die der Wertermittler schreiben muss",
         "Sobald der Blattschutz (F32) wieder eingeschaltet wird, kann man die Vorgangs-Nr., die Beschreibung und die Art der Außenanlagen, alle Beseitigungspositionen und die „nicht berechnete Baufläche“ nicht mehr eintragen.",
         "Bedienung", "sicher", "Zellschutz je Zelle (Eigenschaft „Gesperrt“).", "Diese Zellen vor dem Einschalten des Schutzes entsperren"),
    Fund("F34", "Baulichkeiten, Mängelliste", "Baulichkeiten J35, J44, F101; Mängelliste E1, E2", "Formelzellen sind als „nicht gesperrt“ markiert",
         "Formeln in entsperrten Zellen",
         "Auch bei aktivem Blattschutz könnte man diese Formeln überschreiben.",
         "Bedienung", "sicher", "Zellschutz je Zelle.", "Diese Zellen sperren"),
    Fund("F35", "Baulichkeiten", "G12, G21, G30, G47; I13, I22, I31, I48", "Leere/ungültige Auswahl führt zu #NV im ganzen Protokoll",
         "Auswahl- und Zahlenfelder ohne Prüfmeldung bei Löschen/Text/negativ",
         "Wird die Bauklasse-Auswahl gelöscht (Entf-Taste) oder stehen die Restnutzungsjahre als Text/negativ da, wird die Rechenkette zu #NV; der Fehler läuft bis K81 und auf das Deckblatt. Eine Fehlermeldung gibt es nicht.",
         "Bedienung", "sicher", "Testfälle T5, T9.", "Datenüberprüfung für I13, I22, I31, I48 (ganze Zahl 0–80); bei leerer Typ-Auswahl WENNFEHLER in Tabelle1 D88/F88/B90 (und die drei weiteren Blöcke)"),
    Fund("F36", "Deckbl UPV, alle Köpfe", "z. B. C7:C10, C11, E13:E17", "Leere Formalien erscheinen als „0“ und als Datum 1899/1900",
         "=Formalien!$E$6 ohne Abfrage",
         "Verweise auf leere Zellen liefern 0. Auf allen Blättern stehen „0“ statt leer, und Datumsfelder zeigen „30.12.1899“ (Excel: 00.01.1900).",
         "Bedienung", "sicher", "Testfall T0, Druckvorschau.", "=WENN(Formalien!$E$6=\"\";\"\";Formalien!$E$6); beim Datum ebenso"),
    Fund("F37", "Mängelliste", "F5:F44, F46", "Zahlenformat „0“ zeigt Prozent als 0",
         "Format „0“ in Zellen, die Bruchteile (0,05) erwarten",
         "Eine Eingabe 0,05 (5 %) erscheint als „0“, die Summe F46 ebenso; ohne Blick in die Eingabezeile kann man nicht sehen, was eingetragen ist.",
         "Bedienung", "sicher", "Zahlenformat der Zellen.", "Format „0 %“"),
    Fund("F38", "Baulichkeiten, Formalien, Aufwuchs II", "Seitenumbruch Zeile 54; Druckbereiche", "Druckbild: leere Seite, fehlende Druckbereiche",
         "Baulichkeiten: Skalierung 86 %, manueller Umbruch nach Zeile 54; Formalien und Aufwuchs II ohne Druckbereich",
         "Seite 1 der Baulichkeiten ist rechnerisch länger als die Druckseite; die Zeile 54 („Seite 2“) landet auf einer eigenen, fast leeren Seite (in LibreOffice sichtbar; bitte in der Excel-Seitenansicht prüfen). Formalien druckt die Hinweise rechts mit und kommt auf 2–3 Seiten; Aufwuchs II hat keinen Druckbereich.",
         "Bedienung", "wahrscheinlich", "Berechnung Zeilenhöhen × Skalierung; PDF-Ausdruck der Vorlage über LibreOffice.", "Skalierung 82 % oder „Auf eine Seite“; Druckbereiche setzen"),
    Fund("F39", "versteckte Blätter", "10 versteckte Blätter (Anl IV Schuppen/Dach, Deckbl-Varianten, Beseitigungsauflagen, Bewertung …, Tabelle1)", "Versteckte Altblätter mit über 300 Fehlerwerten (#BEZUG!)",
         "z. B. „Anl IV Schuppen“ 167, „Anl IV Dach“ 165 Zellen mit #BEZUG!; „Beseitigungsauflagen“ H22 =SUMME(E22*G33) (falsche Zeile)",
         "Die Blätter sind versteckt und werden nicht gedruckt, sind aber in der Datei und täuschen Fehler vor (z. B. bei Suche nach Fehlerwerten). Die Deckblatt-Varianten verweisen auf zerstörte Bezüge; Beseitigungsauflagen rechnet in H22 mit der falschen Zeile.",
         "Bedienung", "sicher", "Bestandsaufnahme (alle Formeln mit openpyxl gelesen).", "Blätter löschen (Original behalten) oder reparieren; ungenutzte Namen entfernen"),
    # ---------------------------------------------------------------------------- kosmetisch
    Fund("F40", "Baulichkeiten", "J52, B19, B28, B37, E34, H7-Bezüge, H42, N79:N80", "Falsche Beschriftung beim Schuppen, feste Kopien statt Verweise, Reste",
         "J52 =$B$21 („Wert BT A1“ beim Schuppen); B19 „L“, B28 „A1“, B37 „A2“ als Text; E34 =$E$16; E16/E25/E43/E51 =H7/100 ohne $; H42 = 5; N79 „Summe K56:K74“ mit anderer Formel",
         "Die Beschriftung neben dem Schuppenwert (J52) zeigt das Kürzel von A1 statt „S oder U“. Ändert jemand die Bezeichnung (B12, B21, B30), bleiben die Kopien in B19/B28/B37 alt. E34 holt den Index von der Laube statt aus H7. Die 5 in H42 hat keine Funktion, steht aber auf dem Ausdruck neben „% Altersabzug“; der Rest in N79/N80 beschreibt einen anderen Bereich, als er rechnet.",
         "kosmetisch", "sicher", "Formelvergleich der vier Bauteilblöcke; Ausdruck Blatt 1.", "J52 =$B$47; B19 =B12, B28 =B21, B37 =B30; E34 =$H$7/100; H42, N79:N80 löschen"),
    Fund("F41", "Tabelle1", "D103, F5, Namen", "Versteckte Hilfstabelle mit kaputtem Rest",
         "D103: 22-fach verschachteltes WENN, das mit =J19:J109 endet; F5 =-I10; ungenutzte Namen (BKL.Beschreibung, Jürgen, rose, ungeprüft, uwe.2)",
         "Funktioniert, solange eine Auswahl aus der Liste getroffen ist; alle anderen Bauteile nutzen SVERWEIS. Der Rest kann bei Änderungen der Liste zu falschen Werten führen.",
         "kosmetisch", "sicher", "Formelvergleich Tabelle1 D88/D93/D98/D103.", "D103 =SVERWEIS(D105;uwe;2) wie die anderen Bauteile; Namen löschen"),
    Fund("F42", "mehrere", "siehe Beleg", "Tippfehler und Reste in Beschriftungen",
         "u. a. „horizentral“, „maßgrecht“, „Gladichtung“, „enstsprechend“, „Sepzember“, „entfernene“, „Anfalll“, „Aussenhecke“, „lotgerecht“, „nich ordnungsgemäß“, „unddicht“; „wir.“ in Formalien H6:H19",
         "Rechtschreibfehler in Mängelliste, Beseitigungskosten, Baulichkeiten; die „wir.“ in Formalien sind Reste (vermutlich verlorene Aufzählungszeichen).",
         "kosmetisch", "sicher", "Textsuche in allen Zellen.", "Texte korrigieren (siehe Beleg)"),
    Fund("F43", "Beseitigungskosten", "A21:A24, A42:A44", "Nummerierung fehlerhaft",
         "„9.“ steht zweimal (A42, A43); „12.“ steht in der Summenzeile (A24)",
         "Abschnitt 2 hat die Nummern 1–9, 9, 10. In Abschnitt 1 sind die Positionen 1–11 nummeriert, die „12.“ steht in der Summenzeile.",
         "kosmetisch", "sicher", "Blattansicht.", "A43 „10.“, A44 „11.“, A24 leeren"),
    Fund("F44", "Skizze Außenanlage-Aufwuchs", "A1", "Überschrift lautet „Skizze – Baulichkeiten“",
         "„Skizze – Baulichkeiten“",
         "Das Blatt ist für Außenanlagen und Aufwuchs gedacht; die Überschrift ist von „Skizze Bau“ kopiert.",
         "kosmetisch", "sicher", "Blattansicht.", "„Skizze – Außenanlagen und Aufwuchs“"),
    Fund("F45", "Baulichkeiten", "L62", "Hinweis zum Laubenvorplatz nennt falsche Zellen",
         "„… unter E58 bzw. B62 …“",
         "Die Felder heißen E63 (NHK a) und B67 (Güte b); „Anfalll“ ist ein Tippfehler.",
         "kosmetisch", "sicher", "Blattansicht.", "„… unter E63 bzw. B67 bei einem Anfall von a) oder b) …“"),
    Fund("F46", "Deckblatt, Formalien, Abkürzungen", "Deckbl A34/A38, Formalien I16/I17, Abkürzungen C12, C38 u. a.", "Widersprüchliche Texte und uneinheitliche Begriffe",
         "Widerspruch „direkt an die Wertermittler“ (A34) und „an den Bezirksverband“ (A38); „Dichtheits-/Dichtigkeitsgutachten“; „Aufwuchs 2 ggf. 3“ statt „Aufwuchs I/II“; „Anlage 4 Blatt 1“ vs. „Seite 1 und 2“; „Anlage IV“ vs. „Anlage 4“; U = Unterstand/Toilettenhaus; „RiLW“ nennt die Richtlinie anders als das Deckblatt",
         "Der Unterpächter erfährt zwei verschiedene Stellen für den Widerspruch (die Richtlinie sagt: der Bezirksverband bearbeitet ihn, S. 23). Die übrigen Begriffe sind uneinheitlich.",
         "kosmetisch", "bitte BV fragen", "Richtlinie S. 23 (Einspruch); Textvergleich.", "Eine Stelle festlegen, Begriffe vereinheitlichen"),
]
FUNDE_NACH_NR = {f.nr: f for f in FUNDE}
assert set(FUNDE_NACH_NR) == set(PRUEFUNGEN), set(FUNDE_NACH_NR) ^ set(PRUEFUNGEN)


# ---------------------------------------------------------------------------------------------------
# 5. Weitere, nicht katalogisierte Prüfungen (Hinweise)
# ---------------------------------------------------------------------------------------------------

def pruefe_tabellen(m: Mappe) -> list[str]:
    """Tabelle1 (Restnutzung, NHK, Abriss) gegen die Richtlinie. Liefert Abweichungen (leer = alles gleich)."""
    ws = m.blatt(T1)
    abw = []
    for name, col, ref in (("I", "C", RN_I), ("II/III", "D", RN_III), ("IV", "E", RN_IV)):
        for y, soll in enumerate(ref):
            ist = ws[f"{col}{3 + y}"].value
            if not fast_gleich(ist, soll, 1e-9):
                abw.append(f"Restnutzung Bauklasse {name}, {y} Jahre: Vorlage {ist}, Richtlinie {soll}")
    for r in range(54, 76):
        stufe, nhk = ws[f"K{r}"].value, ws[f"L{r}"].value
        if stufe in NHK and not fast_gleich(nhk, NHK[stufe], 1e-9):
            abw.append(f"NHK {stufe} ({ws[f'J{r}'].value}): Vorlage {nhk}, Richtlinie {NHK[stufe]}")
    for k, v in zip(("I", "II", "III", "IV"), (ws[f"M{r}"].value for r in range(67, 71))):
        if not fast_gleich(v, ABRISS[k], 1e-9):
            abw.append(f"Abriss Bauklasse {k}: Vorlage {v}, Richtlinie {ABRISS[k]}")
    return abw


def konstanten_in_formeln(m: Mappe) -> list[str]:
    """Formeln mit eingebauten Zahlen (außer 0, 1, 2, 100): als Hinweis, ob sie zur Richtlinie passen."""
    res = []
    for blatt in (FO, DB, BL, MG, AA, AW1, AW2, BE):
        ws = m.blatt(blatt)
        for z in ws.iter_rows():
            for c in z:
                if ist_formel(c.value):
                    ohne = re.sub(r'"[^"]*"', "", c.value)
                    ohne = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", ohne)
                    zahlen = [x for x in re.findall(r"\d+(?:\.\d+)?", ohne) if x not in ("0", "1", "2", "100", "99")]
                    if zahlen:
                        res.append(f"{blatt}!{c.coordinate}: {de(c.value)} (Zahlen {', '.join(zahlen)})")
    return res


def zirkelbezuege(m: Mappe) -> list[str]:
    """Sucht Zirkelbezüge (Zelle hängt über Formeln von sich selbst ab). Bereiche werden aufgelöst, Namen nicht."""
    from openpyxl.utils import get_column_letter, range_boundaries
    ref = re.compile(r"(?:(?:'([^']+)'|([A-Za-z0-9_ÄÖÜäöüß.]+))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?")
    graph: dict[tuple[str, str], list[tuple[str, str]]] = {}
    for ws in m.wb:
        for zeile in ws.iter_rows():
            for c in zeile:
                if not ist_formel(c.value):
                    continue
                f = re.sub(r'"[^"]*"', "", c.value)
                ziele = []
                for treffer in ref.finditer(f):
                    blatt = treffer.group(1) or treffer.group(2) or ws.title
                    a = treffer.group(3).replace("$", "")
                    e = (treffer.group(4) or treffer.group(3)).replace("$", "")
                    c1, r1, c2, r2 = range_boundaries(f"{a}:{e}")
                    if (c2 - c1 + 1) * (r2 - r1 + 1) > 3000:
                        continue
                    ziele += [(blatt, f"{get_column_letter(cc)}{rr}") for rr in range(r1, r2 + 1) for cc in range(c1, c2 + 1)]
                graph[(ws.title, c.coordinate)] = ziele
    zyklen, farbe = [], {}
    for start in graph:
        if start in farbe:
            continue
        stapel = [(start, iter(graph.get(start, [])))]
        farbe[start] = 1
        pfad = [start]
        while stapel:
            knoten, it = stapel[-1]
            weiter = next(it, None)
            if weiter is None:
                farbe[knoten] = 2
                stapel.pop()
                pfad.pop()
                continue
            if weiter not in graph:
                continue
            if farbe.get(weiter) == 1:
                zyklen.append(" → ".join(f"{b}!{z}" for b, z in pfad[pfad.index(weiter):] + [weiter]))
            elif weiter not in farbe:
                farbe[weiter] = 1
                stapel.append((weiter, iter(graph[weiter])))
                pfad.append(weiter)
    return zyklen


def formel_zellen_mit_fehlerwerten(m: Mappe) -> list[str]:
    """Zellen mit zwischengespeichertem Fehlerwert (#BEZUG! usw.) in sichtbaren Blättern."""
    wb = openpyxl.load_workbook(m.pfad, data_only=True)
    res = []
    for ws in wb:
        if ws.sheet_state != "visible":
            continue
        for z in ws.iter_rows():
            for c in z:
                if ist_fehlerwert(c.value):
                    res.append(f"{ws.title}!{c.coordinate} = {c.value}")
    return res


def bestandsaufnahme(m: Mappe) -> list[dict[str, Any]]:
    schutz = m.blattschutz()
    zeilen = []
    for ws in m.wb:
        formeln = sum(1 for z in ws.iter_rows() for c in z if ist_formel(c.value))
        zeilen.append({
            "blatt": ws.title, "sichtbar": ws.sheet_state == "visible", "formeln": formeln,
            "verbunden": len(ws.merged_cells.ranges), "listen": len(ws.data_validations.dataValidation),
            "bedingte_formate": len(ws.conditional_formatting), "druckbereich": ws.print_area or "–",
            "schutz": schutz.get(ws.title, "?"),
        })
    return zeilen


# ---------------------------------------------------------------------------------------------------
# 6. Ablauf: prüfen, ausgeben
# ---------------------------------------------------------------------------------------------------

def pruefe_alles(pfad: Path, mit_rechnung: bool = True, lo_pfad: str | None = None):
    m = Mappe(pfad)
    lauf, soll = None, {}
    hinweis = ""
    if mit_rechnung:
        lo = finde_libreoffice(lo_pfad)
        if lo is None:
            hinweis = "LibreOffice nicht gefunden – nur statische Prüfungen (mit --lo PFAD angeben oder --ohne-rechentests)."
        else:
            lauf = Lauf(m, lo)
            soll = baue_szenarien(m, lauf)
            lauf.rechne()
    ctx = Ctx(m, lauf, soll)
    befunde = {f.nr: PRUEFUNGEN[f.nr](ctx) for f in FUNDE}
    return m, ctx, befunde, hinweis


def ausgabe_konsole(m: Mappe, ctx: Ctx, befunde: dict[str, Befund], hinweis: str, ausfuehrlich: bool) -> int:
    print(f"Datei: {m.pfad}")
    if hinweis:
        print(hinweis)
    print(f"{'Nr.':4} {'Zustand':26} {'Gewissheit':16} Fund")
    offen_sicher = 0
    for f in FUNDE:
        b = befunde[f.nr]
        zustand = {True: "VORHANDEN", False: "behoben", None: "nicht automatisch prüfbar"}[b.vorhanden]
        print(f"{f.nr:4} {zustand:26} {f.gewissheit:16} {f.blatt}: {f.titel}")
        if ausfuehrlich:
            print(f"       {b.beleg}")
        if b.vorhanden and f.gewissheit == "sicher":
            offen_sicher += 1
    anz = {k: sum(1 for f in FUNDE if befunde[f.nr].vorhanden is k) for k in (True, False, None)}
    print(f"\nVorhanden: {anz[True]}, behoben: {anz[False]}, nicht automatisch prüfbar: {anz[None]} "
          f"(davon sicher und noch vorhanden: {offen_sicher})")
    abw = pruefe_tabellen(m)
    print(f"Tabelle1 (Restnutzung, NHK, Abriss) gegen Richtlinie: {'alles gleich' if not abw else str(len(abw)) + ' Abweichungen'}")
    for a in abw[:20]:
        print("   ", a)
    return offen_sicher


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Prüft die WEP-Vorlage des BV Pankow auf Fehler (siehe docs/VORLAGE-FEHLERBERICHT.md).")
    ap.add_argument("datei", nargs="?", default=str(STANDARD_VORLAGE), help="zu prüfende Excel-Datei (Standard: docs/WEP_Vorlage.xlsx)")
    ap.add_argument("--ohne-rechentests", action="store_true", help="nur statische Prüfungen (ohne LibreOffice)")
    ap.add_argument("--lo", help="Pfad zu soffice (LibreOffice), falls nicht automatisch gefunden")
    ap.add_argument("--ausfuehrlich", "-v", action="store_true", help="Beleg je Fund anzeigen")
    ap.add_argument("--bericht", help="Fehlerbericht als Markdown schreiben")
    ap.add_argument("--pdf", help="Fehlerbericht zusätzlich als PDF schreiben (braucht --bericht und playwright)")
    ap.add_argument("--korrekturvorschlag", help="korrigierte Kopie der Datei schreiben (Änderungen gelb, mit Kommentar)")
    args = ap.parse_args(argv)

    pfad = Path(args.datei)
    if not pfad.exists():
        print(f"Datei nicht gefunden: {pfad}", file=sys.stderr)
        return 2
    m, ctx, befunde, hinweis = pruefe_alles(pfad, not args.ohne_rechentests, args.lo)
    offen = ausgabe_konsole(m, ctx, befunde, hinweis, args.ausfuehrlich)

    if args.bericht:
        schreibe_bericht(m, ctx, befunde, Path(args.bericht), Path(args.pdf) if args.pdf else None)
    if args.korrekturvorschlag:
        schreibe_korrekturvorschlag(pfad, Path(args.korrekturvorschlag))
    return offen


# ---------------------------------------------------------------------------------------------------
# 7. Korrekturvorschlag: korrigierte KOPIE (gelb markiert, mit Kommentar). Das Original bleibt unverändert.
#    Die Datei wird direkt im XML geändert (nicht über openpyxl), damit Formularfelder (Kontrollkästchen),
#    Zeichnungen, Skizzen und Formatierung erhalten bleiben.
# ---------------------------------------------------------------------------------------------------

NS_M = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
M_ = f"{{{NS_M}}}"
# Reihenfolge der Elemente eines Arbeitsblatts (ECMA-376, CT_Worksheet) – für das Einfügen von <legacyDrawing>
BLATT_REIHENFOLGE = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr",
                     "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
                     "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations",
                     "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
                     "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
                     "legacyDrawingHF", "drawingHF", "picture", "oleObjects", "controls", "webPublishItems",
                     "tableParts", "extLst"]
AUTOR = "Prüfskript Vorlagenfehler"


@dataclass
class Korrektur:
    fund: str
    blatt: str
    zelle: str            # Zelle, Bereich (A1:B2) oder bei art=„dv_*“ der Bereich der Auswahlliste
    art: str              # formel | text | zahl | leer | format | entsperren | sperren | dv_zahl | dv_ganz | seite | druckbereich | name_loeschen
    neu: Any = None
    kommentar: str = ""
    gelb: bool = True
    zahlenformat: int | None = None
    gruppe: str = ""      # gleichartige Änderungen (z. B. dieselbe Formel je Zeile) erscheinen im Bericht als eine Zeile


def _zellen(bereich: str) -> list[str]:
    from openpyxl.utils import get_column_letter, range_boundaries
    if ":" not in bereich:
        return [bereich]
    c1, r1, c2, r2 = range_boundaries(bereich)
    return [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def _korrekturen() -> list[Korrektur]:
    K = Korrektur
    k: list[Korrektur] = []
    # ---- Baulichkeiten
    k += [
        K("F01", BL, "C33", "formel", "$H$32", "F01: war =$H$23 (Wert von Anbau A1). Anbau A2 rechnet jetzt mit seinem eigenen Wert nach Mängelabzug (H32)."),
        K("F40", BL, "J52", "formel", "$B$47", "F40: war =$B$21 – neben dem Schuppenwert stand das Kürzel „A1“ statt „S oder U“."),
        K("F02", BL, "K81", "formel", "K61+K71+K75+K77+K79", "F02: war =K52+K61+K71+K75+K77+K79. Der Schuppen (K52) steckt schon im Übertrag K61 und wurde doppelt gezählt."),
        K("F03", BL, "C64", "formel", "I63", "F03: war leer. Verbindet „NHK nach Mängelabzug“ (I63) mit dem Altersabzug (I64); vorher wurde der integrierte Laubenvorplatz mit 0 € bewertet."),
        K("F05", BL, "G68", "formel", "MIN(100,E66*H66)", "F05: war die feste Zahl 100. Jetzt Altersabzug = „Altersminderung pro Jahr“ (E66) × „Alter“ (H66), höchstens 100 % – wie beim Gewächshaus (E42). Bitte BV bestätigen."),
        K("F06", BL, "C75", "formel", "(1-0.97^C74)*100", "F06/F29: war =C74*3 (linear, ohne Obergrenze – ab 34 Jahren negativer Wert). Vorschlag nach der Rechnung der App (3 % jährlich vom Restwert, PR #95). Alternative linear: =MIN(100,C74*3). Bitte BV entscheiden (F29). Die 16-Jahre-Regel für Tanks ist weiter nicht umgesetzt (F23).", zahlenformat=2),
        K("F07", BL, "H79", "formel", "1-0.97^G78", "F07/F29: war =G78/100*3 (linear, ohne Obergrenze). Vorschlag wie C75; Alternative linear: =MIN(1,G78*3%). Bitte BV entscheiden (F29).", zahlenformat=10),
        K("F08", BL, "E79", "formel", 'IF(G78="",0,MIN(2200,C79))', "F08: war =IF(G78>0,MIN(2200,C79),0) – ein Anschluss mit Alter 0 wurde mit 0 € bewertet. Bitte BV bestätigen."),
        K("F08", BL, "K79", "formel", "ROUND(E79-E79*H79,2)", "F08: war =IF(C79>74,…,0) – Kosten bis 74 € ergaben immer 0 €; woher die 74 stammen, ist unklar. Bitte BV bestätigen."),
        K("F24", BL, "F101", "formel", "A17+A26+A35+A52+F71", "F24: war =A17+A26+A35+A52. Der überdachte Laubenvorplatz (F71) zählt zu den 24 m² (Richtlinie S. 11 d)."),
        K("F30", BL, "H105", "formel", "-K77", "F30: war die feste Zahl 150, nirgends verrechnet. Zeigt jetzt den Abzug aus K77. Ob zusätzlich die 150 € Abfuhrkosten in den Beseitigungskosten gelten, muss der BV entscheiden."),
        K("F13", BL, "K77", "dv_zahl", (-150, 0), "F13: Auswahlliste „-100“ ersetzt durch Zahl zwischen −150 und 0 (Richtlinie S. 14: Senkung um 150 €)."),
        K("F14", BL, "E63", "dv_zahl", (0, 275), "F14: Auswahlliste „0,275,225,175“ ersetzt durch Zahl zwischen 0 und 275 (Richtlinie S. 11: verlängertes Dach bis 275 €). Bitte BV bestätigen."),
    ]
    for z in ("E14", "E23", "E32", "E49"):
        k.append(K("F15", BL, z, "dv_zahl", (0, 100), "F15: Eingabeprüfung: Mängelabzug in Prozent, 0 bis 100 (0,1 bedeutet 0,1 %, nicht 10 %)."))
    for z in ("I13", "I22", "I31", "I48"):
        k.append(K("F35", BL, z, "dv_ganz", (0, 80), "F35: Eingabeprüfung: ganze Zahl 0 bis 80 (längste Restnutzungstabelle der Richtlinie)."))
    for z, alt in (("F104", "5"), ("G104", "50"), ("I19", "1"), ("I28", "2"), ("I37", "3")):
        k.append(K("F12", BL, z, "leer", None, f"F12: war der Beispielwert {alt}; leere Vorlage soll keine Abzüge enthalten."))
    for z, quelle in (("B19", "B12"), ("B28", "B21"), ("B37", "B30")):
        k.append(K("F40", BL, z, "formel", quelle, f"F40: war ein fest getippter Text; jetzt Verweis auf {quelle}."))
    k += [
        K("F40", BL, "E34", "formel", "$H$7/100", "F40: war =$E$16 (Index über die Laube); jetzt wie bei den anderen Bauteilen aus H7."),
        K("F40", BL, "H42", "leer", None, "F40: war die Zahl 5 ohne Funktion (stand auf dem Ausdruck neben „% Altersabzug“)."),
        K("F40", BL, "N79", "leer", None, "F40: Rest („Summe K56:K74“) mit anderer Formel in N80; gelöscht."),
        K("F40", BL, "N80", "leer", None, "F40: Rest-Summe, nirgends verwendet; gelöscht."),
        K("F45", BL, "L62", "text", "Hinweis: Beim Laubenvorplatz immer darauf achten, dass unter E63 bzw. B67 bei einem Anfall von a) oder b) beim anderen Bauteil 0 bzw. ohne (siehe Dropdown-Listen) eingetragen ist.", "F45: nannte E58 und B62; die Felder sind E63 und B67. „Anfalll“ korrigiert."),
        K("F38", BL, "K54", "seite", 82, "F38: Skalierung von 86 % auf 82 %, damit Seite 1 (Zeilen 1–54) auf eine Druckseite passt und „Seite 2“ nicht auf einer leeren Seite landet. Bitte in der Excel-Seitenansicht prüfen."),
        K("F33", BL, "I19", "entsperren", None, "F33: Eingabefeld „nicht berechnete Baufläche“ (I19, I28, I37) war gesperrt – bei aktivem Blattschutz nicht bedienbar.", gelb=False),
        K("F33", BL, "I28", "entsperren", None, "", gelb=False), K("F33", BL, "I37", "entsperren", None, "", gelb=False),
        K("F34", BL, "J35", "sperren", None, "F34: Formelzellen J35, J44, F101 waren nicht gesperrt.", gelb=False),
        K("F34", BL, "J44", "sperren", None, "", gelb=False), K("F34", BL, "F101", "sperren", None, "", gelb=False),
        K("F34", BL, "H105", "sperren", None, "", gelb=False),
        K("F40", BL, "E16", "formel", "$H$7/100", "F40: war =H7/100 ohne $ (der Bezug verschiebt sich beim Kopieren). E25, E43, E51 ebenso geändert.", gruppe="H7"),
        K("F40", BL, "E25", "formel", "$H$7/100", "", gruppe="H7"), K("F40", BL, "E43", "formel", "$H$7/100", "", gruppe="H7"), K("F40", BL, "E51", "formel", "$H$7/100", "", gruppe="H7"),
    ]
    # ---- Deckblatt
    k.append(K("F10", DB, "E23", "formel", "'Aufwuchs II'!L46", "F10: war =MAX('Aufwuchs I'!L46,'Aufwuchs II'!L46). Aufwuchs II enthält Aufwuchs I als Übertrag; MAX übersah Abzüge in Aufwuchs II."))
    leer_wenn = (("C7", "E6"), ("C8", "E3"), ("C10", "E17"), ("E13", "E12"), ("E14", "E13"), ("E15", "E10"), ("E16", "E11"), ("E17", "E9"))
    for i, (z, q) in enumerate(leer_wenn):
        k.append(K("F36", DB, z, "formel", f'IF(Formalien!${q[0]}${q[1:]}="","",Formalien!${q[0]}${q[1:]})',
                   "F36: zeigte bei leeren Formalien „0“; jetzt leer. Dasselbe gilt für die Köpfe der übrigen Blätter (Vorschlag, dort nicht geändert)." if i == 0 else "", gruppe="IF"))
    k.append(K("F36", DB, "C9", "formel", 'IF(Formalien!$E$4="","",IF(Formalien!$E$5="",Formalien!$E$4,Formalien!$E$4&" / "&Formalien!$E$5))', ""))
    k.append(K("F36", DB, "C11", "formel", 'IF(Formalien!$E$14="","",Formalien!$E$14)', "F36: zeigte bei leerem Datum „00.01.1900“; jetzt leer."))
    # ---- Außenanlagen
    k.append(K("F16", AA, "J3", "formel", "MAX(0,ROUND((E3-24)/100*6,2))", "F16: war =(E3-24)/100*6 (negativ/ungerundet bei kleiner oder leerer Parzelle)."))
    # ---- Beseitigungskosten
    for r in list(range(12, 24)) + list(range(34, 45)):
        kom = ""
        if r == 12:
            kom = "F11: Zeilensumme jetzt =ROUND(F*H*J;2) wie auf allen anderen Blättern (K12:K23 und K34:K44); vorher ohne Rundung."
        if r == 21:
            kom = "F04: hier fehlte die Formel (Zeile 9 zählte nicht in K24 und damit nicht auf das Deckblatt). F11: zusätzlich gerundet wie die anderen Zeilen."
        if r == 34:
            kom = "F11: Zeilensumme jetzt =ROUND(F*H*J;2); vorher ohne Rundung."
        k.append(K("F04" if r == 21 else "F11", BE, f"K{r}", "formel", f"ROUND(F{r}*H{r}*J{r},2)", kom, gruppe="K"))
    k += [
        K("F12", BE, "H12", "leer", None, "F12: war 150 € vorbelegt („Einbehalt Fäkalienabfuhr“). Nur eintragen, wenn die Grube nicht geleert ist (Richtlinie S. 20). Siehe auch F30."),
        K("F12", BE, "A26", "leer", None, "F12: vorbelegter Text zu bitumenhaltigen Stoffen stand in jedem Protokoll."),
        K("F17", BE, "M29", "text", "Außenhecke auf zul. Höhe (1,25 m) schneiden", "F17: war „Aussenhecke … (1,80m)“; die Richtlinie erlaubt höchstens 1,25 m."),
        K("F42", BE, "M12", "text", "bitumenhaltige Stoffe sind entsprechend der GefStoffV i.V.m. der TRGS 519 zu demontieren und zu entsorgen.", "F42: „enstsprechend“ korrigiert."),
        K("F42", BE, "M25", "text", "Well-/Ebenasbest ist Sondermüll, es ist entsprechend der GefStoffV i.V.m. der TRGS 519 zu demontieren und zu entsorgen.", "F42: „enstsprechend“ korrigiert."),
        K("F42", BE, "A32", "text", "2. Zu entfernende Außenanlagen laut Skizze Außenanlagen-Aufwuchs", "F42: „entfernene“ korrigiert."),
        K("F42", BE, "A46", "text", "In der Zeit vom 1. März bis 30. September eines jeden Jahres dürfen aus Gründen des Naturschutzes, insbesondere des", "F42: „Sepzember“ korrigiert."),
        K("F43", BE, "A43", "text", "10.", "F43: stand „9.“ (doppelt)."),
        K("F43", BE, "A44", "text", "11.", "F43: stand „10.“."),
        K("F43", BE, "A24", "leer", None, "F43: „12.“ stand in der Summenzeile."),
    ]
    for r in list(range(12, 24)) + list(range(34, 45)):
        for sp in "BFGHJ":
            k.append(K("F33", BE, f"{sp}{r}", "entsperren", None,
                       "F33: Eingabefelder B, F, G, H, J der Zeilen 12–23 und 34–44 waren gesperrt – bei aktivem Blattschutz hätte man nichts eintragen können." if (sp, r) == ("B", 12) else "", gelb=False))
    # ---- Außenanlagen: Eingabefelder entsperren
    for z in ["A8"] + [f"B{r}" for r in range(15, 19)] + [f"C{r}" for r in range(15, 46)]:
        k.append(K("F33", AA, z, "entsperren", None,
                   "F33: Beschreibung (A8), Maßeinheit (B15:B18) und „Art“ (C15:C45) waren gesperrt – bei aktivem Blattschutz nicht bedienbar." if z == "C15" else "", gelb=False))
    k.append(K("F33", FO, "E17", "entsperren", None, "F33: Feld „Vorgangs-Nr.“ war gesperrt.", gelb=False))
    # ---- Mängelliste
    k += [
        K("F42", MG, "D6", "text", "Dach nicht horizontal", "F42: „horizentral“ korrigiert."),
        K("F42", MG, "D26", "text", "nicht maßgerecht", "F42: „maßgrecht“ korrigiert."),
        K("F42", MG, "C34", "text", "Fenster - Glasdichtung fehlerhaft (Kitt)", "F42: „Gladichtung“ korrigiert."),
        K("F42", MG, "C36", "text", "Fensterläden nicht ordnungsgemäß", "F42: „nich“ korrigiert."),
        K("F42", MG, "C38", "text", "Tür undicht", "F42: „unddicht“ korrigiert."),
        K("F42", MG, "C19", "text", "Wände nicht lotrecht", "F42: „lotgerecht“ korrigiert (Richtlinie: „lotrecht“)."),
        K("F18", MG, "D10", "text", "Schornstein fehlerhaft", "F18: „Feuchtigkeit in der Dachkonstruktion“ stand zweimal (C12 und D10); hier ersetzt durch „Schornstein fehlerhaft“ (fehlte). Es fehlen weiter „Schornstein ohne Kehrbescheinigung“, „Moose und Flechten“ (Dach), „Risse“ (Decken), „Belag verschlissen“ (Fußboden), „Fliesen fehlerhaft“ (Ausbau innen) – dafür fehlt Platz im Blatt, bitte BV entscheiden."),
        K("F12", MG, "F44", "leer", None, "F12: war 0,05 vorbelegt (5 % Abzug „nicht fertiggestellt“) und floss in die Summe F46."),
        K("F34", MG, "E1", "sperren", None, "F34: E1 und E2 (Formeln) waren nicht gesperrt.", gelb=False),
        K("F34", MG, "E2", "sperren", None, "", gelb=False),
    ]
    for z in ("F5", "F13", "F16", "F21", "F24", "F28", "F30", "F34", "F40", "F42", "F43", "F44", "F46"):
        k.append(K("F37", MG, z, "format", None, "F37: Zahlenformat „0“ zeigte 5 % (=0,05) als „0“; jetzt „0 %“." if z == "F5" else "", zahlenformat=9))
    # ---- Skizze, Formalien
    k.append(K("F44", SK2, "A1", "text", "Skizze – Außenanlagen und Aufwuchs", "F44: stand „Skizze – Baulichkeiten“ (von Skizze Bau kopiert)."))
    for z in ("H6", "H7", "H8", "H10", "H15", "H16", "H17", "H19"):
        k.append(K("F42", FO, z, "leer", None, "F42: Rest „wir.“ (vermutlich verlorene Aufzählungszeichen) gelöscht." if z == "H6" else ""))
    k.append(K("F38", FO, "A1", "druckbereich", "$A$1:$F$24", "F38: Druckbereich gesetzt (ohne die Hinweise rechts und ohne „Seite 1“ in G1); Formalien druckt sonst auf 2–3 Seiten."))
    k.append(K("F38", AW2, "L48", "druckbereich", "$A$1:$L$48", "F38: Druckbereich gesetzt wie bei Aufwuchs I."))
    # ---- Tabelle1 (versteckt): Tabellenende, SVERWEIS statt verschachteltem WENN, Auswahl leer
    for r in range(64, 84):
        k.append(K("F09", T1, f"D{r}", "zahl", 10, "F09: Restnutzungstabelle Bauklasse II/III nach 60 Jahren leer (→ 0 % Abzug); jetzt 10 % (letzter Tabellenwert D63). D64:D83. Bitte BV bestätigen." if r == 64 else ""))
    for r in range(44, 84):
        k.append(K("F09", T1, f"E{r}", "zahl", 10, "F09: Restnutzungstabelle Bauklasse IV nach 40 Jahren leer (→ 0 % Abzug); jetzt 10 % (letzter Tabellenwert E43). E44:E83. Bitte BV bestätigen." if r == 44 else ""))
    for b0, d_ref in ((88, 90), (93, 95), (98, 100), (103, 105)):
        k.append(K("F35", T1, f"D{b0}", "formel", f"IFERROR(VLOOKUP(D{d_ref},uwe,2),0)", "F35: bei leerer Typ-Auswahl kam #NV im ganzen Protokoll; jetzt Bauklasse 0 / NHK 0 / kein Altersabzug." if b0 == 88 else "", gruppe="IFERROR"))
        k.append(K("F35", T1, f"F{b0}", "formel", f"IFERROR(VLOOKUP(D{d_ref},uwe,5),5)", "", gruppe="IFERROR"))
        k.append(K("F35", T1, f"B{d_ref}", "formel", f"IFERROR(VLOOKUP(D{d_ref},uwe,3),0)", "", gruppe="IFERROR"))
    k.append(K("F41", T1, "D103", "formel", "IFERROR(VLOOKUP(D105,uwe,2),0)", "F41: war ein 22-fach verschachteltes WENN mit kaputtem Ende (=J19:J109); jetzt wie die anderen Bauteile."))
    k.append(K("F41", T1, "F5", "leer", None, "F41: war =-I10 (Rest ohne Funktion)."))
    for n in ("BKL.Beschreibung", "Jürgen", "rose", "ungeprüft", "uwe.2", "Gutachten", "Kinderspielhaus", "wasserlos"):
        k.append(K("F41", T1, "A1", "name_loeschen", n, "", gelb=False))
    k.append(K("F39", "Beseitigungsauflagen", "H22", "formel", "SUM(E22*G22)", "F39: war =SUM(E22*G33) (falsche Zeile). Die übrigen Fehler der versteckten Altblätter (#BEZUG!) bleiben; Blätter besser löschen."))
    return k


class XlsxPatcher:
    """Ändert Zellen, Stile, Kommentare und Auswahllisten direkt im XML einer .xlsx-Datei."""

    def __init__(self, quelle: Path):
        from lxml import etree
        self.etree = etree
        with zipfile.ZipFile(quelle) as z:
            self.reihenfolge = [i.filename for i in z.infolist()]
            self.teile: dict[str, bytes] = {i.filename: z.read(i.filename) for i in z.infolist()}
        self.roots: dict[str, Any] = {}
        self.neue_teile: dict[str, bytes] = {}
        self._styles_cache: dict[tuple, int] = {}
        wb = self.xml("xl/workbook.xml")
        rels = self.xml("xl/_rels/workbook.xml.rels")
        ziele = {r.get("Id"): r.get("Target") for r in rels}
        self.blattdatei: dict[str, str] = {}
        self.blattindex: dict[str, int] = {}
        for i, s in enumerate(wb.find(f"{M_}sheets")):
            ziel = ziele[s.get(f"{{{NS_R}}}id")]
            self.blattdatei[s.get("name")] = "xl/" + ziel.lstrip("/").removeprefix("xl/")
            self.blattindex[s.get("name")] = i

    # -- XML-Teile
    def xml(self, name: str):
        if name not in self.roots:
            self.roots[name] = self.etree.fromstring(self.teile[name])
        return self.roots[name]

    def speichern(self, ziel: Path) -> None:
        self.pruefe_gemeinsame_formeln()
        with zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as z:
            fertig = set()
            for name in self.reihenfolge:
                if name in self.entfernt:
                    continue
                z.writestr(name, self._bytes(name))
                fertig.add(name)
            for name in list(self.roots) + list(self.neue_teile):
                if name not in fertig and name not in self.entfernt:
                    z.writestr(name, self._bytes(name))
                    fertig.add(name)

    entfernt: set[str] = set()

    def _bytes(self, name: str) -> bytes:
        if name in self.roots:
            return self.etree.tostring(self.roots[name], xml_declaration=True, encoding="UTF-8", standalone=True)
        if name in self.neue_teile:
            return self.neue_teile[name]
        return self.teile[name]

    # -- Stile
    def _stile(self):
        return self.xml("xl/styles.xml")

    def stil_variante(self, s: int, gelb: bool = False, zahlenformat: int | None = None, gesperrt: bool | None = None) -> int:
        key = (s, gelb, zahlenformat, gesperrt)
        if key in self._styles_cache:
            return self._styles_cache[key]
        st = self._stile()
        xfs = st.find(f"{M_}cellXfs")
        basis = xfs[s]
        neu = copy.deepcopy(basis)
        if gelb:
            fills = st.find(f"{M_}fills")
            if not hasattr(self, "_gelb_id"):
                f = self.etree.SubElement(fills, f"{M_}fill")
                p = self.etree.SubElement(f, f"{M_}patternFill", patternType="solid")
                self.etree.SubElement(p, f"{M_}fgColor", rgb="FFFFFF00")
                self.etree.SubElement(p, f"{M_}bgColor", indexed="64")
                fills.set("count", str(len(fills)))
                self._gelb_id = len(fills) - 1
            neu.set("fillId", str(self._gelb_id))
            neu.set("applyFill", "1")
        if zahlenformat is not None:
            neu.set("numFmtId", str(zahlenformat))
            neu.set("applyNumberFormat", "1")
        if gesperrt is not None:
            for p in neu.findall(f"{M_}protection"):
                neu.remove(p)
            prot = self.etree.Element(f"{M_}protection", locked="1" if gesperrt else "0")
            ext = neu.find(f"{M_}extLst")
            if ext is not None:
                ext.addprevious(prot)
            else:
                neu.append(prot)
            neu.set("applyProtection", "1")
        xfs.append(neu)
        xfs.set("count", str(len(xfs)))
        self._styles_cache[key] = len(xfs) - 1
        return len(xfs) - 1

    # -- Zellen
    def _zeile(self, blatt: str, nr: int):
        sd = self.xml(self.blattdatei[blatt]).find(f"{M_}sheetData")
        vor = None
        for row in sd:
            r = int(row.get("r"))
            if r == nr:
                return row
            if r > nr:
                vor = row
                break
        neu = self.etree.Element(f"{M_}row", r=str(nr))
        if vor is not None:
            vor.addprevious(neu)
        else:
            sd.append(neu)
        return neu

    def zelle(self, blatt: str, ref: str):
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        spalte, nr = coordinate_from_string(ref)
        ci = column_index_from_string(spalte)
        row = self._zeile(blatt, nr)
        vor = None
        for c in row:
            cc = column_index_from_string(coordinate_from_string(c.get("r"))[0])
            if cc == ci:
                return c
            if cc > ci:
                vor = c
                break
        neu = self.etree.Element(f"{M_}c", r=ref)
        row.attrib.pop("spans", None)   # Bereichsangabe der Zeile stimmt nach dem Einfügen evtl. nicht mehr
        if vor is not None:
            vor.addprevious(neu)
        else:
            row.append(neu)
        return neu

    def _leeren(self, c) -> None:
        for kind in list(c):
            c.remove(kind)
        if "t" in c.attrib:
            del c.attrib["t"]

    def setze(self, blatt: str, ref: str, art: str, neu: Any = None) -> None:
        c = self.zelle(blatt, ref)
        self._leeren(c)
        if art == "formel":
            f = self.etree.SubElement(c, f"{M_}f")
            f.text = str(neu)
        elif art == "text":
            c.set("t", "inlineStr")
            is_ = self.etree.SubElement(c, f"{M_}is")
            t = self.etree.SubElement(is_, f"{M_}t")
            t.text = str(neu)
        elif art == "zahl":
            v = self.etree.SubElement(c, f"{M_}v")
            v.text = repr(neu) if isinstance(neu, float) else str(neu)
        elif art == "leer":
            pass
        else:
            raise ValueError(art)

    def stil(self, blatt: str, ref: str, **kw) -> None:
        c = self.zelle(blatt, ref)
        s = int(c.get("s", "0"))
        c.set("s", str(self.stil_variante(s, **kw)))

    def pruefe_gemeinsame_formeln(self) -> None:
        for blatt, datei in self.blattdatei.items():
            root = self.xml(datei) if datei in self.roots else None
            if root is None:
                continue
            master, kinder = set(), set()
            for f in root.iter(f"{M_}f"):
                if f.get("t") == "shared":
                    (master if f.get("ref") else kinder).add(f.get("si"))
            fehlt = kinder - master
            if fehlt:
                raise RuntimeError(f"{blatt}: gemeinsame Formel ohne Master (si={sorted(fehlt)})")

    # -- Kommentare
    def _rels(self, datei: str):
        pfad = posix_join_rels(datei)
        if pfad in self.roots or pfad in self.teile:
            return pfad, self.xml(pfad)
        root = self.etree.Element(f"{{{NS_PR}}}Relationships", nsmap={None: NS_PR})
        self.roots[pfad] = root
        return pfad, root

    def kommentar(self, blatt: str, ref: str, text: str) -> None:
        import posixpath
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        datei = self.blattdatei[blatt]
        rp, rels = self._rels(datei)
        base = posixpath.dirname(datei)

        def finde(typ: str):
            for r in rels:
                if r.get("Type").endswith("/" + typ):
                    return posixpath.normpath(posixpath.join(base, r.get("Target")))
            return None

        def neue_rel(typ: str, ziel_abs: str) -> str:
            ids = [int(re.sub(r"\D", "", r.get("Id")) or 0) for r in rels]
            rid = f"rId{max(ids + [0]) + 1}"
            ziel = posixpath.relpath(ziel_abs, base)
            self.etree.SubElement(rels, f"{{{NS_PR}}}Relationship", Id=rid,
                                  Type=f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{typ}", Target=ziel)
            return rid

        cpart = finde("comments")
        if cpart is None:
            n = 1
            while f"xl/comments{n}.xml" in self.teile or f"xl/comments{n}.xml" in self.neue_teile or f"xl/comments{n}.xml" in self.roots:
                n += 1
            cpart = f"xl/comments{n}.xml"
            root = self.etree.fromstring(f'<comments xmlns="{NS_M}"><authors/><commentList/></comments>'.encode())
            self.roots[cpart] = root
            neue_rel("comments", cpart)
            ct = self.xml("[Content_Types].xml")
            self.etree.SubElement(ct, f"{{{NS_CT}}}Override", PartName="/" + cpart,
                                  ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml")
            self.reihenfolge.append(cpart)
        croot = self.xml(cpart)
        autoren = croot.find(f"{M_}authors")
        namen = [a.text for a in autoren]
        if AUTOR not in namen:
            a = self.etree.SubElement(autoren, f"{M_}author")
            a.text = AUTOR
            namen.append(AUTOR)
        aid = namen.index(AUTOR)
        cl = croot.find(f"{M_}commentList")
        vorhanden = {c.get("ref"): c for c in cl}
        run_kopf = ('<r xmlns="%s"><rPr><b/><sz val="9"/><color indexed="81"/><rFont val="Segoe UI"/><family val="2"/></rPr>'
                    '<t xml:space="preserve">%s</t></r>')
        run_text = ('<r xmlns="%s"><rPr><sz val="9"/><color indexed="81"/><rFont val="Segoe UI"/><family val="2"/></rPr>'
                    '<t xml:space="preserve">%s</t></r>')
        kopf, _, rest = text.partition(":")
        if ref in vorhanden:
            t = vorhanden[ref].find(f"{M_}text")
            t.append(self.etree.fromstring((run_kopf % (NS_M, html.escape("\n\n" + kopf + ":"))).encode()))
            t.append(self.etree.fromstring((run_text % (NS_M, html.escape(rest))).encode()))
            return
        com = self.etree.Element(f"{M_}comment", ref=ref, authorId=str(aid), shapeId="0")
        t = self.etree.SubElement(com, f"{M_}text")
        t.append(self.etree.fromstring((run_kopf % (NS_M, html.escape(kopf + ":"))).encode()))
        t.append(self.etree.fromstring((run_text % (NS_M, html.escape(rest))).encode()))
        # in Reihenfolge Zeile/Spalte einsortieren
        def schluessel(r):
            sp, nr = coordinate_from_string(r)
            return nr, column_index_from_string(sp)
        einsortiert = False
        for c in cl:
            if schluessel(c.get("ref")) > schluessel(ref):
                c.addprevious(com)
                einsortiert = True
                break
        if not einsortiert:
            cl.append(com)
        # VML-Notizfeld
        vpart = finde("vmlDrawing")
        sp, nr = coordinate_from_string(ref)
        col0, row0 = column_index_from_string(sp) - 1, nr - 1
        if vpart is None:
            n = 1
            while f"xl/drawings/vmlDrawing{n}.vml" in self.teile or f"xl/drawings/vmlDrawing{n}.vml" in self.neue_teile:
                n += 1
            vpart = f"xl/drawings/vmlDrawing{n}.vml"
            idmap = 40 + n
            self.neue_teile[vpart] = (
                '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
                'xmlns:x="urn:schemas-microsoft-com:office:excel">\n'
                f' <o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="{idmap}"/></o:shapelayout>\n'
                ' <v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe">\n'
                '  <v:stroke joinstyle="miter"/><v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>\n</xml>').encode("utf8")
            self.reihenfolge.append(vpart)
            rid = neue_rel("vmlDrawing", vpart)
            sheet = self.xml(datei)
            ld = self.etree.Element(f"{M_}legacyDrawing", {f"{{{NS_R}}}id": rid})
            self._einfuegen_nach_reihenfolge(sheet, ld, "legacyDrawing")
        daten = (self.neue_teile.get(vpart) or self.teile[vpart]).decode("utf8")
        ids = [int(x) for x in re.findall(r"_x0000_s(\d+)", daten)]
        sid = max(ids + [int(re.search(r'idmap[^>]*data="(\d+)"', daten).group(1)) * 1024 if re.search(r'idmap[^>]*data="(\d+)"', daten) else 1024]) + 1
        shape = (
            f'<v:shape id="_x0000_s{sid}" type="#_x0000_t202" style="position:absolute;margin-left:80pt;margin-top:20pt;'
            f'width:300pt;height:110pt;z-index:{sid % 1000};visibility:hidden" fillcolor="#ffffe1" o:insetmode="auto">'
            '<v:fill color2="#ffffe1"/><v:shadow on="t" color="black" obscured="t"/><v:path o:connecttype="none"/>'
            '<v:textbox style="mso-direction-alt:auto"><div style="text-align:left"></div></v:textbox>'
            '<x:ClientData ObjectType="Note"><x:MoveWithCells/><x:SizeWithCells/>'
            f'<x:Anchor>{col0 + 1}, 15, {max(row0 - 1, 0)}, 10, {col0 + 5}, 15, {row0 + 6}, 4</x:Anchor>'
            f'<x:AutoFill>False</x:AutoFill><x:Row>{row0}</x:Row><x:Column>{col0}</x:Column></x:ClientData></v:shape>\n')
        ende = daten.rfind("</xml>")
        daten = daten[:ende] + shape + daten[ende:]
        if vpart in self.teile:
            self.teile[vpart] = daten.encode("utf8")
        else:
            self.neue_teile[vpart] = daten.encode("utf8")

    def _einfuegen_nach_reihenfolge(self, root, element, name: str) -> None:
        pos = BLATT_REIHENFOLGE.index(name)
        letztes = None
        for kind in root:
            tag = self.etree.QName(kind).localname
            if tag in BLATT_REIHENFOLGE and BLATT_REIHENFOLGE.index(tag) < pos:
                letztes = kind
        if letztes is None:
            root.insert(0, element)
        else:
            letztes.addnext(element)

    # -- Auswahllisten / Eingabeprüfung
    def eingabepruefung(self, blatt: str, bereich: str, art: str, von: float, bis: float, hinweis: str) -> None:
        root = self.xml(self.blattdatei[blatt])
        dvs = root.find(f"{M_}dataValidations")
        if dvs is None:
            dvs = self.etree.Element(f"{M_}dataValidations", count="0")
            self._einfuegen_nach_reihenfolge(root, dvs, "dataValidations")
        # vorhandene Liste für diesen Bereich ersetzen
        for dv in list(dvs):
            if set(dv.get("sqref", "").split()) & set(bereich.split()):
                rest = set(dv.get("sqref").split()) - set(bereich.split())
                if rest:
                    dv.set("sqref", " ".join(sorted(rest)))
                else:
                    dvs.remove(dv)
        dv = self.etree.SubElement(dvs, f"{M_}dataValidation", type=art, operator="between", allowBlank="1",
                                   showInputMessage="1", showErrorMessage="1", errorTitle="Eingabeprüfung",
                                   error=hinweis, sqref=bereich)
        f1 = self.etree.SubElement(dv, f"{M_}formula1")
        f1.text = str(von)
        f2 = self.etree.SubElement(dv, f"{M_}formula2")
        f2.text = str(bis)
        dvs.set("count", str(len(dvs)))

    def skalierung(self, blatt: str, wert: int) -> None:
        ps = self.xml(self.blattdatei[blatt]).find(f"{M_}pageSetup")
        ps.set("scale", str(wert))

    def druckbereich(self, blatt: str, bereich: str) -> None:
        wb = self.xml("xl/workbook.xml")
        dn = wb.find(f"{M_}definedNames")
        if dn is None:
            dn = self.etree.Element(f"{M_}definedNames")
            wb.find(f"{M_}sheets").addnext(dn)
        el = self.etree.SubElement(dn, f"{M_}definedName", name="_xlnm.Print_Area", localSheetId=str(self.blattindex[blatt]))
        name = blatt if re.fullmatch(r"\w+", blatt) else f"'{blatt}'"
        el.text = f"{name}!{bereich}"

    def name_loeschen(self, name: str) -> None:
        dn = self.xml("xl/workbook.xml").find(f"{M_}definedNames")
        for el in list(dn):
            if el.get("name") == name:
                dn.remove(el)

    def neu_berechnen(self) -> None:
        """Excel rechnet beim Öffnen alles neu; die Berechnungskette (calcChain) wird verworfen und von Excel neu aufgebaut."""
        wb = self.xml("xl/workbook.xml")
        cp = wb.find(f"{M_}calcPr")
        if cp is None:
            cp = self.etree.SubElement(wb, f"{M_}calcPr")
        cp.set("fullCalcOnLoad", "1")
        # Zwischengespeicherte Ergebnisse der Formeln entfernen, damit auch Programme ohne Neuberechnung
        # (Vorschau, Handy) keine veralteten Zahlen zeigen.
        for datei in self.blattdatei.values():
            for c in self.xml(datei).iter(f"{M_}c"):
                if c.find(f"{M_}f") is not None:
                    for v in c.findall(f"{M_}v"):
                        c.remove(v)
                    if c.get("t") in ("str", "e", "b", "n"):
                        del c.attrib["t"]
        if "xl/calcChain.xml" in self.teile:
            self.entfernt = self.entfernt | {"xl/calcChain.xml"}
            rels = self.xml("xl/_rels/workbook.xml.rels")
            for r in list(rels):
                if r.get("Type").endswith("/calcChain"):
                    rels.remove(r)
            ct = self.xml("[Content_Types].xml")
            for o in list(ct):
                if o.get("PartName") == "/xl/calcChain.xml":
                    ct.remove(o)


def posix_join_rels(datei: str) -> str:
    import posixpath
    return posixpath.join(posixpath.dirname(datei), "_rels", posixpath.basename(datei) + ".rels")


def schreibe_korrekturvorschlag(quelle: Path, ziel: Path) -> list[tuple[Korrektur, str]]:
    """Schreibt die korrigierte Kopie. Gibt (Korrektur, alter Inhalt) für den Bericht zurück."""
    if ziel.resolve() == quelle.resolve():
        raise ValueError("Ziel darf nicht das Original sein")
    m = Mappe(quelle)
    p = XlsxPatcher(quelle)
    protokoll: list[tuple[Korrektur, str]] = []
    for k in _korrekturen():
        zellen = _zellen(k.zelle)
        alt = ""
        if k.art in ("formel", "text", "zahl", "leer", "format", "entsperren", "sperren"):
            wert = m.blatt(k.blatt)[k.zelle].value
            alt = "" if wert is None else str(wert)
        protokoll.append((k, alt))
        if k.art in ("formel", "text", "zahl", "leer"):
            for z in zellen:
                p.setze(k.blatt, z, k.art, k.neu)
                p.stil(k.blatt, z, gelb=k.gelb, zahlenformat=k.zahlenformat)
        elif k.art == "format":
            for z in zellen:
                p.stil(k.blatt, z, gelb=k.gelb, zahlenformat=k.zahlenformat)
        elif k.art == "entsperren":
            p.stil(k.blatt, k.zelle, gesperrt=False)
        elif k.art == "sperren":
            p.stil(k.blatt, k.zelle, gesperrt=True)
        elif k.art in ("dv_zahl", "dv_ganz"):
            von, bis = k.neu
            hinweis = {"dv_zahl": f"Bitte eine Zahl zwischen {von} und {bis} eingeben.", "dv_ganz": f"Bitte eine ganze Zahl zwischen {von} und {bis} eingeben."}[k.art]
            p.eingabepruefung(k.blatt, k.zelle, "decimal" if k.art == "dv_zahl" else "whole", von, bis, hinweis)
            if k.gelb:
                p.stil(k.blatt, k.zelle, gelb=True)
        elif k.art == "seite":
            p.skalierung(k.blatt, k.neu)
            if k.gelb:
                p.stil(k.blatt, k.zelle, gelb=True)
        elif k.art == "druckbereich":
            p.druckbereich(k.blatt, k.neu)
        elif k.art == "name_loeschen":
            p.name_loeschen(k.neu)
        else:
            raise ValueError(k.art)
        if k.kommentar and k.art != "name_loeschen":
            p.kommentar(k.blatt, zellen[0], k.kommentar)
    p.neu_berechnen()
    ziel.parent.mkdir(parents=True, exist_ok=True)
    p.speichern(ziel)
    return protokoll


# ---------------------------------------------------------------------------------------------------
# 8. Fehlerbericht (Markdown und PDF) – aus dem Katalog und den Ergebnissen der Prüfung erzeugt
# ---------------------------------------------------------------------------------------------------

SELBST = {
    "F02": "Blatt „Baulichkeiten“, Zelle K81 anklicken: In der Bearbeitungsleiste steht =K52+K61+K71+K75+K77+K79. K61 ist der Übertrag von K53 (=SUMME(K14:K52)), und diese Summe enthält K52 (Schuppen) schon. Einen Schuppen eintragen: K81 ist um dessen Wert höher als K53 + K71 + K75 + K77 + K79.",
    "F05": "Blatt „Baulichkeiten“, Zelle G68 anklicken: Dort steht die Zahl 100, keine Formel. Unter „2.3 b) angebaut“ Güte „hochwertig“, Altersminderung 3 %, Alter 8 Jahre, Fläche 6 m² eintragen: K71 bleibt 0,00 €.",
    "F03": "Blatt „Baulichkeiten“, „2.3 a) integriert“: 275 € bei NHK und 6 m² bei F71 eintragen (Güte bei b) auf „ohne“ lassen): K71 bleibt 0,00 €. Zelle C64 anklicken: Sie ist leer.",
    "F12": "Leere Vorlage öffnen, Blatt „Deckbl UPV“: E25 zeigt 250,00 €, E26 150,00 €, E28 −400,00 €.",
    "F01": "Blatt „Baulichkeiten“, Zelle C33 (Anbau A2) anklicken: =$H$23. Bei A1 steht in C24 richtig =$H$23, bei A2 muss in C33 der Bezug auf H32 stehen (die eigene Zeile „Mängelabzug“).",
}

FRAGEN = {
    "F05": "Soll der angebaute Laubenvorplatz (b) wie das Gewächshaus mit „Altersminderung pro Jahr × Alter“ gerechnet werden (Vorschlag G68 =MIN(100;E66*H66))?",
    "F08": "Woher kommt die Grenze „Kosten > 74 €“ im Kanalanschluss (K79), und soll ein Anschluss mit Alter 0 den vollen Erstellungspreis bringen?",
    "F09": "Was soll gelten, wenn die Restnutzungsjahre über das Ende der Tabelle gehen (Bauklasse IV über 40, II/III über 60 Jahre)? Vorschlag: 10 % wie der letzte Tabellenwert.",
    "F14": "Warum bietet die Liste für den integrierten Laubenvorplatz 225 € an? Die Richtlinie kennt 275 / 175 / 80 €.",
    "F18": "Sollen die fehlenden Mängel der Richtlinie (Schornstein, Moose und Flechten, Risse bei Decken, Belag verschlissen, Fliesen) in die Mängelliste, und ist „unzulässige Höhe der Laube“ gewollt?",
    "F20": "Soll das Blatt Baulichkeiten Felder für die Begründung der Restnutzungsdauer und „Baujahr nachgewiesen/geschätzt“ bekommen (Richtlinie S. 9)?",
    "F21": "Wie werden Kellerräume eingetragen (als Anbau A1/A2, oder soll es einen eigenen Block geben)?",
    "F25": "Soll das Gewächshaus mit NHK 160 € vorbelegt werden und die Grenzen (12 m², 2,20 m First) im Blatt stehen? Woher kommen die 10 % Mängelabzug?",
    "F26": "Gilt für Schuppen der Mindestabzug von 3 % je Jahr (Richtlinie S. 11) zusätzlich zur Restnutzungstabelle?",
    "F27": "Sollen Außenanlagen und Aufwuchs Spalten für Alter/Zustand bekommen (Mindestabzug 3 % je Jahr, Richtlinie S. 12)?",
    "F28": "Welcher Index gilt für Wertermittlungen 2026 – der Wert des IV. Quartals 2025? In H7 steht 179,2 „für 2025“.",
    "F29": "Wird der Altersabzug von Abwassergrube und Kanal linear (Alter × 3 %, 16 Jahre = 48 %) oder vom Restwert gerechnet (16 Jahre = 38,57 %, so rechnet die App)?",
    "F30": "Gilt die Absenkung bei nicht geleerter Grube (150 €, S. 14) zusätzlich zum Einbehalt Fäkalienabfuhr (150 €, S. 20), oder ist einer von beiden gemeint?",
    "F31": "Auf welche Regel stützt sich „Beetflächen auf 10 % Parzellenfläche erhöhen“?",
    "F46": "An wen geht der Widerspruch – an die Wertermittler (Deckblatt A34) oder an den Bezirksverband (A38; Richtlinie S. 23)?",
    "F32": "Soll der Blattschutz (ohne Passwort) wieder eingeschaltet werden? Dann müssen vorher F33 und F34 behoben sein.",
}


def _bild(x: Any, einheit: str = "€") -> str:
    """Zahl für den Bericht: Euro mit zwei Stellen (drei, wenn ein halber Cent vorkommt), sonst mit Einheit."""
    if not isinstance(x, (int, float)) or isinstance(x, bool):
        return str(x)
    if einheit == "€":
        s = f"{x:,.3f}" if abs(x * 100 - round(x * 100)) > 1e-6 else f"{x:,.2f}"
    else:
        s = f"{x:,.2f}".rstrip("0").rstrip(".")
    s = s.replace(",", "X").replace(".", ",").replace("X", ".").replace("-", "−")
    return f"{s} {einheit}"


def vergleiche(ctx: Ctx) -> list[dict[str, Any]]:
    """Zeilen für den Anhang „Testfälle“: Vorlage gegen Handrechnung."""
    if ctx.lauf is None:
        return []
    r, s = ctx.r, ctx.soll
    zeilen: list[dict[str, Any]] = []

    def z(fall, blatt, zelle, was, hand, tol=0.0051, einheit="€"):
        v = r(fall, blatt, zelle)
        zeilen.append({"fall": fall, "zelle": f"{blatt.strip()}!{zelle}", "was": was, "vorlage": v, "hand": hand,
                       "gleich": fast_gleich(v, hand, tol), "einheit": einheit})

    z("T0", DB, "E28", "Gesamtbetrag mit Abzügen (leere Vorlage)", 0)
    for k, name in (("L", "Laube"), ("A1", "Anbau A1"), ("A2", "Anbau A2"), ("S", "Schuppen")):
        z("T1", BL, BAUTEIL_ZELLEN[k][4] if k != "S" else "E52", f"Wert {name}", s["T1"]["teile"][k]["betrag"])
    z("T1", BL, "K44", "Gewächshaus 10 m²", s["T1"]["gwh"])
    wie = "linear, 30 %" if s["modus"] == "linear" else "vom Restwert, 26,26 %"
    z("T1", BL, "K75", f"Abwassergrube 6 m³, 10 Jahre ({wie})", s["T1"]["grube"])
    z("T1", BL, "K79", f"Kanalanschluss 2.000 €, 10 Jahre ({wie})", s["T1"]["kanal"])
    z("T1", BL, "K53", "Zwischensumme Baulichkeiten", s["T1"]["K53"])
    z("T1", BL, "K81", "Baulichkeiten insgesamt", s["T1"]["K81"])
    z("T1", BL, "F101", "Bewertete bebaute Fläche", s["T1"]["F101"], einheit="m²")
    z("T2", AA, "L46", "Außenanlagen insgesamt (31 Zeilen)", s["T2"]["AA"])
    z("T2", AW1, "L46", "Aufwuchs I (28 Zeilen)", s["T2"]["AW1"])
    z("T2", AW2, "L46", "Aufwuchs II einschließlich Übertrag (36 + 28 Zeilen)", s["T2"]["AW2"])
    z("T2", BE, "K24", "Beseitigung Bau, Summe (12 Zeilen)", s["T2"]["BE1"])
    z("T2", BE, "K45", "Beseitigung Außenanlagen, Summe (11 Zeilen)", s["T2"]["BE2"], 0.0001)
    z("T2", DB, "E28", "Gesamtbetrag mit Abzügen", s["T2"]["E28"])
    z("T3", BL, "K71", "Laubenvorplatz a) integriert", s["T3"]["K71"])
    z("T4", BL, "K71", "Laubenvorplatz b) angebaut", s["T4"]["K71"])
    z("T6", BL, "K79", "Kanal Alter 0, 2.000 €", kanal_betrag(0, 2000, s["modus"]))
    z("T7", BL, "K79", "Kanal Alter 10, 70 €", kanal_betrag(10, 70, s["modus"]))
    z("T8", DB, "E23", "Aufwuchs (Aufwuchs II enthält einen Abzug)", r("T8", AW2, "L46"))
    z("T8", BE, "K34", "Beseitigung 3 × 10,05 € × 150 %", 45.23, 0.0001)
    z("T8", AA, "J3", "versiegelbare Fläche bei 20 m² Parzelle", 0.0, einheit="m²")
    z("T9", BL, "E33", "Altersabzug bei 7,5 Restnutzungsjahren (Bauklasse IV; Tabelle: 7 Jahre)", s["T9"]["E33_7_5"], einheit="%")
    return zeilen


def baue_bericht(m: Mappe, ctx: Ctx, befunde: dict[str, Befund]) -> list[tuple]:
    """Der Bericht als Liste von Blöcken: ("h1"|"h2"|"h3"|"p"|"ul"|"tabelle"|"hinweis"|"neueSeite", …)."""
    import hashlib
    sha = hashlib.sha256(m.pfad.read_bytes()).hexdigest()
    geaendert = ""
    try:
        geaendert = re.search(r"<dcterms:modified[^>]*>([^<]+)<", m.zip.read("docProps/core.xml").decode("utf8")).group(1)[:10]
    except Exception:
        pass
    b: list[tuple] = []
    anz = {a: [f for f in FUNDE if f.auswirkung == a] for a in AUSWIRKUNGEN}
    gew = {g: [f for f in FUNDE if f.gewissheit == g] for g in GEWISSHEITEN}
    behoben = [f.nr for f in FUNDE if befunde[f.nr].vorhanden is False]

    b.append(("h1", "Fehlerbericht zur WEP-Vorlage des BV Pankow"))
    b.append(("meta", [
        ("Geprüfte Datei", f"`docs/WEP_Vorlage.xlsx` (Stand der Datei: {geaendert or 'unbekannt'}, SHA-256 `{sha[:16]}…`)"),
        ("Fachliche Grundlage", "Wertermittlungsrichtlinie für Kleingärten (BewRL), Landesverband Berlin der Gartenfreunde, gültig ab 01.01.2023 (`docs/Richtlinie.pdf`)"),
        ("Geprüft am", dt.date.today().strftime("%d.%m.%Y")),
        ("Prüfskript", f"`tools/vorlage-pruefen.py` (Version {__version__}) – prüft auch jede neue Fassung der Vorlage erneut"),
        ("Korrekturvorschlag", "`docs/WEP_Vorlage_Korrekturvorschlag.xlsx` (geänderte Zellen gelb, mit Kommentar; das Original bleibt unverändert)"),
    ]))

    # ---- 1. Zusammenfassung
    b.append(("h2", "1. Zusammenfassung für den BV"))
    b.append(("p", "Wir haben die Excel-Vorlage Zelle für Zelle durchgesehen, alle Formeln gelesen und mit erfundenen Beispielen "
                   "(ohne echte Personendaten) durchgerechnet. Jede Zahl haben wir mit einer eigenen Handrechnung verglichen und die "
                   "Texte und Werte gegen die Richtlinie gehalten. Alles, was wir gefunden haben, steht in der Tabelle in Abschnitt 2 – "
                   "mit der Zelle, was dort steht, was daran falsch ist und wie es richtig heißen muss."))
    b.append(("p", f"**Ergebnis: {len(FUNDE)} Funde.** "
                   f"{len(anz['falscher Betrag'])} davon ändern einen Betrag im Protokoll, "
                   f"{len(anz['falsche oder fehlende Angabe'])} betreffen falsche oder fehlende Angaben, "
                   f"{len(anz['Bedienung'])} die Bedienung und {len(anz['kosmetisch'])} sind Schönheitsfehler. "
                   f"{len(gew['sicher'])} Funde sind sicher (nachgerechnet oder an der Datei ablesbar), "
                   f"{len(gew['wahrscheinlich'])} sind sehr wahrscheinlich, bei {len(gew['bitte BV fragen'])} brauchen wir Ihre Entscheidung "
                   "(Abschnitt 3: „Bitte BV fragen“) – das sind keine bestätigten Fehler, sondern Punkte, bei denen die Richtlinie mehrdeutig ist oder der BV festlegen muss, was gelten soll."))
    tab = [[a] + [str(sum(1 for f in anz[a] if f.gewissheit == g)) for g in GEWISSHEITEN] + [str(len(anz[a]))] for a in AUSWIRKUNGEN]
    tab.append(["**zusammen**"] + [f"**{len(gew[g])}**" for g in GEWISSHEITEN] + [f"**{len(FUNDE)}**"])
    b.append(("tabelle", {"kopf": ["Auswirkung"] + GEWISSHEITEN + ["zusammen"], "zeilen": tab, "klasse": "klein"}))
    b.append(("h3", "Die fünf wichtigsten Funde"))
    top = sorted((f for f in FUNDE if f.top), key=lambda f: f.top)
    wirkung = {}
    if ctx.lauf:
        s = ctx.soll
        wirkung = {
            "F02": f"Im Testfall T1 (Laube, zwei Anbauten, Schuppen …) zeigt die Vorlage {eur(ctx.r('T1', BL, 'K81'))} statt {eur(s['T1']['K81'])}; "
                   f"davon {eur(ctx.r('T1', BL, 'K52'))} allein durch den doppelt gezählten Schuppen.",
            "F05": f"Im Testfall T4 (angebauter Vorplatz, 6 m²) bleibt der Wert bei {eur(ctx.r('T4', BL, 'K71'))} statt {eur(s['T4']['K71'])}.",
            "F03": f"Im Testfall T3 (integrierter Vorplatz, 6 m²) bleibt der Wert bei {eur(ctx.r('T3', BL, 'K71'))} statt {eur(s['T3']['K71'])}.",
            "F12": f"Die leere Vorlage zeigt als Gesamtbetrag {eur(ctx.r('T0', DB, 'E28'))}.",
            "F01": f"Im Testfall T1 rechnet Anbau A2 mit {eur(ctx.r('T1', BL, 'E35'))} statt {eur(s['T1']['teile']['A2']['betrag'])}.",
        }
    for i, f in enumerate(top, 1):
        ort = f.zellen if "," in f.blatt else f"{f.blatt}, {f.zellen}"
        b.append(("p", f"**{i}. {f.nr} – {f.titel}** ({ort})"))
        punkte = [x for x in (f"*Wirkung:* {wirkung[f.nr]}" if f.nr in wirkung else "",
                              f"*So sehen Sie es selbst:* {SELBST[f.nr]}" if f.nr in SELBST else "",
                              f"*Korrektur:* {f.vorschlag}") if x]
        b.append(("ul", punkte))
    b.append(("h3", "Was jetzt zu tun ist"))
    b.append(("ul", [
        "**Formelfehler beheben** (F01–F05, F10, F11): Die Datei `WEP_Vorlage_Korrekturvorschlag.xlsx` enthält die Änderungen, jede gelb markiert und mit Kommentar. Sie kann als Vorlage dienen oder zeigt, was in der eigenen Vorlage zu ändern ist.",
        "**Rückfragen beantworten** (Abschnitt 3): Dort stehen die Punkte, bei denen die Richtlinie mehrdeutig ist oder die Vorlage etwas anderes tut als die App.",
        "**Leere Vorlage säubern** (F12): Beispielwerte löschen, damit ein neues Protokoll nicht mit Abzügen beginnt.",
        "**Neue Fassung prüfen:** `python tools/vorlage-pruefen.py NEUE_DATEI.xlsx` zeigt je Fund „VORHANDEN“ oder „behoben“ (Anhang D).",
    ]))
    b.append(("h3", "Was in Ordnung ist"))
    b.append(("ul", [
        "Die versteckte Tabelle „Tabelle1“ stimmt mit der Richtlinie überein: Restnutzungstabelle aller drei Bauklassen (Seite 10), NHK 2015 je Bauklasse (Seite 5) und Abrisskosten je m³ (Seite 19). Jede Zahl wurde einzeln verglichen.",
        "Die Rechenkette Laube / Anbau A1 / Schuppen (Mängelabzug → Altersabzug → Baupreisindex → Betrag) rechnet bei richtigen Eingaben genau wie die Handrechnung.",
        "Die Summen der Blätter Außenanlagen, Aufwuchs I und Aufwuchs II sowie der Übertrag von Aufwuchs I nach II sind richtig, auch mit allen Zeilen voll.",
        "Die Verweise vom Deckblatt auf die Anlagen 1 bis 4 zeigen auf die richtigen Zellen; im sichtbaren Teil der Vorlage gibt es keinen Fehlerwert wie #BEZUG!.",
        "Feste Zahlen in Formeln: 2.200 € (Kanal), 24 m² und 6 % (versiegelte Fläche) passen zur Richtlinie; Auffälligkeiten stehen in F08 und F29.",
    ]))

    # ---- 2. Funde
    b.append(("neueSeite", "wide"))
    b.append(("h2", "2. Alle Funde"))
    b.append(("p", "Die Nummern gelten auch im Prüfskript und im Korrekturvorschlag (dort steht die Nummer am Anfang jedes Kommentars). "
                   "**Gewissheit:** *sicher* = an der Datei ablesbar oder nachgerechnet; *wahrscheinlich* = sehr wahrscheinlich ein Fehler, "
                   "bitte bestätigen; *bitte BV fragen* = kein Fehler, solange der BV nicht entschieden hat (siehe Abschnitt 3). "
                   "**Formeln** stehen in der Schreibweise des deutschen Excel (Funktionsnamen deutsch, Semikolon als Trennzeichen); "
                   "die Datei selbst speichert sie englisch."))
    for a in AUSWIRKUNGEN:
        fl = [f for f in FUNDE if f.auswirkung == a]
        b.append(("h3", f"{a} ({len(fl)})"))
        zeilen = []
        for f in fl:
            bel = f.beleg
            dyn = befunde[f.nr].beleg
            zeilen.append([f"**{f.nr}**", f.blatt, f.zellen, f.ist, f"**{f.titel}.** {f.problem}",
                           f"{bel}<br>*Prüfung:* {dyn}" if dyn else bel, f.vorschlag, f.gewissheit])
        b.append(("tabelle", {"kopf": ["Nr.", "Blatt", "Zelle(n)", "Was steht da", "Was ist falsch", "Beleg (Richtlinie / Testfall)",
                                       "Korrekturvorschlag", "Gewissheit"], "zeilen": zeilen, "klasse": "funde",
                                 "breiten": [3.5, 8, 9, 14, 26, 19, 14, 6.5]}))

    # ---- 3. Fragen
    b.append(("neueSeite", "normal"))
    b.append(("h2", "3. Bitte BV fragen"))
    b.append(("p", "Hier braucht es eine Entscheidung des BV: Die Richtlinie ist nicht eindeutig, die Vorlage tut etwas anderes als die App, "
                   "oder wir wissen nicht, was gewollt ist. Solange nicht entschieden ist, behandeln wir diese Punkte nicht als bestätigte Fehler."))
    fr = [[f"**{f.nr}**", f.titel, FRAGEN.get(f.nr, "")] for f in FUNDE if f.nr in FRAGEN]
    b.append(("tabelle", {"kopf": ["Nr.", "Thema", "Frage"], "zeilen": fr, "klasse": "klein"}))

    # ---- Anhang A
    b.append(("h2", "Anhang A – Testfälle (alle Zahlen erfunden)"))
    b.append(("p", "Die Eingaben wurden in eine Kopie der Vorlage geschrieben, mit LibreOffice neu berechnet und mit einer unabhängigen "
                   "Handrechnung (Python, Excel-Rundung) verglichen. Das Original wurde nicht verändert. „Vorlage“ ist das Ergebnis der "
                   "Vorlage, „Handrechnung“ das richtige Ergebnis."))
    if ctx.lauf:
        szen = [[f"**{s.name}**", s.beschreibung] for s in ctx.lauf.szenarien.values()]
        b.append(("tabelle", {"kopf": ["Fall", "Eingabe"], "zeilen": szen, "klasse": "klein"}))
        vz = [[v["fall"], v["zelle"], v["was"], _bild(v["vorlage"], v["einheit"]), _bild(v["hand"], v["einheit"]),
               "✔ gleich" if v["gleich"] else "✘ weicht ab"] for v in vergleiche(ctx)]
        b.append(("tabelle", {"kopf": ["Fall", "Zelle", "Was", "Vorlage", "Handrechnung", "Ergebnis"], "zeilen": vz, "klasse": "klein"}))
        ohne_fehler = [s.name for s in ctx.lauf.szenarien.values() if not s.fehlerzellen]
        mit_fehler = [f"{s.name} ({len(s.fehlerzellen)})" for s in ctx.lauf.szenarien.values() if s.fehlerzellen]
        b.append(("p", f"**Fehlerwerte nach der Neuberechnung** (alle Zellen der sichtbaren Blätter, #NV, #WERT!, #BEZUG!, #DIV/0! …): keine in {', '.join(ohne_fehler)}"
                       + (f"; Fehlerwerte in {', '.join(mit_fehler)} (Randfälle F35)." if mit_fehler else ".")))
        s5 = ctx.lauf["T5"]
        b.append(("p", "**Randfälle ohne Handrechnung** (Beobachtung): "
                       f"T5 Bauteil-Auswahl leer → K53 = {s5[(BL, 'K53')]}, K81 = {s5[(BL, 'K81')]}. "
                       f"Restnutzungsjahre 61 (Bauklasse II) → Altersabzug {s5[(BL, 'E24')]} %; 41 (Bauklasse IV) → {s5[(BL, 'E33')]} %. "
                       f"Mängel 120 % → Betrag {_bild(s5[(BL, 'E52')])}. Grube und Kanal 40 Jahre → {_bild(s5[(BL, 'K75')])} und {_bild(s5[(BL, 'K79')])}. "
                       f"T9 Restnutzungsjahre „zwanzig“ → E15 = {ctx.lauf['T9'][(BL, 'E15')]}, −3 → E24 = {ctx.lauf['T9'][(BL, 'E24')]}."))
    else:
        b.append(("p", "Ohne Rechentests erzeugt (LibreOffice nicht verfügbar)."))

    # ---- Anhang B
    b.append(("h2", "Anhang B – Bestandsaufnahme"))
    bz = [[z["blatt"].strip() + ("" if z["sichtbar"] else " (versteckt)"), str(z["formeln"]), str(z["verbunden"]), str(z["listen"]),
           z["druckbereich"].split("!")[-1] if z["druckbereich"] != "–" else "–", z["schutz"]] for z in bestandsaufnahme(m)]
    b.append(("tabelle", {"kopf": ["Blatt", "Formeln", "verbundene Zellen", "Auswahllisten", "Druckbereich", "Blattschutz"], "zeilen": bz, "klasse": "klein"}))
    b.append(("p", "**Blattschutz:** „aktiv“ nur bei Tabelle1 (mit Passwort). Bei den anderen Blättern ist ein Schutz-Element ohne den Schalter vorhanden "
                   "(„aus“) oder gar keines („nicht vorhanden“). **Bedingte Formatierungen:** keine. **Definierte Namen:** `uwe` (Typen-Tabelle), "
                   "`nelke` (Restnutzungstabelle), `typen` (Auswahlliste) werden benutzt; `BKL.Beschreibung`, `Jürgen`, `rose`, `ungeprüft`, `uwe.2` und drei "
                   "blattbezogene Namen nicht. **Auswahllisten** in Baulichkeiten: E63, K77, B67:C67, E76, G12:J12, G21/G30/G47/J47, E66/G40, D76 (Datum)."))
    zyk = zirkelbezuege(m)
    b.append(("p", f"**Zirkelbezüge:** {'keine gefunden (Abhängigkeitsgraph aller Formeln)' if not zyk else str(len(zyk)) + ' gefunden: ' + '; '.join(zyk[:5])}. "
                   f"**Fehlerwerte** in sichtbaren Blättern (in der Datei gespeicherte Werte): {len(formel_zellen_mit_fehlerwerten(m))}. "
                   "**Versteckte Blätter** mit #BEZUG!: siehe F39."))

    # ---- Anhang C
    b.append(("h2", "Anhang C – Geprüft und ohne Befund"))
    abw = pruefe_tabellen(m)
    b.append(("ul", [
        f"Tabelle1 gegen Richtlinie: Restnutzungstabelle (3 Bauklassen, bis 80/60/40 Jahre), NHK-Stufen, Abrisskosten – {'alle Zahlen gleich' if not abw else str(len(abw)) + ' Abweichungen'}.",
        "Zeilenformeln Außenanlagen L15:L45, Aufwuchs I L18:L45, Aufwuchs II L10:L45: überall `=RUNDEN(A*K;2)`, Summen mit richtigem Bereich (bei allen Zeilen voll nachgerechnet, T2).",
        "Rechenkette Laube / A1 / Schuppen: Formeln der vier Bauteil-Blöcke bis auf C33 und J52 gleich aufgebaut (Mustervergleich).",
        "Deckblatt: E21 = K81, E22 = L46 (Außenanlagen), E24 = Summe, E25 = H104, E26 = K24, E27 = K45; Gesamtbetrag E28 = E24 − E25 − E26 − E27 (T2).",
        "Gewächshaus (E42 = Alter × Minderung je Jahr, „abgeschrieben“ ab 100 %), Laubenvorplatz-Güte (b: 175 / 80 € wie Richtlinie S. 11), Grundwert 470 €, Höchstwert Kanal 2.200 €.",
        "Mängeltabelle: alle „Abzug maximal“-Werte der Richtlinie stimmen (Dach 20 %, Decken 10 %, Wände 13 %, Fußboden 10 %, Fundament 20 %, Ausbau außen 5 %, innen 7 %, Bauelemente 5 %, Schädlinge 10 %, Sonderabfall 10 %, Laube nicht beräumt 10 %).",
        "Rundung: Die Vorlage rundet bei jedem Zwischenschritt (€/m² nach Mängelabzug, nach Altersabzug, nach Index) auf Cent; gegenüber einer ungerundeten Rechnung sind das wenige Cent je m², bei 24 m² höchstens etwa 0,50 €. Das ist kein Fehler, muss aber in der App genauso gerundet werden (der Export lässt die Vorlage rechnen).",
        "Formeln mit festen Zahlen: " + "; ".join(konstanten_in_formeln(m)),
    ]))

    # ---- Anhang D
    b.append(("h2", "Anhang D – Prüfskript und Korrekturvorschlag"))
    b.append(("p", "**Neue Fassung der Vorlage prüfen:**"))
    b.append(("code", "python tools/vorlage-pruefen.py NEUE_VORLAGE.xlsx          # Zustand je Fund: VORHANDEN / behoben\n"
                      "python tools/vorlage-pruefen.py NEUE_VORLAGE.xlsx -v       # mit Beleg je Fund\n"
                      "python tools/vorlage-pruefen.py --ohne-rechentests         # ohne LibreOffice (nur Formeln, Formate, Texte)\n"
                      "python tools/vorlage-pruefen.py --bericht docs/VORLAGE-FEHLERBERICHT.md --pdf docs/VORLAGE-FEHLERBERICHT.pdf\n"
                      "python tools/vorlage-pruefen.py --korrekturvorschlag docs/WEP_Vorlage_Korrekturvorschlag.xlsx"))
    b.append(("p", "Voraussetzungen: Python 3.10+, `pip install openpyxl lxml`, für die Rechentests LibreOffice, für `--pdf` Chrome, Edge oder Chromium. "
                   "Der Exit-Code ist die Zahl der noch vorhandenen sicheren Funde (0 = alles behoben). Das Skript liest die Datei nur; Rechentests "
                   "laufen in Kopien im temporären Ordner."))
    b.append(("p", "**Änderungen im Korrekturvorschlag** (gelb = geändert; Zellsperren sind nicht gelb markiert; Blattschutz bleibt aus; "
                   "Kontrollkästchen, Skizzen und Formatierung bleiben erhalten – es wird direkt im Dateiinhalt geändert):"))
    gruppen: dict[tuple, list[Korrektur]] = {}
    for k in _korrekturen():
        if k.gruppe:
            schluessel = (k.fund, k.blatt, k.art, k.gruppe)
        elif k.art in ("formel", "text"):
            schluessel = (k.fund, k.blatt, k.art, k.zelle)          # jede Zelle für sich
        elif k.art == "name_loeschen":
            schluessel = (k.fund, k.blatt, k.art, "")
        else:
            schluessel = (k.fund, k.blatt, k.art, str(k.neu))
        gruppen.setdefault(schluessel, []).append(k)
    kz = []

    def kurz(s: str, n: int = 90) -> str:
        return s if len(s) <= n else s[: n - 1] + "…"

    for (fund, blatt, art, _), liste in gruppen.items():
        n0 = liste[0].neu
        if art == "name_loeschen":
            kz.append([fund, blatt, "Namen", "vorhanden", "gelöscht: " + ", ".join(str(k.neu) for k in liste)])
            continue
        zellen = [k.zelle for k in liste]
        zt = ", ".join(zellen[:6]) + (f" … ({len(zellen)} Zellen)" if len(zellen) > 6 else "")
        ws_alt = m.blatt(blatt)
        alt = ""
        if art in ("formel", "text", "zahl", "leer"):
            alte = []
            for z in zellen:
                v = ws_alt[z].value
                s = "leer" if v is None else fx(v)
                if s not in alte:
                    alte.append(s)
            alt = kurz(", ".join(alte[:5]) + (" …" if len(alte) > 5 else ""))
            neu = {"formel": de("=" + str(n0)), "text": f"„{n0}“", "zahl": str(n0), "leer": "leer"}[art]
            if len(liste) > 1 and art == "formel":
                neu += " (entsprechend in jeder Zeile)"
            neu = kurz(neu, 120)
        elif art == "format":
            alt = ", ".join(sorted({ws_alt[z].number_format for z in zellen}))
            neu = "Zahlenformat „0 %“"
        elif art in ("entsperren", "sperren"):
            alt = "gesperrt" if art == "entsperren" else "nicht gesperrt"
            neu = "entsperrt" if art == "entsperren" else "gesperrt"
        elif art in ("dv_zahl", "dv_ganz"):
            dv = _dv(m, blatt, zellen[0])
            alt = f"Liste {dv.formula1}" if dv is not None else "keine Prüfung"
            neu = f"Eingabeprüfung {'ganze ' if art == 'dv_ganz' else ''}Zahl {n0[0]} bis {n0[1]}"
        elif art == "seite":
            alt = f"{ws_alt.page_setup.scale} %"
            neu = f"Skalierung {n0} %"
        else:
            alt = "kein Druckbereich"
            neu = f"Druckbereich {n0}"
        kz.append([fund, blatt.strip(), zt, alt, neu])
    b.append(("tabelle", {"kopf": ["Fund", "Blatt", "Zelle(n)", "vorher", "nachher"], "zeilen": kz, "klasse": "klein"}))
    if behoben:
        b.append(("hinweis", "In der geprüften Datei bereits behoben: " + ", ".join(behoben)))
    return b


def _inline_md(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", "<br>")


def als_markdown(blocks: list[tuple]) -> str:
    out: list[str] = []
    for blk in blocks:
        typ = blk[0]
        if typ == "h1":
            out.append(f"# {blk[1]}\n")
        elif typ == "h2":
            out.append(f"\n## {blk[1]}\n")
        elif typ == "h3":
            out.append(f"\n### {blk[1]}\n")
        elif typ == "p":
            out.append(blk[1] + "\n")
        elif typ == "hinweis":
            out.append("> " + blk[1] + "\n")
        elif typ == "ul":
            out.append("\n".join(f"- {x}" for x in blk[1]) + "\n")
        elif typ == "code":
            out.append("```\n" + blk[1] + "\n```\n")
        elif typ == "meta":
            out.append("\n".join(f"- **{k}:** {v}" for k, v in blk[1]) + "\n")
        elif typ == "tabelle":
            t = blk[1]
            out.append("| " + " | ".join(t["kopf"]) + " |")
            out.append("|" + "|".join("---" for _ in t["kopf"]) + "|")
            for z in t["zeilen"]:
                out.append("| " + " | ".join(_inline_md(str(c)) for c in z) + " |")
            out.append("")
        # „neueSeite“ hat im Markdown keine Entsprechung
    return "\n".join(out).rstrip() + "\n"


CSS = """
@page { size: A4; margin: 14mm 12mm 16mm 12mm; @bottom-center { content: "Fehlerbericht WEP-Vorlage – Seite " counter(page) " von " counter(pages); font: 8pt sans-serif; color: #555; } }
@page wide { size: A4 landscape; margin: 12mm 10mm 14mm 10mm; @bottom-center { content: "Fehlerbericht WEP-Vorlage – Seite " counter(page) " von " counter(pages); font: 8pt sans-serif; color: #555; } }
html { font-family: "Segoe UI", "Liberation Sans", Arial, sans-serif; font-size: 10pt; color: #1a1a1a; line-height: 1.38; }
h1 { font-size: 20pt; margin: 0 0 4mm; } h2 { font-size: 14pt; margin: 7mm 0 2mm; border-bottom: 1.5px solid #333; padding-bottom: 1mm; }
h3 { font-size: 11.5pt; margin: 5mm 0 1.5mm; } p { margin: 1.6mm 0; } ul { margin: 1mm 0 2mm 5mm; padding: 0; } li { margin: 0.8mm 0; }
code, pre { font-family: Consolas, "DejaVu Sans Mono", monospace; font-size: 8.5pt; background: #f1f1f1; } pre { padding: 2mm; white-space: pre-wrap; }
table { border-collapse: collapse; width: 100%; margin: 2mm 0 3mm; } th, td { border: 0.6px solid #999; padding: 1.2mm 1.6mm; vertical-align: top; text-align: left; }
th { background: #e8e8e8; } tr { page-break-inside: avoid; }
table.klein { font-size: 8.5pt; } table.funde { font-size: 7.6pt; table-layout: fixed; }
.wide { page: wide; } blockquote { border-left: 3px solid #999; margin: 2mm 0; padding-left: 3mm; }
ul.meta { list-style: none; margin-left: 0; }
"""


def _html_inline(s: str) -> str:
    s = html.escape(s, quote=False)
    s = re.sub(r"&lt;br&gt;", "<br>", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", s)
    s = re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
    return s.replace("\n", "<br>")


def als_html(blocks: list[tuple]) -> str:
    teile = []
    offen = False
    for blk in blocks:
        typ = blk[0]
        if typ == "neueSeite":
            if offen:
                teile.append("</section>")
            teile.append(f'<section class="{"wide" if blk[1] == "wide" else ""}" style="break-before: page">')
            offen = True
        elif typ == "h1":
            teile.append(f"<h1>{_html_inline(blk[1])}</h1>")
        elif typ in ("h2", "h3"):
            teile.append(f"<{typ}>{_html_inline(blk[1])}</{typ}>")
        elif typ == "p":
            teile.append(f"<p>{_html_inline(blk[1])}</p>")
        elif typ == "hinweis":
            teile.append(f"<blockquote>{_html_inline(blk[1])}</blockquote>")
        elif typ == "ul":
            teile.append("<ul>" + "".join(f"<li>{_html_inline(x)}</li>" for x in blk[1]) + "</ul>")
        elif typ == "meta":
            teile.append('<ul class="meta">' + "".join(f"<li><b>{html.escape(k)}:</b> {_html_inline(v)}</li>" for k, v in blk[1]) + "</ul>")
        elif typ == "code":
            teile.append(f"<pre>{html.escape(blk[1])}</pre>")
        elif typ == "tabelle":
            t = blk[1]
            kopf = "".join(f"<th>{_html_inline(h)}</th>" for h in t["kopf"])
            zeilen = "".join("<tr>" + "".join(f"<td>{_html_inline(str(c))}</td>" for c in z) + "</tr>" for z in t["zeilen"])
            cols = "".join(f'<col style="width:{w}%">' for w in t["breiten"]) if t.get("breiten") else ""
            teile.append(f'<table class="{t.get("klasse", "")}"><colgroup>{cols}</colgroup><thead><tr>{kopf}</tr></thead><tbody>{zeilen}</tbody></table>')
    if offen:
        teile.append("</section>")
    return f'<!doctype html><html lang="de"><head><meta charset="utf-8"><title>Fehlerbericht WEP-Vorlage</title><style>{CSS}</style></head><body>{"".join(teile)}</body></html>'


def finde_browser() -> list[str]:
    """Chrome, Edge oder Chromium (auch die „headless shell“ von Playwright) – bevorzugt die headless shell."""
    kandidaten: list[Any] = []
    for wurzel in (Path("/opt/pw-browsers"), Path.home() / ".cache" / "ms-playwright"):
        kandidaten += sorted(wurzel.glob("chromium_headless_shell-*/chrome-headless-shell-*/chrome-headless-shell"))
        kandidaten += sorted(wurzel.glob("chromium-*/chrome-linux*/chrome"))
    kandidaten += [shutil.which(n) for n in ("chromium", "chromium-browser", "google-chrome", "chrome", "msedge")]
    kandidaten += [r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                   r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
                   "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"]
    return [str(k) for k in kandidaten if k and Path(k).exists()]


def html_zu_pdf(html_text: str, ziel: Path) -> None:
    browser = finde_browser()
    if not browser:
        raise RuntimeError("Kein Chrome/Edge/Chromium gefunden – PDF nicht erzeugt (die Markdown-Datei ist fertig).")
    tmp = Path(tempfile.mkdtemp(prefix="vorlage-bericht-"))
    try:
        quelle = tmp / "bericht.html"
        quelle.write_text(html_text, encoding="utf8")
        ziel.parent.mkdir(parents=True, exist_ok=True)
        for exe in browser:
            cmd = [exe, "--headless", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                   f"--user-data-dir={tmp / 'profil'}", f"--print-to-pdf={ziel.resolve()}", quelle.as_uri()]
            try:
                subprocess.run(cmd, capture_output=True, timeout=90, check=False)
            except subprocess.TimeoutExpired:
                continue
            if ziel.exists():
                return
        raise RuntimeError("Der Browser hat kein PDF erzeugt")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def schreibe_bericht(m: Mappe, ctx: Ctx, befunde: dict[str, Befund], md_pfad: Path, pdf_pfad: Path | None) -> None:
    blocks = baue_bericht(m, ctx, befunde)
    md_pfad.parent.mkdir(parents=True, exist_ok=True)
    md_pfad.write_text(als_markdown(blocks), encoding="utf8")
    print(f"Bericht geschrieben: {md_pfad}")
    if pdf_pfad:
        html_zu_pdf(als_html(blocks), pdf_pfad)
        print(f"PDF geschrieben: {pdf_pfad}")



if __name__ == "__main__":
    sys.exit(main())
