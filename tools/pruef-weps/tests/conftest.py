"""Erfundene Testdateien, aus der Vorlage erzeugt und mit LibreOffice berechnet (gespeicherte Werte).

Es liegen keine WEPs im Repository: Alles entsteht zur Laufzeit im temporären Ordner aus
docs/WEP_Vorlage.xlsx (die Vorlage selbst wird nur gelesen).
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

HIER = Path(__file__).resolve().parent
sys.path.insert(0, str(HIER.parent))

VORLAGE = HIER.parents[2] / "docs" / "WEP_Vorlage.xlsx"
LO = shutil.which("libreoffice") or shutil.which("soffice")

# Die fünf Fehler und ihre richtigen Formeln: (Blatt, Zelle, richtig)
KORREKTUREN = {
    "C33": ("Baulichkeiten", "C33", "=$H$32"),
    "K81": ("Baulichkeiten", "K81", "=K61+K71+K75+K77+K79"),
    "J52": ("Baulichkeiten", "J52", "=$B$47"),
    "C64": ("Baulichkeiten", "C64", "=I63"),
    "K21": ("Beseitigungskosten", "K21", "=(F21*H21*J21)"),
}


def eingaben_voll(wb, typen: dict[str, str], beispielwerte_loeschen: bool = True) -> None:
    """Erfundener Fall, in dem jeder Fehler wirkt: Laube, Anbau A1/A2, Schuppen, integrierter Vorplatz, Beseitigung Zeile 9."""
    fo = wb["Formalien"]
    fo["E3"], fo["E4"], fo["E5"], fo["E6"], fo["E17"] = "Musteranlage Nord", "12", "Weg 3", "Erika Musterfrau", "V-2026-001"
    bl = wb["Baulichkeiten"]
    bl["G12"], bl["C13"], bl["I13"], bl["E14"], bl["A17"] = typen["stein1"], 1990, 20, 5, 12
    bl["G21"], bl["C22"], bl["I22"], bl["E23"], bl["A26"] = typen["stein1"], 1995, 25, 10, 6
    bl["G30"], bl["C31"], bl["I31"], bl["E32"], bl["A35"] = typen["stein2"], 2005, 35, 20, 5
    bl["G47"], bl["C48"], bl["I48"], bl["E49"], bl["A52"] = typen["block"], 2000, 18, 15, 8
    bl["E63"], bl["F71"] = 275, 10
    be = wb["Beseitigungskosten"]
    be["H21"] = 80
    if beispielwerte_loeschen:
        bl["F104"], bl["G104"] = 0, 0
        bl["I19"], bl["I28"], bl["I37"] = 0, 0, 0


def eingaben_nur_laube(wb, typen: dict[str, str]) -> None:
    """Nur eine Laube: die fehlerhaften Formeln bleiben, wirken aber nicht."""
    fo = wb["Formalien"]
    fo["E3"], fo["E4"], fo["E17"] = "Musteranlage Süd", "7", "V-2026-002"
    bl = wb["Baulichkeiten"]
    bl["G12"], bl["C13"], bl["I13"], bl["E14"], bl["A17"] = typen["stein1"], 1990, 20, 5, 12
    bl["F104"], bl["G104"] = 0, 0
    bl["I19"], bl["I28"], bl["I37"] = 0, 0, 0


@pytest.fixture(scope="session")
def dateien(tmp_path_factory):
    """Dict Name → Pfad der von LibreOffice berechneten Testdateien."""
    if not LO:
        pytest.skip("LibreOffice nicht gefunden (für die gespeicherten Werte nötig)")
    import openpyxl

    roh = tmp_path_factory.mktemp("roh")
    ziel = tmp_path_factory.mktemp("berechnet")
    cache = openpyxl.load_workbook(VORLAGE, data_only=True)["Baulichkeiten"]
    typen = {"stein1": cache["N45"].value, "stein2": cache["N46"].value, "block": cache["N50"].value}

    def neu():
        return openpyxl.load_workbook(VORLAGE)

    def speichere(name: str, wb) -> None:
        wb.save(roh / f"{name}.xlsx")

    # A: alle Fehler wirken (Beispielwerte überschrieben) – Referenz: dieselben Eingaben mit einzeln/allen korrigierten Formeln
    wb = neu(); eingaben_voll(wb, typen); speichere("A_alle_fehler", wb)
    wb = neu(); eingaben_voll(wb, typen)
    for blatt, zelle, formel in KORREKTUREN.values():
        wb[blatt][zelle] = formel
    speichere("A_alle_korrigiert", wb)
    for stelle, (blatt, zelle, formel) in KORREKTUREN.items():
        wb = neu(); eingaben_voll(wb, typen); wb[blatt][zelle] = formel
        speichere(f"A_nur_{stelle}_korrigiert", wb)
    # B: Beispielwerte „nicht beräumt“ unverändert
    wb = neu(); eingaben_voll(wb, typen, beispielwerte_loeschen=False); speichere("B_beispielwerte", wb)
    # C: fehlerhafte Formeln, aber nur eine Laube – ohne Wirkung
    wb = neu(); eingaben_nur_laube(wb, typen); speichere("C_ohne_wirkung", wb)
    # D: Vorlage korrigiert, ohne Fehler und ohne Beispielwerte
    wb = neu(); eingaben_nur_laube(wb, typen)
    for blatt, zelle, formel in KORREKTUREN.values():
        wb[blatt][zelle] = formel
    speichere("D_korrekt", wb)

    profil = roh / "lo-profil"
    namen = sorted(p.name for p in roh.glob("*.xlsx"))
    subprocess.run(
        [LO, f"-env:UserInstallation=file://{profil}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(ziel), *[str(roh / n) for n in namen]],
        check=True, capture_output=True, timeout=280,
    )
    ergebnis = {Path(n).stem: ziel / n for n in namen}

    # E: ohne gespeicherte Werte (von openpyxl gespeichert, nie berechnet) – gleicher Inhalt wie A
    ergebnis["E_ohne_werte"] = roh / "A_alle_fehler.xlsx"
    return ergebnis


@pytest.fixture()
def ordner(tmp_path, dateien):
    """Prüfordner mit Kopien der Testdateien, einer Unterordner-Datei und Störfällen."""
    import openpyxl

    for name, pfad in dateien.items():
        shutil.copy(pfad, tmp_path / f"{name}.xlsx")
    unter = tmp_path / "unterordner" / "tief"
    unter.mkdir(parents=True)
    shutil.copy(dateien["A_alle_fehler"], unter / "A_im_unterordner.xlsx")
    # keine Vorlage
    wb = openpyxl.Workbook(); wb.active["A1"] = "Kein WEP"; wb.save(tmp_path / "F_keine_vorlage.xlsx")
    # beschädigt
    (tmp_path / "G_kaputt.xlsx").write_bytes(b"das ist keine Excel-Datei")
    # Excel-Sperrdatei und Fremdformat: müssen ignoriert werden
    (tmp_path / "~$A_alle_fehler.xlsx").write_bytes(b"x")
    (tmp_path / "notizen.txt").write_text("egal")
    return tmp_path
