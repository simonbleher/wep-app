from __future__ import annotations

import csv
import hashlib
import os
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import pytest

import pruefe_weps as p
from conftest import KORREKTUREN

SKRIPT = Path(p.__file__)


def e28(pfad) -> Decimal:
    """Deckblatt E28, von LibreOffice berechnet."""
    import openpyxl
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb[[n for n in wb.sheetnames if n.strip() == "Deckbl UPV"][0]]
    return p.zahl(ws["E28"].value)


def deckblatt_wert(pfad, zelle) -> Decimal:
    import openpyxl
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb[[n for n in wb.sheetnames if n.strip() == "Deckbl UPV"][0]]
    return p.zahl(ws[zelle].value)


def befund(erg, id_):
    return next(b for b in erg.befunde if b.id == id_)


# --- Rechenhilfen ---------------------------------------------------------------------------

def test_runde_wie_excel():
    assert p.runde(Decimal("2.675")) == Decimal("2.68")
    assert p.runde(Decimal("-2.675")) == Decimal("-2.68")
    assert p.zahl(1.7919999999999998) == Decimal("1.792")
    assert p.zahl(True) is None and p.zahl("5") is None and p.zahl(None) is None


def test_euro_format():
    assert p.eur(Decimal("1234.5")) == "1.234,50 €"
    assert p.eur(Decimal("-80")) == "−80,00 €"
    assert p.eur(Decimal("80"), True) == "+80,00 €"
    assert p.csv_zahl(Decimal("1234.5")) == "1234,50"


def test_csv_text_gegen_formeln():
    assert p.csv_text("=1+1") == "'=1+1"
    assert p.csv_text("Musteranlage") == "Musteranlage"


def test_norm_formel():
    assert p.norm_formel("= $H$23 ") == "H23"
    assert p.norm_formel(None) == ""
    assert p.norm_formel(5).startswith("WERT:")


# --- Prüfung gegen LibreOffice-Rechnung --------------------------------------------------------

def test_alle_fehler_wirken_und_stimmen_mit_korrigierter_datei(dateien):
    erg = p.pruefe_datei(str(dateien["A_alle_fehler"]))
    assert erg.status == "betroffen"
    assert (erg.anlage, erg.parzelle, erg.vorgang) == ("Musteranlage Nord", "12 / Weg 3", "V-2026-001")
    ist = e28(dateien["A_alle_fehler"])
    soll = e28(dateien["A_alle_korrigiert"])
    assert ist != soll
    # Deckblatt: Ist = gespeichert, Soll = LibreOffice-Rechnung mit korrigierten Formeln
    for zelle in ("E21", "E25", "E26", "E28"):
        w = erg.deckblatt[zelle]
        assert w["ist"] == deckblatt_wert(dateien["A_alle_fehler"], zelle)
        assert abs(w["soll"] - deckblatt_wert(dateien["A_alle_korrigiert"], zelle)) <= Decimal("0.01"), zelle
    assert abs(erg.gesamt_diff - (soll - ist)) <= Decimal("0.01")


@pytest.mark.parametrize("stelle", ["C33", "K81", "C64", "K21"])
def test_einzelner_fehler_betrag_wie_libreoffice(dateien, stelle):
    erg = p.pruefe_datei(str(dateien["A_alle_fehler"]))
    b = befund(erg, stelle)
    assert b.zustand == "fehlerhaft" and b.wirkung == "euro"
    erwartet = e28(dateien[f"A_nur_{stelle}_korrigiert"]) - e28(dateien["A_alle_fehler"])
    assert abs(b.betrag - erwartet) <= Decimal("0.01"), (stelle, b.betrag, erwartet)
    assert abs(b.betrag) > Decimal("1")


def test_j52_nur_anzeige(dateien):
    b = befund(p.pruefe_datei(str(dateien["A_alle_fehler"])), "J52")
    assert b.zustand == "fehlerhaft" and b.wirkung == "anzeige" and b.betrag == 0


def test_vorzeichen_und_richtung(dateien):
    erg = p.pruefe_datei(str(dateien["A_alle_fehler"]))
    k81 = befund(erg, "K81")   # Schuppen doppelt: Gesamtbetrag zu hoch
    k21 = befund(erg, "K21")   # Zeile fehlt: Abzug zu niedrig, Gesamtbetrag zu hoch
    assert k81.betrag < 0 and "zu hoch" in p.richtung(k81.betrag)
    assert k21.betrag == Decimal("-80.00") and k21.delta == {"E26": Decimal("80.00"), "E28": Decimal("-80.00")}


def test_beispielwerte_als_verdacht(dateien):
    erg = p.pruefe_datei(str(dateien["B_beispielwerte"]))
    b = befund(erg, "F104")
    assert b.zustand == "beispielwert" and b.betrag == Decimal("250")
    assert befund(erg, "I19").zustand == "beispielwert"
    assert erg.status == "betroffen"  # die Formelfehler wirken auch
    assert erg.deckblatt["E28"]["soll_verdacht"] == erg.deckblatt["E28"]["soll"] + 250
    assert erg.deckblatt["E25"]["soll_verdacht"] == erg.deckblatt["E25"]["soll"] - 250


def test_fehler_ohne_wirkung(dateien):
    erg = p.pruefe_datei(str(dateien["C_ohne_wirkung"]))
    assert erg.status == "hinweis"
    for stelle in ("C33", "K81", "C64", "K21"):
        b = befund(erg, stelle)
        assert b.zustand == "fehlerhaft" and b.wirkung == "keine" and ist_null(b.betrag), stelle
    assert befund(erg, "J52").wirkung == "keine"
    assert ist_null(erg.gesamt_diff)
    assert all(ist_null(w["soll"] - w["ist"]) for w in erg.deckblatt.values())


def ist_null(x):
    return p.ist_null(x)


def test_korrigierte_datei_ohne_befund(dateien):
    erg = p.pruefe_datei(str(dateien["D_korrekt"]))
    assert erg.status == "ohne_befund"
    assert all(b.zustand == "korrekt" for b in erg.befunde)


def test_von_hand_geaenderte_formel_wird_gemeldet(dateien, tmp_path):
    import openpyxl, shutil
    ziel = tmp_path / "x.xlsx"
    shutil.copy(dateien["A_alle_fehler"], ziel)
    wb = openpyxl.load_workbook(ziel)
    wb["Baulichkeiten"]["K81"] = "=K61+K71"
    wb.save(ziel)
    b = befund(p.pruefe_datei(str(ziel)), "K81")
    assert b.zustand == "abweichend" and b.betrag is None


def test_ohne_gespeicherte_werte(dateien):
    erg = p.pruefe_datei(str(dateien["E_ohne_werte"]))
    assert erg.status == "werte_fehlen"
    assert "keine gespeicherten Werte" in erg.grund and "C33" in erg.grund
    assert erg.deckblatt == {}


# --- Ordnerlauf, Berichte, nur lesen -----------------------------------------------------------

def sha(ordner: Path) -> dict[str, str]:
    return {str(f): hashlib.sha256(f.read_bytes()).hexdigest() for f in sorted(ordner.rglob("*")) if f.is_file()}


def test_ordnerlauf_berichte_und_originale_unveraendert(ordner):
    vorher = sha(ordner)
    r = subprocess.run([sys.executable, str(SKRIPT), str(ordner)], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stderr
    nachher = sha(ordner)
    neu = set(nachher) - set(vorher)
    assert {Path(f).name for f in neu} == {"pruefbericht.html", "pruefbericht.csv"}
    assert all(nachher[f] == h for f, h in vorher.items())  # Originale byte-gleich

    zeilen = list(csv.reader(open(ordner / "pruefbericht.csv", encoding="utf-8-sig", newline=""), delimiter=";"))
    kopf, daten = zeilen[0], zeilen[1:]
    assert kopf[:5] == ["Datei", "Status", "Anlage", "Parzelle", "Vorgangs-Nr."]
    dateien = {z[0] for z in daten}
    assert "A_alle_fehler.xlsx" in dateien
    assert os.path.join("unterordner", "tief", "A_im_unterordner.xlsx") in dateien
    assert not any(d.startswith("~$") for d in dateien)       # Sperrdatei ignoriert
    status = {z[0]: z[1] for z in daten}
    assert status["F_keine_vorlage.xlsx"] == p.STATUS_TEXT["keine_vorlage"]
    assert status["G_kaputt.xlsx"] == p.STATUS_TEXT["nicht_lesbar"]
    assert status["E_ohne_werte.xlsx"] == p.STATUS_TEXT["werte_fehlen"]
    assert status["D_korrekt.xlsx"] == p.STATUS_TEXT["ohne_befund"]
    k21 = next(z for z in daten if z[0] == "A_alle_fehler.xlsx" and z[6].endswith("K21"))
    assert k21[9] == "-80,00" and k21[3] == "12 / Weg 3" and k21[4] == "V-2026-001"

    h = (ordner / "pruefbericht.html").read_text(encoding="utf-8")
    assert "Personendaten" in h and "Musteranlage Nord" in h and "V-2026-001" in h
    assert "http://" not in h and "https://" not in h and "<script" not in h   # nichts nach außen
    # betroffene WEPs stehen vor den übrigen Dateien
    assert h.index("Betroffene WEPs") < h.index("Einzelheiten je Datei") < h.index("Weitere Dateien")
    assert h.index("A_alle_fehler.xlsx") < h.index("F_keine_vorlage.xlsx")


def test_html_maskiert_sonderzeichen(dateien, tmp_path):
    import openpyxl, shutil
    ziel = tmp_path / "x.xlsx"
    shutil.copy(dateien["A_alle_fehler"], ziel)
    wb = openpyxl.load_workbook(ziel)
    wb["Formalien"]["E3"] = "<b>=cmd|' /C calc'!A0</b>"
    wb.save(ziel)
    erg = p.pruefe_ordner(str(tmp_path), ausgabe=lambda *_: None)
    p.schreibe_html(str(tmp_path / "b.html"), str(tmp_path), erg)
    p.schreibe_csv(str(tmp_path / "b.csv"), erg)
    h = (tmp_path / "b.html").read_text(encoding="utf-8")
    assert "<b>=cmd" not in h and "&lt;b&gt;" in h
    assert "'<b>=cmd" in (tmp_path / "b.csv").read_text(encoding="utf-8-sig") or "<b>=cmd" in (tmp_path / "b.csv").read_text(encoding="utf-8-sig")


def test_leerer_und_fehlender_ordner(tmp_path):
    r = subprocess.run([sys.executable, str(SKRIPT), str(tmp_path)], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0 and (tmp_path / "pruefbericht.html").exists()
    r = subprocess.run([sys.executable, str(SKRIPT), str(tmp_path / "gibtsnicht")], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 2
