#!/usr/bin/env python3
"""Prüft fertige WEPs (Excel) auf die Fehler der WEP-Vorlage des BV Pankow.

Die Vorlage hatte bis Beta 4 fünf Formelfehler (Blätter „Baulichkeiten“ und „Beseitigungskosten“)
und Beispielwerte, die im Protokoll stehen bleiben konnten. Das Skript sucht in einem Ordner
(einschließlich Unterordnern) alle .xlsx/.xlsm, erkennt, ob sie auf der Vorlage beruhen, und rechnet
für jede Stelle aus, ob der Fehler das Ergebnis verfälscht hat und um wie viel Euro – bis zum
Deckblatt (Deckbl UPV E21/E25/E26/E28).

Gerechnet wird mit den **gespeicherten Werten** der Dateien (so, wie Excel sie zuletzt gespeichert hat)
und einer eigenen Nachrechnung der richtigen Formel. Die Dateien werden nur gelesen, nie verändert,
und es wird nichts ins Netz gesendet. Ergebnis: pruefbericht.html und pruefbericht.csv im geprüften Ordner.

Aufruf:  python pruefe_weps.py ORDNER [--ausgabe ANDERER_ORDNER]
Siehe README.md und docs/ENTSCHEIDUNGEN.md (Beta 4, Paket 2) im Repository.
"""
from __future__ import annotations

import argparse
import csv
import html
import os
import sys
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Iterable

__version__ = "1.0"

BERICHT_HTML = "pruefbericht.html"
BERICHT_CSV = "pruefbericht.csv"

BL = "Baulichkeiten"
BE = "Beseitigungskosten"
DB = "Deckblatt"  # Blatt „Deckbl UPV “ (mit Leerzeichen am Ende)
FO = "Formalien"

NULL = Decimal(0)
CENT = Decimal("0.01")

# Bereiche, die gelesen werden (Blatt, erste/letzte Zeile, erste/letzte Spalte)
BEREICHE = {
    FO: (1, 20, 1, 6),
    DB: (20, 30, 1, 6),
    BL: (12, 105, 1, 11),
    BE: (12, 24, 1, 11),
}


# ---------------------------------------------------------------------------------------------
# Rechenhilfen (wie Excel: ROUND rundet .5 weg von null)
# ---------------------------------------------------------------------------------------------

def zahl(v: Any) -> Decimal | None:
    """Gespeicherter Zellwert als Decimal, sonst None (leer, Text, Fehlerwert, Datum)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(round(float(v), 10)))
    if isinstance(v, Decimal):
        return v
    return None


def z0(v: Any) -> Decimal:
    """Wie zahl(), aber leer/Text = 0 (so rechnet Excel in Produkten und Summen)."""
    d = zahl(v)
    return NULL if d is None else d


def runde(x: Decimal, stellen: int = 2) -> Decimal:
    return x.quantize(Decimal(1).scaleb(-stellen), rounding=ROUND_HALF_UP)


def ist_null(x: Decimal | None) -> bool:
    return x is None or abs(x) < Decimal("0.005")


def eur(x: Decimal | None, vorzeichen: bool = False) -> str:
    """1.234,56 € (deutsches Format)."""
    if x is None:
        return "–"
    x = runde(x)
    s = f"{abs(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if x < 0:
        s = "−" + s
    elif vorzeichen and x > 0:
        s = "+" + s
    return s + " €"


def csv_zahl(x: Decimal | None) -> str:
    return "" if x is None else str(runde(x)).replace(".", ",")


def csv_text(s: Any) -> str:
    """Schutz vor Formel-Einschleusung in Excel: Text, der mit = + - @ beginnt, bekommt ein Hochkomma."""
    t = "" if s is None else str(s)
    return "'" + t if t[:1] in ("=", "+", "-", "@", "\t", "\r") else t


def norm_formel(raw: Any) -> str:
    """Formel ohne „=“, Leerzeichen und $, in Großbuchstaben. Leere Zelle = ''. Fester Wert = 'WERT:…'."""
    if raw is None:
        return ""
    if isinstance(raw, str) and raw.startswith("="):
        return raw[1:].replace(" ", "").replace("$", "").upper()
    return f"WERT:{raw}"


# ---------------------------------------------------------------------------------------------
# Datei lesen
# ---------------------------------------------------------------------------------------------

@dataclass
class Mappe:
    """Die benötigten Zellen einer Datei: Formeln/Eingaben und gespeicherte Werte."""
    formeln: dict[tuple[str, str], Any] = field(default_factory=dict)
    werte: dict[tuple[str, str], Any] = field(default_factory=dict)

    def roh(self, blatt: str, zelle: str) -> Any:
        return self.formeln.get((blatt, zelle))

    def wert(self, blatt: str, zelle: str) -> Any:
        return self.werte.get((blatt, zelle))

    def n(self, blatt: str, zelle: str) -> Decimal | None:
        return zahl(self.wert(blatt, zelle))

    def n0(self, blatt: str, zelle: str) -> Decimal:
        return z0(self.wert(blatt, zelle))


class KeineVorlage(Exception):
    pass


def _blatt_zuordnen(namen: Iterable[str]) -> dict[str, str]:
    """Kanonischer Name → tatsächlicher Blattname (Groß-/Kleinschreibung und Leerzeichen am Rand egal)."""
    ziel = {BL: "baulichkeiten", BE: "beseitigungskosten", FO: "formalien", DB: "deckbl upv"}
    gefunden: dict[str, str] = {}
    for name in namen:
        k = " ".join(name.split()).casefold()
        for kanon, soll in ziel.items():
            if k == soll and kanon not in gefunden:
                gefunden[kanon] = name
    return gefunden


def lese_mappe(pfad: str) -> Mappe:
    """Liest die benötigten Zellen (nur lesend). Wirft KeineVorlage, wenn die Blätter fehlen."""
    from openpyxl import load_workbook

    mappe = Mappe()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        for datenwerte, ziel in ((False, mappe.formeln), (True, mappe.werte)):
            wb = load_workbook(pfad, read_only=True, data_only=datenwerte, keep_links=False)
            try:
                zuordnung = _blatt_zuordnen(wb.sheetnames)
                fehlend = [k for k in BEREICHE if k not in zuordnung]
                if fehlend:
                    raise KeineVorlage("Blatt fehlt: " + ", ".join(
                        {BL: "Baulichkeiten", BE: "Beseitigungskosten", FO: "Formalien", DB: "Deckbl UPV"}[k] for k in fehlend))
                for kanon, (r1, r2, c1, c2) in BEREICHE.items():
                    ws = wb[zuordnung[kanon]]
                    for zeile in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                        for c in zeile:
                            v = getattr(c, "value", None)
                            if v is not None:
                                ziel[(kanon, c.coordinate)] = v
            finally:
                wb.close()
    return mappe


# ---------------------------------------------------------------------------------------------
# Befunde
# ---------------------------------------------------------------------------------------------

@dataclass
class Befund:
    id: str
    stelle: str            # Kurzbeschreibung
    zellen: str            # z. B. „Baulichkeiten!C33“
    art: str               # 'fehler' (Formel) | 'beispiel' (Beispielwert) | 'hinweis'
    zustand: str           # 'fehlerhaft' | 'korrekt' | 'abweichend' | 'beispielwert'
    wirkung: str           # 'euro' | 'anzeige' | 'keine' | 'unbekannt'
    betrag: Decimal | None = None          # Soll − Ist beim Gesamtbetrag (Deckbl UPV E28)
    delta: dict[str, Decimal] = field(default_factory=dict)   # Deckblatt-Zelle → Soll − Ist
    text: str = ""

    @property
    def auffaellig(self) -> bool:
        return self.zustand != "korrekt"


@dataclass
class DateiErgebnis:
    pfad: str
    rel: str
    status: str = "ohne_befund"
    grund: str = ""
    anlage: str = ""
    parzelle: str = ""
    vorgang: str = ""
    befunde: list[Befund] = field(default_factory=list)
    deckblatt: dict[str, dict[str, Decimal | None]] = field(default_factory=dict)  # Zelle → ist/soll/soll_verdacht
    warnungen: list[str] = field(default_factory=list)

    @property
    def gesamt_diff(self) -> Decimal:
        return sum((b.betrag for b in self.befunde if b.art == "fehler" and b.betrag), NULL)

    @property
    def verdacht_diff(self) -> Decimal:
        return sum((b.betrag for b in self.befunde if b.art == "beispiel" and b.betrag), NULL)


STATUS_TEXT = {
    "betroffen": "Ergebnis betroffen",
    "verdacht": "Verdacht (Beispielwert)",
    "hinweis": "Hinweise, ohne Euro-Wirkung",
    "ohne_befund": "ohne Befund",
    "werte_fehlen": "keine gespeicherten Werte",
    "keine_vorlage": "keine WEP-Vorlage",
    "nicht_lesbar": "nicht lesbar",
}
STATUS_REIHE = ["betroffen", "verdacht", "hinweis", "werte_fehlen", "ohne_befund", "keine_vorlage", "nicht_lesbar"]


def _klass(raw: Any, fehler: set[str], richtig: set[str]) -> str:
    n = norm_formel(raw)
    if n in fehler:
        return "fehlerhaft"
    if n in richtig:
        return "korrekt"
    return "abweichend"


def _abweichend_text(raw: Any) -> str:
    return f"enthält weder die fehlerhafte noch die korrekte Formel ({raw!r}); von Hand geändert, nicht geprüft"


def _wirkung(betrag: Decimal | None) -> str:
    if betrag is None:
        return "unbekannt"
    return "keine" if ist_null(betrag) else "euro"


# --- 1. Baulichkeiten!C33 --------------------------------------------------------------------

def pruefe_c33(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    z = f"{BL}!C33"
    b = Befund("C33", "Anbau A2 rechnet mit dem Ergebnis von Anbau A1", z, "fehler", "", "unbekannt")
    b.zustand = _klass(m.roh(BL, "C33"), {"H23"}, {"H32"})
    if b.zustand == "abweichend":
        b.text = _abweichend_text(m.roh(BL, "C33"))
        return b
    if b.zustand == "korrekt":
        return b
    b.text = "C33 verweist auf H23 (Anbau A1) statt auf H32 (Anbau A2 nach Mängelabzug)."
    if not werte_da:
        return b
    h32, e33, e34, a35, k35 = m.n(BL, "H32"), m.n(BL, "E33"), m.n(BL, "E34"), m.n0(BL, "A35"), m.n(BL, "K35")
    if h32 is None or k35 is None or e34 is None:
        return b
    e33 = e33 or NULL
    h33 = runde(h32 - h32 * e33 / 100)
    h34 = runde(h33 * e34)
    soll = runde(a35 * h34)
    d = soll - k35
    b.delta = {"E21": d}
    b.betrag = d
    b.wirkung = _wirkung(d)
    c33, h23 = m.n(BL, "C33"), m.n(BL, "H23")
    if c33 is not None and h23 is not None and abs(c33 - h23) > CENT:
        warn.append("Baulichkeiten!C33 passt nicht zu H23: gespeicherte Werte sind vermutlich nicht aktuell (Datei nach der letzten Änderung nicht neu berechnet).")
    if ist_null(d):
        b.text += " Bei den eingetragenen Werten bleibt das Ergebnis gleich (kein zweiter Anbau oder gleiche Werte wie A1)."
    else:
        b.text += f" Wert des Anbaus A2 (K35): gespeichert {eur(k35)}, richtig {eur(soll)}."
    return b


# --- 2. Baulichkeiten!K81 --------------------------------------------------------------------

def pruefe_k81(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("K81", "Schuppen (K52) wird in „Baulichkeiten insgesamt“ doppelt gezählt", f"{BL}!K81", "fehler", "", "unbekannt")
    b.zustand = _klass(m.roh(BL, "K81"), {"K52+K61+K71+K75+K77+K79"}, {"K61+K71+K75+K77+K79"})
    if b.zustand == "abweichend":
        b.text = _abweichend_text(m.roh(BL, "K81"))
        return b
    if b.zustand == "korrekt":
        return b
    b.text = "K81 addiert den Schuppen (K52) noch einmal zur Zwischensumme K61, in der er schon steckt."
    if not werte_da:
        return b
    k52 = m.n0(BL, "K52")
    b.delta = {"E21": -k52}
    b.betrag = -k52
    b.wirkung = _wirkung(b.betrag)
    k81 = m.n(BL, "K81")
    teile = sum((m.n0(BL, c) for c in ("K52", "K61", "K71", "K75", "K77", "K79")), NULL)
    if k81 is not None and abs(k81 - teile) > CENT:
        warn.append("Baulichkeiten!K81 passt nicht zu seinen Summanden: gespeicherte Werte sind vermutlich nicht aktuell.")
    b.text += (" Kein Schuppen bewertet – bleibt ohne Wirkung." if ist_null(k52)
               else f" Der Schuppen ist mit {eur(k52)} bewertet; der Gesamtbetrag ist um diesen Betrag zu hoch.")
    return b


# --- 3. Baulichkeiten!J52 --------------------------------------------------------------------

def pruefe_j52(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("J52", "Beschriftung „Wert BT“ beim Schuppen zeigt das Kürzel von A1", f"{BL}!J52", "fehler", "", "keine")
    b.zustand = _klass(m.roh(BL, "J52"), {"B21"}, {"B47"})
    if b.zustand == "abweichend":
        b.text = _abweichend_text(m.roh(BL, "J52"))
        return b
    if b.zustand == "korrekt":
        return b
    b.text = "J52 verweist auf B21 (Kürzel Anbau A1) statt auf B47 (Kürzel Schuppen). Betrifft nur die Anzeige, keinen Betrag."
    if not werte_da:
        b.wirkung = "unbekannt"
        return b
    b.betrag = NULL
    schuppen = not ist_null(m.n(BL, "K52")) or not ist_null(m.n(BL, "A52"))
    if schuppen and m.wert(BL, "J52") != m.wert(BL, "B47"):
        b.wirkung = "anzeige"
        b.text += f" Hier sichtbar falsch: steht „{m.wert(BL, 'J52')}“ statt „{m.wert(BL, 'B47')}“."
    return b


# --- 4. Baulichkeiten!C64 --------------------------------------------------------------------

def pruefe_c64(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("C64", "Integrierter Laubenvorplatz: Wert nach Mängelabzug geht verloren", f"{BL}!C64", "fehler", "", "unbekannt")
    b.zustand = _klass(m.roh(BL, "C64"), {""}, {"I63"})
    if b.zustand == "abweichend":
        b.text = _abweichend_text(m.roh(BL, "C64"))
        return b
    if b.zustand == "korrekt":
        return b
    b.text = "C64 ist leer; es fehlt der Bezug auf I63 (NHK nach Mängelabzug). Betrifft den überdachten Laubenvorplatz „a) integriert“."
    if not werte_da:
        return b
    i63, e15, i68, e70, f71, k71 = (m.n(BL, "I63"), m.n(BL, "E15"), m.n(BL, "I68"), m.n(BL, "E70"), m.n0(BL, "F71"), m.n(BL, "K71"))
    if i63 is None or k71 is None or e70 is None:
        return b
    e64 = (e15 or NULL) if i63 > 0 else NULL
    i64 = runde(i63 - i63 * e64 / 100)
    c70 = i64 if ist_null(i68) else i68
    soll = runde(f71 * runde(c70 * e70))
    d = soll - k71
    b.delta = {"E21": d}
    b.betrag = d
    b.wirkung = _wirkung(d)
    if ist_null(d):
        b.text += " Bei den eingetragenen Werten ohne Wirkung (Vorplatz nicht integriert, angebaut berechnet oder ohne Fläche)."
    else:
        b.text += f" Wert des Vorplatzes (K71): gespeichert {eur(k71)}, richtig {eur(soll)}."
    return b


# --- 5. Beseitigungskosten!K21 ---------------------------------------------------------------

def pruefe_k21(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("K21", "Beseitigung Bau, Zeile 9: Summenformel fehlt, Zeile zählt nicht", f"{BE}!K21", "fehler", "", "unbekannt")
    b.zustand = _klass(m.roh(BE, "K21"), {""}, {"(F21*H21*J21)", "F21*H21*J21"})
    if b.zustand == "abweichend":
        b.text = _abweichend_text(m.roh(BE, "K21"))
        return b
    if b.zustand == "korrekt":
        return b
    b.text = "K21 ist leer (K12:K20 und K22:K23 haben =(Menge×Einzelpreis×Faktor)); die Zeile fehlt in der Summe K24 und damit in Deckbl UPV E26."
    if not werte_da:
        return b
    # bei leerem Einzelpreis ist das Produkt 0 – ohne Wirkung
    produkt = z0(m.wert(BE, "F21")) * z0(m.wert(BE, "H21")) * z0(m.wert(BE, "J21"))
    produkt = runde(produkt)
    b.delta = {"E26": produkt, "E28": -produkt}
    b.betrag = -produkt
    b.wirkung = _wirkung(b.betrag)
    if ist_null(produkt):
        b.text += " Zeile 9 ist nicht ausgefüllt – ohne Wirkung."
    else:
        b.text += f" Zeile 9 ergibt {eur(produkt)}; die Beseitigungskosten sind um diesen Betrag zu niedrig, der Gesamtbetrag zu hoch."
    return b


# --- 6. Baulichkeiten!F104/G104 (Beispielwert „nicht beräumt“) ------------------------------------

def pruefe_beispiel_beraeumt(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("F104", "Beispielwert der Vorlage „nicht beräumt“ (5 m³ × 50 €) unverändert", f"{BL}!F104, {BL}!G104", "beispiel", "", "unbekannt")
    f, g = m.roh(BL, "F104"), m.roh(BL, "G104")
    gleich = (not isinstance(f, str) and f == 5 and not isinstance(g, str) and g == 50)
    b.zustand = "beispielwert" if gleich else "korrekt"
    if not gleich:
        return b
    b.text = "F104 = 5 m³ und G104 = 50 € sind die Beispielwerte der Vorlage. Waren tatsächlich 5 m³ nicht beräumt, ist das richtig; sonst werden 250 € zu Unrecht abgezogen."
    if not werte_da:
        return b
    h104 = m.n(BL, "H104")
    if h104 is None:
        return b
    b.delta = {"E25": -h104, "E28": h104}
    b.betrag = h104
    b.wirkung = _wirkung(h104)
    return b


# --- Hinweis: I19/I28/I37 (nicht berechnete Baufläche) ------------------------------------------

def pruefe_beispiel_flaeche(m: Mappe, werte_da: bool, warn: list[str]) -> Befund:
    b = Befund("I19", "Beispielwerte „nicht berechnete Baufläche“ (1/2/3 m²) unverändert", f"{BL}!I19, {BL}!I28, {BL}!I37", "hinweis", "", "keine")
    ist = tuple(m.roh(BL, c) for c in ("I19", "I28", "I37"))
    gleich = all(not isinstance(v, str) and v == s for v, s in zip(ist, (1, 2, 3)))
    b.zustand = "beispielwert" if gleich else "korrekt"
    if gleich:
        b.text = "Die drei Zellen haben noch genau die Beispielwerte der Vorlage (1, 2, 3 m²). Sie gehen nur als Fläche in das Protokoll ein, nicht in einen Euro-Betrag."
    return b


PRUEFUNGEN = [pruefe_c33, pruefe_k81, pruefe_j52, pruefe_c64, pruefe_k21, pruefe_beispiel_beraeumt, pruefe_beispiel_flaeche]


# ---------------------------------------------------------------------------------------------
# Eine Datei prüfen
# ---------------------------------------------------------------------------------------------

def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def pruefe_datei(pfad: str, basis: str | None = None) -> DateiErgebnis:
    rel = os.path.relpath(pfad, basis) if basis else os.path.basename(pfad)
    erg = DateiErgebnis(pfad=pfad, rel=rel)
    try:
        m = lese_mappe(pfad)
    except KeineVorlage as e:
        erg.status, erg.grund = "keine_vorlage", f"Keine WEP-Vorlage ({e})"
        return erg
    except zipfile.BadZipFile:
        erg.status, erg.grund = "nicht_lesbar", "Keine gültige Excel-Datei (beschädigt oder verschlüsselt)"
        return erg
    except PermissionError:
        erg.status, erg.grund = "nicht_lesbar", "Zugriff verweigert (Datei evtl. in Excel geöffnet?)"
        return erg
    except Exception as e:  # noqa: BLE001 – eine kaputte Datei darf den Lauf nicht abbrechen
        erg.status, erg.grund = "nicht_lesbar", f"Nicht lesbar: {type(e).__name__}: {e}"
        return erg

    # Erkennung: das Deckblatt muss auf Baulichkeiten!K81 und Beseitigungskosten!K24 zeigen
    if norm_formel(m.roh(DB, "E21")) != "BAULICHKEITEN!K81" or norm_formel(m.roh(DB, "E26")) != "BESEITIGUNGSKOSTEN!K24":
        erg.status, erg.grund = "keine_vorlage", "Keine (passende) WEP-Vorlage: Deckblatt verweist nicht auf Baulichkeiten!K81 / Beseitigungskosten!K24"
        return erg

    erg.anlage = _text(m.wert(FO, "E3") if m.wert(FO, "E3") is not None else m.roh(FO, "E3"))
    parz = _text(m.wert(FO, "E4") if m.wert(FO, "E4") is not None else m.roh(FO, "E4"))
    weg = _text(m.wert(FO, "E5") if m.wert(FO, "E5") is not None else m.roh(FO, "E5"))
    erg.parzelle = f"{parz} / {weg}" if weg and parz else (parz or weg)
    erg.vorgang = _text(m.wert(FO, "E17") if m.wert(FO, "E17") is not None else m.roh(FO, "E17"))

    werte_da = all(m.n(*k) is not None for k in ((BL, "K81"), (DB, "E21"), (DB, "E28")))
    for pruefung in PRUEFUNGEN:
        erg.befunde.append(pruefung(m, werte_da, erg.warnungen))

    auff = [b for b in erg.befunde if b.auffaellig]
    fehler_formel = [b for b in auff if b.art == "fehler" and b.zustand == "fehlerhaft"]
    if not werte_da:
        erg.status = "werte_fehlen"
        erg.grund = ("Die Datei enthält keine gespeicherten Werte (nie in Excel berechnet/gespeichert). "
                     "Die Wirkung in Euro lässt sich nicht berechnen – Datei in Excel öffnen, speichern und erneut prüfen.")
        if fehler_formel:
            erg.grund += " Fehlerhafte Formeln sind enthalten: " + ", ".join(b.id for b in fehler_formel) + "."
        return erg

    erg.deckblatt = deckblatt(m, erg.befunde)
    if any(b.art == "fehler" and b.zustand == "fehlerhaft" and b.wirkung == "euro" for b in erg.befunde):
        erg.status = "betroffen"
    elif any(b.art == "beispiel" and b.wirkung == "euro" for b in erg.befunde):
        erg.status = "verdacht"
    elif auff:
        erg.status = "hinweis"
    else:
        erg.status = "ohne_befund"
    return erg


def deckblatt(m: Mappe, befunde: list[Befund]) -> dict[str, dict[str, Decimal | None]]:
    """Ist (gespeichert) und Soll (Formelfehler korrigiert) für Deckbl UPV E21/E25/E26/E28.
    „soll_verdacht“ nimmt zusätzlich den unveränderten Beispielwert „nicht beräumt“ heraus."""
    summe: dict[str, Decimal] = {}
    verdacht: dict[str, Decimal] = {}
    for b in befunde:
        if b.zustand not in ("fehlerhaft", "beispielwert") or b.betrag is None:
            continue
        ziel = summe if b.art == "fehler" else verdacht
        for zelle, d in b.delta.items():
            ziel[zelle] = ziel.get(zelle, NULL) + d
    # Änderungen am Gesamtbetrag ergeben sich aus den Teilen
    def e28_delta(d: dict[str, Decimal]) -> Decimal:
        return d.get("E21", NULL) - d.get("E25", NULL) - d.get("E26", NULL)
    for d in (summe, verdacht):
        d.pop("E28", None)
    ergebnis: dict[str, dict[str, Decimal | None]] = {}
    for zelle in ("E21", "E25", "E26", "E28"):
        ist = m.n(DB, zelle)
        if zelle == "E28":
            d1, d2 = e28_delta(summe), e28_delta(summe) + e28_delta(verdacht)
        else:
            d1, d2 = summe.get(zelle, NULL), summe.get(zelle, NULL) + verdacht.get(zelle, NULL)
        ergebnis[zelle] = {
            "ist": ist,
            "soll": None if ist is None else ist + d1,
            "soll_verdacht": None if ist is None else ist + d2,
        }
    return ergebnis


# ---------------------------------------------------------------------------------------------
# Ordner prüfen
# ---------------------------------------------------------------------------------------------

def finde_dateien(ordner: str) -> list[str]:
    ergebnis = []
    for wurzel, unterordner, dateien in os.walk(ordner):
        unterordner[:] = sorted(d for d in unterordner if not d.startswith("."))
        for name in sorted(dateien):
            if name.startswith("~$") or name.startswith("."):
                continue
            if name.lower().endswith((".xlsx", ".xlsm")):
                ergebnis.append(os.path.join(wurzel, name))
    return ergebnis


def pruefe_ordner(ordner: str, ausgabe=print) -> list[DateiErgebnis]:
    dateien = finde_dateien(ordner)
    ergebnisse = []
    for i, pfad in enumerate(dateien, 1):
        ausgabe(f"[{i}/{len(dateien)}] {os.path.relpath(pfad, ordner)}")
        ergebnisse.append(pruefe_datei(pfad, ordner))
    return ergebnisse


def sortiert(ergebnisse: list[DateiErgebnis]) -> list[DateiErgebnis]:
    def schluessel(e: DateiErgebnis):
        return (STATUS_REIHE.index(e.status), -abs(e.gesamt_diff + (e.verdacht_diff if e.status == "verdacht" else NULL)), e.rel.casefold())
    return sorted(ergebnisse, key=schluessel)


# ---------------------------------------------------------------------------------------------
# Bericht
# ---------------------------------------------------------------------------------------------

def richtung(d: Decimal | None) -> str:
    """Soll − Ist am Gesamtbetrag in Worten."""
    if d is None or ist_null(d):
        return "keine Abweichung"
    return f"Gesamtbetrag zu hoch um {eur(abs(d))}" if d < 0 else f"Gesamtbetrag zu niedrig um {eur(d)}"


def befund_kurz(b: Befund) -> str:
    if b.wirkung == "euro" and b.art == "beispiel":
        return f"{b.id}: Beispielwert unverändert – Verdacht: {richtung(b.betrag)}"
    if b.wirkung == "euro":
        return f"{b.id}: {richtung(b.betrag)}"
    if b.wirkung == "anzeige":
        return f"{b.id}: falsche Beschriftung (kein Betrag)"
    if b.zustand == "abweichend":
        return f"{b.id}: Formel von Hand geändert"
    if b.zustand == "beispielwert":
        return f"{b.id}: Beispielwert unverändert"
    return f"{b.id}: fehlerhafte Formel, ohne Wirkung" + ("" if b.wirkung == "keine" else " (nicht berechenbar)")


WIRKUNG_TEXT = {"euro": "wirkt auf das Ergebnis", "anzeige": "nur Anzeige", "keine": "ohne Wirkung", "unbekannt": "nicht berechenbar"}
ZUSTAND_TEXT = {"fehlerhaft": "fehlerhafte Formel", "abweichend": "Formel geändert", "beispielwert": "Beispielwert unverändert", "korrekt": "korrekt"}

CSS = """
body{font-family:Segoe UI,Arial,sans-serif;max-width:1100px;margin:1.5rem auto;padding:0 1rem;color:#1b1b1b;line-height:1.4}
h1{font-size:1.6rem} h2{font-size:1.25rem;margin-top:2rem;border-bottom:2px solid #ccc;padding-bottom:.2rem} h3{font-size:1.05rem;margin:1.5rem 0 .3rem}
table{border-collapse:collapse;width:100%;margin:.5rem 0;font-size:.92rem} th,td{border:1px solid #bbb;padding:.3rem .5rem;text-align:left;vertical-align:top}
th{background:#eee} td.zahl,th.zahl{text-align:right;white-space:nowrap}
.warn{background:#fff3cd;border:1px solid #e0b100;padding:.6rem .8rem;margin:1rem 0}
.rot{color:#a40000;font-weight:600} .grau{color:#555}
ul{margin:.2rem 0;padding-left:1.2rem}
@media print{body{max-width:none}h3{page-break-after:avoid}}
"""


def _td(s: Any, klasse: str = "") -> str:
    return f"<td{' class=\"' + klasse + '\"' if klasse else ''}>{html.escape(str(s))}</td>"


def schreibe_html(pfad: str, ordner: str, ergebnisse: list[DateiErgebnis]) -> None:
    erg = sortiert(ergebnisse)
    zaehl = {s: sum(1 for e in erg if e.status == s) for s in STATUS_REIHE}
    h: list[str] = []
    a = h.append
    a(f"<!doctype html><html lang='de'><head><meta charset='utf-8'><title>Prüfbericht WEP-Vorlagenfehler</title><style>{CSS}</style></head><body>")
    a("<h1>Prüfbericht: Vorlagenfehler in WEPs</h1>")
    a(f"<p class='grau'>Geprüfter Ordner: {html.escape(os.path.abspath(ordner))}<br>Erstellt am {datetime.now():%d.%m.%Y %H:%M} · pruefe_weps {__version__}</p>")
    a("<div class='warn'><strong>Achtung, Personendaten:</strong> Dieser Bericht nennt Kleingartenanlage, Parzelle und Vorgangs-Nr. der geprüften WEPs. "
      "Nicht öffentlich ablegen und nur an Berechtigte weitergeben. Die Originaldateien wurden nur gelesen, nicht verändert; es wurde nichts ins Netz gesendet.</div>")

    a("<h2>Zusammenfassung</h2><ul>")
    a(f"<li>{len(erg)} Datei(en) geprüft</li>")
    for s in STATUS_REIHE:
        if zaehl[s]:
            klasse = " class='rot'" if s == "betroffen" else ""
            a(f"<li{klasse}>{zaehl[s]} × {html.escape(STATUS_TEXT[s])}</li>")
    a("</ul>")
    betroffen = [e for e in erg if e.status in ("betroffen", "verdacht")]
    if betroffen:
        summe = sum((e.gesamt_diff for e in betroffen), NULL)
        a(f"<p>Summe der Abweichungen durch Formelfehler über alle Dateien (Soll − Ist am Gesamtbetrag): <strong>{eur(summe, True)}</strong>. "
          "Minus = der Gesamtbetrag im Protokoll ist zu hoch, plus = zu niedrig.</p>")

    a("<h2>Betroffene WEPs</h2>")
    if not betroffen:
        a("<p>Keine Datei mit Auswirkung auf das Ergebnis gefunden.</p>")
    else:
        a("<table><tr><th>Datei</th><th>Anlage</th><th>Parzelle</th><th>Vorgangs-Nr.</th><th>Gefundene Fehler</th><th class='zahl'>Abweichung Gesamtbetrag</th></tr>")
        for e in betroffen:
            fehler = "<ul>" + "".join(f"<li>{html.escape(befund_kurz(b))}</li>" for b in e.befunde
                                      if b.auffaellig and (b.wirkung == "euro" or b.art != "hinweis")) + "</ul>"
            diff = eur(e.gesamt_diff, True) if e.status == "betroffen" else "Verdacht: " + eur(e.verdacht_diff, True)
            a(f"<tr>{_td(e.rel)}{_td(e.anlage)}{_td(e.parzelle)}{_td(e.vorgang)}<td>{fehler}</td>{_td(diff, 'zahl')}</tr>")
        a("</table>")
        a("<p class='grau'>„Abweichung“ = was der Gesamtbetrag mit richtiger Formel wäre, minus was in der Datei steht. "
          "„Verdacht“: nur der Beispielwert „nicht beräumt“ (5 m³ × 50 €) ist unverändert – prüfen, ob er tatsächlich so gemeint war.</p>")

    a("<h2>Einzelheiten je Datei</h2>")
    for e in erg:
        if e.status not in ("betroffen", "verdacht", "hinweis"):
            continue
        a(f"<h3>{html.escape(e.rel)}</h3>")
        a(f"<p>Anlage: {html.escape(e.anlage or '–')} · Parzelle: {html.escape(e.parzelle or '–')} · Vorgangs-Nr.: {html.escape(e.vorgang or '–')} · Status: {html.escape(STATUS_TEXT[e.status])}</p>")
        a("<table><tr><th>Stelle</th><th>Zelle</th><th>Befund</th><th>Wirkung</th><th class='zahl'>Betrag (Soll − Ist)</th></tr>")
        for b in e.befunde:
            if not b.auffaellig:
                continue
            betrag = eur(b.betrag, True) if b.betrag is not None and b.wirkung == "euro" else "–"
            a(f"<tr>{_td(b.stelle)}{_td(b.zellen)}<td>{html.escape(ZUSTAND_TEXT[b.zustand])}<br><span class='grau'>{html.escape(b.text)}</span></td>{_td(WIRKUNG_TEXT[b.wirkung])}{_td(betrag, 'zahl')}</tr>")
        a("</table>")
        if e.deckblatt:
            verdacht = any(b.art == "beispiel" and b.wirkung == "euro" for b in e.befunde)
            namen = {"E21": "Baulichkeiten (E21)", "E25": "Kosten nicht beräumter Baulichkeiten (E25)", "E26": "Beseitigungskosten Bau (E26)", "E28": "Gesamtbetrag mit Abzügen (E28)"}
            a("<table><tr><th>Deckblatt „Deckbl UPV“</th><th class='zahl'>in der Datei</th><th class='zahl'>mit richtigen Formeln</th><th class='zahl'>Differenz</th>"
              + ("<th class='zahl'>zusätzlich ohne Beispielwert</th>" if verdacht else "") + "</tr>")
            for zelle, w in e.deckblatt.items():
                d = None if w["ist"] is None else w["soll"] - w["ist"]
                a(f"<tr>{_td(namen[zelle])}{_td(eur(w['ist']), 'zahl')}{_td(eur(w['soll']), 'zahl')}{_td(eur(d, True), 'zahl')}"
                  + (_td(eur(w['soll_verdacht']), 'zahl') if verdacht else "") + "</tr>")
            a("</table>")
        for w in e.warnungen:
            a(f"<p class='rot'>Hinweis: {html.escape(w)}</p>")

    rest = [e for e in erg if e.status in ("werte_fehlen", "ohne_befund", "keine_vorlage", "nicht_lesbar")]
    a("<h2>Weitere Dateien</h2>")
    if not rest:
        a("<p>Keine.</p>")
    else:
        a("<table><tr><th>Datei</th><th>Anlage</th><th>Parzelle</th><th>Vorgangs-Nr.</th><th>Ergebnis</th></tr>")
        for e in rest:
            text = STATUS_TEXT[e.status] + (f": {e.grund}" if e.grund else "")
            if e.warnungen:
                text += " · Hinweis: " + " ".join(e.warnungen)
            a(f"<tr>{_td(e.rel)}{_td(e.anlage)}{_td(e.parzelle)}{_td(e.vorgang)}{_td(text)}</tr>")
        a("</table>")

    a("<h2>Was wird geprüft?</h2><ul>")
    a("<li><b>C33</b> (Baulichkeiten): Anbau A2 rechnete mit dem Ergebnis von Anbau A1.</li>")
    a("<li><b>K81</b> (Baulichkeiten): der Schuppen wurde in „Baulichkeiten insgesamt“ doppelt gezählt.</li>")
    a("<li><b>J52</b> (Baulichkeiten): falsche Beschriftung beim Schuppen (kein Betrag).</li>")
    a("<li><b>C64</b> (Baulichkeiten): beim integrierten überdachten Laubenvorplatz ging der Wert verloren.</li>")
    a("<li><b>K21</b> (Beseitigungskosten): Zeile 9 fehlte in der Summe.</li>")
    a("<li><b>F104/G104</b> (Baulichkeiten): Beispielwert „nicht beräumt“ 5 m³ × 50 € unverändert (Verdacht, nur wenn nicht tatsächlich so gemeint).</li>")
    a("<li><b>I19/I28/I37</b> (Baulichkeiten): Beispielwerte 1/2/3 m² „nicht berechnete Baufläche“ (nur Hinweis, kein Betrag).</li></ul>")
    a("</body></html>")
    with open(pfad, "w", encoding="utf-8") as f:
        f.write("\n".join(h))


CSV_SPALTEN = ["Datei", "Status", "Anlage", "Parzelle", "Vorgangs-Nr.", "Stelle", "Zelle", "Befund", "Wirkung", "Betrag Soll-Ist (EUR)",
               "Datei gesamt Soll-Ist (EUR)", "Deckblatt E21 Ist", "Deckblatt E21 Soll", "Deckblatt E25 Ist", "Deckblatt E25 Soll",
               "Deckblatt E26 Ist", "Deckblatt E26 Soll", "Deckblatt E28 Ist", "Deckblatt E28 Soll", "Anmerkung"]


def schreibe_csv(pfad: str, ergebnisse: list[DateiErgebnis]) -> None:
    # UTF-8 mit BOM und Semikolon: so öffnet das deutsche Excel die Datei mit Umlauten und Dezimalkomma
    with open(pfad, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        w.writerow(CSV_SPALTEN)
        for e in sortiert(ergebnisse):
            kopf = [csv_text(e.rel), STATUS_TEXT[e.status], csv_text(e.anlage), csv_text(e.parzelle), csv_text(e.vorgang)]
            db: list[str] = []
            for zelle in ("E21", "E25", "E26", "E28"):
                x = e.deckblatt.get(zelle) if e.deckblatt else None
                db += [csv_zahl(x["ist"]) if x else "", csv_zahl(x["soll"]) if x else ""]
            gesamt = csv_zahl(e.gesamt_diff) if e.deckblatt else ""
            zeilen = [b for b in e.befunde if b.auffaellig]
            if not zeilen:
                w.writerow(kopf + ["", "", "", "", "", gesamt] + db + [csv_text(e.grund + (" " + " ".join(e.warnungen) if e.warnungen else ""))])
                continue
            for b in zeilen:
                w.writerow(kopf + [csv_text(b.stelle), csv_text(b.zellen), ZUSTAND_TEXT[b.zustand], WIRKUNG_TEXT[b.wirkung],
                                   csv_zahl(b.betrag) if b.wirkung == "euro" else "", gesamt] + db + [csv_text(b.text)])


# ---------------------------------------------------------------------------------------------
# Kommandozeile
# ---------------------------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    for strom in (sys.stdout, sys.stderr):
        try:
            strom.reconfigure(errors="replace")  # type: ignore[attr-defined]
        except Exception:  # noqa: BLE001
            pass
    p = argparse.ArgumentParser(description="Prüft WEP-Excel-Dateien auf die Fehler der WEP-Vorlage (nur lesend).")
    p.add_argument("ordner", nargs="?", help="Ordner mit den WEPs (.xlsx/.xlsm, auch Unterordner)")
    p.add_argument("--ausgabe", help="Ordner für pruefbericht.html/.csv (Standard: der geprüfte Ordner)")
    args = p.parse_args(argv)

    ordner = args.ordner or input("Ordner mit den WEPs: ").strip().strip('"')
    if not os.path.isdir(ordner):
        print(f"Fehler: Ordner nicht gefunden: {ordner}", file=sys.stderr)
        return 2
    ausgabe = args.ausgabe or ordner
    try:
        os.makedirs(ausgabe, exist_ok=True)
        probe = os.path.join(ausgabe, BERICHT_HTML)
        open(probe, "a").close()
    except OSError as e:
        print(f"Fehler: Im Ordner {ausgabe} kann nicht geschrieben werden ({e}). Mit --ausgabe einen anderen Ordner wählen.", file=sys.stderr)
        return 2

    ergebnisse = pruefe_ordner(ordner)
    schreibe_html(os.path.join(ausgabe, BERICHT_HTML), ordner, ergebnisse)
    schreibe_csv(os.path.join(ausgabe, BERICHT_CSV), ergebnisse)

    zaehl = {s: sum(1 for e in ergebnisse if e.status == s) for s in STATUS_REIHE}
    print()
    print(f"{len(ergebnisse)} Datei(en) geprüft.")
    for s in STATUS_REIHE:
        if zaehl[s]:
            print(f"  {zaehl[s]:4d} × {STATUS_TEXT[s]}")
    print(f"Bericht: {os.path.join(ausgabe, BERICHT_HTML)}")
    print(f"         {os.path.join(ausgabe, BERICHT_CSV)}")
    print("Hinweis: Der Bericht enthält Personendaten (Anlage, Parzelle, Vorgangs-Nr.) – nicht weitergeben.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
